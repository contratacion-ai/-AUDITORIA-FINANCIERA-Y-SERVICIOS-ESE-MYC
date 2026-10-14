"""Exportación de informes y conciliaciones a Excel (.xlsx) con identidad profesional.

Cada hoja lleva: banda de título (institución + nombre de la hoja), línea de autoría «Winston Gerlein Gutiérrez Pastrana – Profesional de Costos»,
encabezado y pie de impresión con la misma identidad, paginación y formato coherente (Calibri, paleta azul sobria)."""
import os, tempfile
from datetime import datetime
import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from . import analysis as A, costos as C
from .db import AUTOR, AZUL, AZUL_CLARO, AZUL_MED, CARGO, FUENTE, GRIS, GRIS_CLARO, IDENTIDAD, INSTITUCION, NIT, VERSION, df

TOP = 5            # filas ocupadas por el bloque de título antes de la tabla (títulos en 1-3, fila 4 en blanco, encabezado en la 5)
DESCRIPCION = {
    "Resumen_Mensual": "Serie mensual de facturación, RIPS y costos.", "Resumen_Empresa": "Facturación y brecha Fact–RIPS desagregadas por empresa (EPS).",
    "Resumen_Empresa_Mes": "Facturación por empresa y mes.", "Comparativo_Sede_Mes": "Facturación por sede y mes con variación del último mes.",
    "Comparativo_Servicio_Mes": "Facturación por servicio y mes con variación del último mes.", "Comparativo_EPS_Mes": "Facturación por EPS y mes.",
    "Comparativo_Contrato_Mes": "Facturación por contrato y mes.", "Validacion_Cruce": "Cobertura de llaves comunes entre las cuatro fuentes.",
    "Contratos": "Comportamiento de contratos.", "Radicacion": "Radicación factura a factura (360 ↔ RIPS).", "Notas_Tecnicas": "Nota técnica vs facturación, RIPS y costo unitario.",
    "Hallazgos": "Reglas de auditoría de línea R01–R09.", "Recomendaciones": "Acciones sugeridas por la herramienta.", "Costos_Cruce": "Cruce Costos × RIPS × Facturación × Notas técnicas por CUPS.",
    "Costos_Validaciones": "Errores y advertencias detectados al cargar costos.", "Regla_Actividad": "Mapa ACTIVIDAD → CUPS usado para identificar el CUPS de cada registro.",
}

def hojas(mes=None):
    m = A.mensual(); h = {"Resumen_Mensual": m if not mes else m[m.mes == mes]}
    h["Resumen_Empresa"] = A.resumen_empresa(mes); h["Resumen_Empresa_Mes"] = A.resumen_empresa_mensual(mes)
    for n, t in (("Sede", "sede"), ("Servicio", "servicio"), ("EPS", "eps"), ("Contrato", "contrato")): h[f"Comparativo_{n}_Mes"] = A.comparativo_mensual(t)
    for k in A.INFORMES:
        h["INF_" + k[:26]] = (pd.DataFrame({"info": ["El informe 1 (detalle) se exporta solo al elegir un mes"]}) if k == "1_general" and not mes else A.informe(k, mes))
    h["Validacion_Cruce"] = A.cobertura(mes)
    for n in ("eps", "contrato", "modalidad", "sede", "servicio", "cups"): h["FvR_" + n] = A.fact_vs_rips(mes, n)
    for n in ("servicio", "sede", "cups"): h["TriEje_" + n] = A.tri_eje(mes, n)
    h["Contratos"] = A.contratos(mes); h["Radicacion"] = A.radicacion(mes); h["Notas_Tecnicas"] = A.notas_vs_fact(mes); h["Hallazgos"] = A.hallazgos(mes)
    h["Costos_Cruce"] = C.cruce(mes); h["Costos_Validaciones"] = C.validaciones()
    h["Recomendaciones"] = pd.DataFrame(A.recomendaciones(mes))
    h["Regla_Actividad"] = df("SELECT mes, actividad, cups, lineas FROM regla_actividad" + (" WHERE mes=?" if mes else "") + " ORDER BY lineas DESC", (mes,) if mes else ())
    h["Param_Contrato"] = df("SELECT * FROM param_contrato"); h["Param_Servicio"] = df("SELECT * FROM param_servicio")
    h["Historial_Cargas"] = df("SELECT * FROM carga ORDER BY id DESC")
    return h

_borde = Border(bottom=Side(style="hair", color="BFBFBF"))

def _identidad(ws, titulo, mes, ncols):
    """Banda de título + autoría en las filas 1-3, encabezado y pie de impresión."""
    ncols = max(ncols, 6)
    for r, (txt, font, fill) in enumerate([
            (f"{INSTITUCION} · NIT {NIT}", Font(name=FUENTE, bold=True, size=13, color="FFFFFF"), AZUL),
            (f"{titulo.replace('_', ' ')} · {'Periodo ' + mes if mes else 'Todos los periodos'}", Font(name=FUENTE, bold=True, size=11, color="FFFFFF"), AZUL_MED),
            (f"{IDENTIDAD} · Generado {datetime.now():%d/%m/%Y %H:%M} · Auditoría v{VERSION}", Font(name=FUENTE, italic=True, size=9, color=GRIS), GRIS_CLARO)], start=1):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        c = ws.cell(r, 1, txt); c.font = font; c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1); ws.row_dimensions[r].height = 22 if r < 3 else 16
    ws.oddHeader.left.text = f"&B{INSTITUCION}&B"; ws.oddHeader.left.size = 9
    ws.oddHeader.right.text = f"{titulo.replace('_', ' ')}"; ws.oddHeader.right.size = 9
    ws.oddFooter.left.text = IDENTIDAD; ws.oddFooter.left.size = 8
    ws.oddFooter.center.text = "Documento de apoyo a la auditoría · uso interno"; ws.oddFooter.center.size = 8
    ws.oddFooter.right.text = "Página &P de &N"; ws.oddFooter.right.size = 8
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.print_title_rows = f"{TOP}:{TOP}"
    ws.sheet_view.showGridLines = False

def _estilo(ws, titulo, mes):
    ncols = ws.max_column; _identidad(ws, titulo, mes, ncols)
    for c in ws[TOP]:
        c.font = Font(name=FUENTE, bold=True, color="FFFFFF", size=10); c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[TOP].height = 30
    ws.freeze_panes = ws.cell(TOP + 1, 1)
    if ws.max_row > TOP: ws.auto_filter.ref = f"A{TOP}:{get_column_letter(ncols)}{ws.max_row}"
    grande = ws.max_row > 20000     # hojas enormes: solo encabezado con estilo (rapidez)
    for col in ws.iter_cols(min_row=TOP, max_row=ws.max_row):
        head = str(col[0].value or ""); letra = get_column_letter(col[0].column)
        ws.column_dimensions[letra].width = min(max(len(head) + 2, 12), 45)
        if grande: continue
        for c in col[1:]:
            c.font = Font(name=FUENTE, size=10); c.border = _borde
            if (c.row - TOP) % 2 == 0: c.fill = PatternFill("solid", fgColor=GRIS_CLARO)
            if isinstance(c.value, float): c.number_format = "0.0%" if head.startswith(("pct", "dif_", "cobertura")) else "#,##0.00"
            elif isinstance(c.value, int) and not head.endswith(("id", "mes", "fila")): c.number_format = "#,##0"
    # colorea el estado (semáforo)
    cols = {str(c.value): c.column for c in ws[TOP]}
    if "estado" in cols and not grande:
        for r in range(TOP + 1, ws.max_row + 1):
            c = ws.cell(r, cols["estado"]); v = str(c.value or "").upper()
            col = "C00000" if ("CRÍTICO" in v or "EXCEDE" in v or "NEGATIVO" in v or "SIN COSTO" in v) else "BF8F00" if (v and v != "OK" and "OK" not in v[:2]) else "548235"
            c.font = Font(name=FUENTE, size=10, bold=True, color=col)

def _grafica_horizontal(ws, titulo, col_etiqueta, col_valor, desde, hasta, ancla):
    ch = BarChart(); ch.type = "bar"; ch.style = 10; ch.title = titulo; ch.legend = None
    ch.add_data(Reference(ws, min_col=col_valor, min_row=TOP, max_row=hasta), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=col_etiqueta, min_row=desde, max_row=hasta))
    ch.y_axis.numFmt = '#,##0'; ch.x_axis.scaling.orientation = "maxMin"; ch.height = 11; ch.width = 22
    ch.y_axis.delete = False; ch.x_axis.delete = False
    ws.add_chart(ch, ancla)

def _portada(wb, mes, hojas_):
    ws = wb.create_sheet("Portada", 0); ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3; ws.column_dimensions["B"].width = 34; ws.column_dimensions["C"].width = 70
    for r in range(1, 8):
        for c in range(1, 4): ws.cell(r, c).fill = PatternFill("solid", fgColor=AZUL)
    ws["B2"] = "AUDITORÍA DE FACTURACIÓN, RIPS Y COSTOS"; ws["B2"].font = Font(name=FUENTE, bold=True, size=20, color="FFFFFF")
    ws["B3"] = INSTITUCION; ws["B3"].font = Font(name=FUENTE, size=13, color="FFFFFF")
    ws["B5"] = f"NIT {NIT} · Tame, Arauca"; ws["B5"].font = Font(name=FUENTE, size=11, color=AZUL_CLARO)
    datos = [("Elaborado por", AUTOR), ("Cargo", CARGO), ("Periodo", mes or "Todos los periodos cargados"),
             ("Fecha de generación", f"{datetime.now():%d/%m/%Y %H:%M}"), ("Versión de la herramienta", VERSION),
             ("Fuentes", "Software 360 (facturación por ámbito) · Consolidado RIPS · Costos mensuales · Notas técnicas EPS")]
    for i, (k, v) in enumerate(datos, start=9):
        ws.cell(i, 2, k).font = Font(name=FUENTE, bold=True, size=11, color=AZUL); ws.cell(i, 3, v).font = Font(name=FUENTE, size=11)
        ws.cell(i, 2).border = ws.cell(i, 3).border = _borde
    ws.cell(16, 2, "CONTENIDO").font = Font(name=FUENTE, bold=True, size=12, color=AZUL)
    for i, n in enumerate(hojas_, start=17):
        c = ws.cell(i, 2, n); c.hyperlink = f"#'{n}'!A1"; c.font = Font(name=FUENTE, size=10, underline="single", color=AZUL_MED)
        ws.cell(i, 3, DESCRIPCION.get(n, "")).font = Font(name=FUENTE, size=10, color=GRIS)
    ws.oddFooter.left.text = IDENTIDAD; ws.oddFooter.right.text = "Página &P de &N"
    ws.page_setup.orientation = "portrait"; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToHeight = 0

def exportar(mes=None, solo=None):
    """Genera el libro. `solo`: nombre de una hoja (o prefijo) para exportar un único informe."""
    h = hojas(mes)
    if solo: h = {k: v for k, v in h.items() if k.lower().startswith(solo.lower())}
    if not h: raise KeyError(solo)
    path = os.path.join(tempfile.mkdtemp(), f"auditoria_ESE_MyC_{mes or 'consolidado'}.xlsx")
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for nombre, d in h.items():
            (d if not d.empty else pd.DataFrame({"info": ["Sin datos cargados para esta sección"]})).to_excel(xw, sheet_name=nombre[:31], index=False, startrow=TOP - 1)
        wb = xw.book
        for ws in list(wb.worksheets):
            _estilo(ws, ws.title, mes)
            if ws.title.startswith("INF_3_servicio") and ws.max_row > TOP + 1:     # Informe 3: gráfica horizontal (reemplaza la barra vertical)
                cab = [str(c.value) for c in ws[TOP]]
                if "Total general" in cab and "SERVICIO" in cab:
                    _grafica_horizontal(ws, "Facturación por servicio (top 20)", cab.index("SERVICIO") + 1, cab.index("Total general") + 1, TOP + 1,
                                        min(ws.max_row, TOP + 20), f"{get_column_letter(ws.max_column + 2)}{TOP}")
        if not solo: _portada(wb, mes, [ws.title for ws in wb.worksheets])
    return path
