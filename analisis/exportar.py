"""Exportación a Excel (openpyxl) del informe, matrices de distribución, costos CUPS e inconsistencias."""
import io
import re
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL = PatternFill("solid", fgColor="1F3864")
GRIS = PatternFill("solid", fgColor="D9E1F2")
BLANCO = Font(color="FFFFFF", bold=True)
NEGRITA = Font(bold=True)
BORDE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
MONEDA = '#,##0;[Red]-#,##0'
PCT = '0.0%'


def _num(v):
    if isinstance(v, Decimal):
        return float(v)
    return v


class Libro:
    def __init__(self):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)

    def hoja(self, titulo, encabezado_titulo, subtitulo, columnas, filas, formatos=None, anchos=None):
        ws = self.wb.create_sheet(re.sub(r"[\\/*?:\[\]]", "-", titulo)[:31])
        ws["A1"] = encabezado_titulo
        ws["A1"].font = Font(bold=True, size=14, color="1F3864")
        ws["A2"] = subtitulo
        ws.append([])
        ws.append(columnas)
        for i in range(1, len(columnas) + 1):
            c = ws.cell(row=4, column=i)
            c.fill, c.font, c.border = AZUL, BLANCO, BORDE
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for fila in filas:
            ws.append([_num(v) for v in fila])
            r = ws.max_row
            for i, fmt in enumerate(formatos or [], start=1):
                if fmt:
                    ws.cell(row=r, column=i).number_format = fmt
            for i in range(1, len(columnas) + 1):
                ws.cell(row=r, column=i).border = BORDE
        for i, ancho in enumerate(anchos or [18] * len(columnas), start=1):
            ws.column_dimensions[get_column_letter(i)].width = ancho
        ws.freeze_panes = "A5"
        return ws

    def total(self, ws):
        for c in ws[ws.max_row]:
            c.font, c.fill = NEGRITA, GRIS

    def bytes(self):
        buf = io.BytesIO()
        self.wb.save(buf)
        return buf.getvalue()


def _p(v):
    return float(v) / 100 if v is not None else None


def informe_excel(d):
    lib = Libro()
    inst, per, comp = d["institucion"].nombre.upper(), d["periodo"], d["periodo_comp"]
    et_a, et_b = per.etiqueta_corta, comp.etiqueta_corta if comp else "—"
    sub = f"Análisis comparativo {et_b} vs {et_a} · Modelo ABC Res. 002/2014"
    ws = lib.hoja("RESUMEN EJECUTIVO", inst, sub,
                  ["ELEMENTO", "RECURSO", f"{et_b} ($)", f"{et_a} ($)", "VAR. $", "VAR. %", f"PART% {et_b}", f"PART% {et_a}"],
                  [[f["elemento"].nombre, f["recurso"].nombre, f["anterior"], f["actual"], f["dif"], _p(f["pct"]),
                    _p(f["part_anterior"]), _p(f["part_actual"])] for f in d["recursos"]]
                  + [["TOTAL GENERAL", "", d["total_anterior"], d["total"], d["total"] - d["total_anterior"], _p(d["variacion"]), 1, 1]],
                  [None, None, MONEDA, MONEDA, MONEDA, PCT, PCT, PCT], [24, 40, 18, 18, 18, 10, 12, 12])
    lib.total(ws)
    filas = []
    for clase in ("FINALES", "INTERMEDIOS", "SERVICIOS", "ESTRUCTURA"):
        for f in d["unidades"][clase]:
            filas.append([clase, f["unidad"].codigo, f["unidad"].nombre, f["anterior"], f["actual"], f["dif"], _p(f["pct"]),
                          _p(f["part_clase"])])
    lib.hoja("UNIDADES SERVICIO", "COSTOS POR UNIDADES DE NEGOCIO", sub + " · costo directo (1ª distribución)",
             ["CLASE", "CÓDIGO", "NOMBRE UNIDAD DE NEGOCIO", f"{et_b} ($)", f"{et_a} ($)", "VAR. $", "VAR. %", "% CLASE"],
             filas, [None, None, None, MONEDA, MONEDA, MONEDA, PCT, PCT], [14, 10, 50, 18, 18, 18, 10, 10])
    ws = lib.hoja("DISTRIBUCIÓN SEDES", "DISTRIBUCIÓN DE COSTOS POR SEDE", sub,
                  ["SEDE / MUNICIPIO"] + [f"{r.nombre} ($)" for r in d["rec_sede"]] + ["TOTAL ($)"],
                  [[s["sede"]] + s["valores"] + [s["total"]] for s in d["sedes"]]
                  + [["TOTAL SEDES"] + d["totales_sede"] + [sum(d["totales_sede"], Decimal(0))]],
                  [None] + [MONEDA] * (len(d["rec_sede"]) + 1), [28] + [18] * (len(d["rec_sede"]) + 1))
    lib.total(ws)
    lib.hoja("BENCHMARKING", "BENCHMARKING · INDICADORES DE REFERENCIA", sub,
             ["INDICADOR", "VALOR", "REFERENCIA MÍN.", "REFERENCIA MÁX.", "ESTADO", "ACCIÓN RECOMENDADA", "FUENTE"],
             [[i["indicador"].nombre, _num(i["valor"]) if i["indicador"].es_valor else _p(i["valor"]),
               _num(i["indicador"].minimo), _num(i["indicador"].maximo), i["estado"], i["indicador"].accion_recomendada,
               i["indicador"].fuente] for i in d["indicadores"]],
             [None, None, None, None, None, None, None], [36, 14, 14, 14, 12, 50, 40])
    lib.hoja("HALLAZGOS", "HALLAZGOS PRINCIPALES", sub, ["NIVEL", "HALLAZGO", "DETALLE"],
             [list(h) for h in d["hallazgos"]], None, [12, 60, 100])
    lib.hoja("RECOMENDACIONES", "PLAN DE ACCIÓN Y RECOMENDACIONES", sub,
             ["PRIORIDAD", "ÁREA", "RECOMENDACIÓN", "IMPACTO ESPERADO", "PLAZO", "RESPONSABLE"], d["recomendaciones"],
             None, [16, 22, 60, 30, 12, 24])
    return lib.bytes()


def matriz_excel(institucion, periodo, recursos, filas, totales, etapa):
    lib = Libro()
    cols = ["CLASE", "UNIDAD DE NEGOCIO", "CÓDIGO CECO", "NOMBRE CENTROS DE COSTOS"] + [f"{r.codigo} {r.nombre}" for r in recursos]
    extra = [x for d, x in ((2, "DE ESTRUCTURA"), (3, "DE SERVICIOS"), (4, "DE INTERMEDIOS")) if d <= etapa]
    cols += extra + ["TOTAL GENERAL", "DISTRIBUIDO", "COSTO NETO"]
    datos = [[f["centro"].clase, f["centro"].unidad_negocio.codigo, f["centro"].codigo, f["centro"].nombre] + f["valores"]
             + f["recibidos"] + [f["total"], f["distribuido"], f["neto"]] for f in filas]
    datos.append(["TOTAL", "", "", ""] + totales["valores"] + totales["recibidos"] + [totales["total"], totales["distribuido"], totales["neto"]])
    nombres = {1: "1ra. DISTRIBUCION", 2: "2da. DISTRIBUCION", 3: "3ra. DISTRIBUCION", 4: "4ta. DISTRIBUCION"}
    ws = lib.hoja(nombres[etapa], f"{institucion.nombre.upper()} · {nombres[etapa]}", f"Periodo {periodo}", cols, datos,
                  [None] * 4 + [MONEDA] * (len(cols) - 4), [12, 12, 10, 40] + [14] * (len(cols) - 4))
    lib.total(ws)
    ws.freeze_panes = "E5"
    return lib.bytes()


def tabla_excel(titulo, subtitulo, columnas, filas, formatos=None):
    lib = Libro()
    lib.hoja(titulo[:31], titulo, subtitulo, columnas, filas, formatos, [20] * len(columnas))
    return lib.bytes()
