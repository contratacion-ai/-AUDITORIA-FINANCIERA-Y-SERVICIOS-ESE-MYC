"""Generación de la plantilla estándar de carga mensual (un libro con una hoja por tipo de archivo).

Se genera a partir de los esquemas, así que siempre coincide con lo que el sistema valida.
"""
import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .esquemas import ESQUEMAS, FECHA, NUMERO

AZUL = PatternFill("solid", fgColor="1F3864")
NARANJA = PatternFill("solid", fgColor="C55A11")

INSTRUCCIONES = [
    ("PLANTILLA ESTÁNDAR DE CARGA MENSUAL · MODELO DE COSTOS ABC", True),
    ("", False),
    ("1. Cada hoja corresponde a un tipo de archivo. Llene solo las hojas que su institución use en el mes.", False),
    ("2. No cambie los encabezados (fila 1). Las columnas en NARANJA son obligatorias; las AZULES son opcionales.", False),
    ("3. Pase el cursor sobre cada encabezado para ver su descripción y un ejemplo. Escriba los datos desde la fila 2.", False),
    ("4. Valores: puede escribir 1.234.567,89 o 1234567.89; el sistema reconoce separadores, $ y negativos entre paréntesis.", False),
    ("5. Fechas: dd/mm/aaaa (también acepta aaaa-mm-dd y fechas de Excel).", False),
    ("6. CENTRO DE COSTOS: use el código del modelo (A10, B01…) o el código contable homologado en Parametrización.", False),
    ("7. CUENTA: el sistema asigna el recurso (R1…R29) por el prefijo más largo configurado en 'Cuentas contables - Recursos'.", False),
    ("8. ESTADO_FINANCIERO es la fuente oficial del costo. NOMINA, CONTRATISTAS, INVENTARIOS, DEPRECIACION y OTROS_AREA", False),
    ("   distribuyen el recurso entre centros; si sus totales difieren del estado financiero, la diferencia se prorratea y se reporta.", False),
    ("9. DOCUMENTO PACIENTE (FACTURACION y RIPS) se seudonimiza al cargar (HMAC-SHA256) y nunca se guarda en claro.", False),
    ("10. Filas duplicadas idénticas se reportan y se omiten. Un archivo idéntico ya aplicado en el mismo periodo se rechaza.", False),
    ("11. En 'Cargas' elija el periodo y el tipo; para este libro use la plantilla 'Estándar' e indique la hoja,", False),
    ("    o use 'Cargar libro estándar completo' para procesar todas las hojas diligenciadas a la vez.", False),
    ("", False),
    ("HOJA · DESCRIPCIÓN", True),
]


def generar_plantilla():
    wb = Workbook()
    ws = wb.active
    ws.title = "INSTRUCCIONES"
    for texto, negrita in INSTRUCCIONES:
        ws.append([texto])
        if negrita:
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12, color="1F3864")
    for esq in ESQUEMAS.values():
        ws.append([f"{esq.codigo} · {esq.nombre}: {esq.descripcion}" + ("" if esq.periodico else " (catálogo, sin periodo)")])
    ws.column_dimensions["A"].width = 140
    for esq in ESQUEMAS.values():
        h = wb.create_sheet(esq.codigo[:31])
        for i, campo in enumerate(esq.campos, start=1):
            celda = h.cell(row=1, column=i, value=campo.etiqueta)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = NARANJA if campo.requerido else AZUL
            celda.alignment = Alignment(horizontal="center", wrap_text=True)
            nota = campo.ayuda or ("Obligatorio." if campo.requerido else "Opcional.")
            if campo.tipo == NUMERO:
                nota += " Numérico."
            elif campo.tipo == FECHA:
                nota += " Fecha dd/mm/aaaa."
            if campo.ejemplo != "":
                nota += f" Ej.: {campo.ejemplo}"
            celda.comment = Comment(nota, "Plantilla ABC")
            h.column_dimensions[get_column_letter(i)].width = max(14, len(campo.etiqueta) + 4)
        h.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
