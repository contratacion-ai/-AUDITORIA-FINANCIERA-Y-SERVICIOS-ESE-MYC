"""Lectura de archivos Excel (.xlsx/.xlsm/.xls) y CSV a una tabla de encabezados + filas."""
import csv
import io
import re

from .limpieza import clave, es_vacio, texto


class ErrorArchivo(Exception):
    pass


def _leer_xlsx(contenido, hoja):
    import openpyxl

    try:
        libro = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as exc:  # archivo corrupto o no es Excel
        raise ErrorArchivo(f"No se pudo abrir el libro de Excel: {exc}") from exc
    try:
        hoja_obj = _hoja(libro.sheetnames, hoja, lambda n: libro[n])
        return [list(f) for f in hoja_obj.iter_rows(values_only=True)]
    finally:
        libro.close()


def _leer_xls(contenido, hoja):
    import xlrd

    try:
        libro = xlrd.open_workbook(file_contents=contenido)
    except Exception as exc:
        raise ErrorArchivo(f"No se pudo abrir el libro .xls: {exc}") from exc
    h = _hoja(libro.sheet_names(), hoja, libro.sheet_by_name)
    filas = []
    for i in range(h.nrows):
        fila = []
        for celda in h.row(i):
            if celda.ctype == xlrd.XL_CELL_DATE:
                fila.append(xlrd.xldate.xldate_as_datetime(celda.value, libro.datemode))
            else:
                fila.append(celda.value)
        filas.append(fila)
    return filas


def _hoja(nombres, hoja, obtener):
    if not hoja:
        return obtener(nombres[0])
    buscada = clave(hoja)
    for n in nombres:
        if clave(n) == buscada:
            return obtener(n)
    raise ErrorArchivo(f"La hoja '{hoja}' no existe. Hojas disponibles: {', '.join(nombres)}")


def _leer_csv(contenido):
    for codificacion in ("utf-8-sig", "latin-1"):
        try:
            texto_csv = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    muestra = texto_csv[:4096]
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=";,\t|")
    except csv.Error:
        dialecto = csv.excel
        dialecto.delimiter = ";" if muestra.count(";") > muestra.count(",") else ","
    return [fila for fila in csv.reader(io.StringIO(texto_csv), dialecto)]


def hojas_del_libro(contenido, nombre):
    ext = nombre.lower().rsplit(".", 1)[-1]
    if ext in ("xlsx", "xlsm"):
        import openpyxl

        libro = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True)
        try:
            return list(libro.sheetnames)
        finally:
            libro.close()
    if ext == "xls":
        import xlrd

        return xlrd.open_workbook(file_contents=contenido, on_demand=True).sheet_names()
    return []


ENCABEZADO_CTRP = ["CUENTA", "NOMBRE CUENTA", "CENTRO DE COSTOS", "DEBITO", "CREDITO", "DESCRIPCION"]


def adaptar_reporte_ctrp(filas, meta=None):
    """Si las filas son el reporte CTRP del software contable (estado de resultados por tercero y centro), las
    convierte al formato de la carga de estado financiero conservando el número de fila. Ver homologacion_ceco.formatos.

    La cuenta de cada fila de detalle se toma de OidCuentaProceso; el tercero queda solo como identificador interno en
    la descripción (distingue filas iguales sin copiar nombres ni documentos). Una cuenta sin detalle por CECO se emite
    sin centro de costo (la reparte la 1.ª distribución). Devuelve las filas originales si no es un reporte CTRP.
    """
    if not filas:
        return filas
    from homologacion_ceco import formatos

    columnas = {clave(v).replace(" ", "") for v in filas[0]}
    if not formatos.COLUMNAS_CTRP <= columnas:
        return filas
    import pandas as pd

    from homologacion_ceco.lectores import Tabla

    ancho = max(len(f) for f in filas)
    tabla = formatos.adaptar(Tabla("", "", pd.DataFrame([list(f) + [None] * (ancho - len(f)) for f in filas], dtype=object), ""))
    i = {c: j for j, c in enumerate(formatos.ENCABEZADO)}
    salida = [list(ENCABEZADO_CTRP)]
    for f in tabla.datos.to_numpy(dtype=object)[1:]:
        f = list(f)
        if all(es_vacio(v) for v in f):
            salida.append([None] * len(ENCABEZADO_CTRP))
            continue
        descripcion = " · ".join(texto(v) for v in (f[i["DESCRIPCION"]], f[i["COMPROBANTE"]]) if not es_vacio(v))
        salida.append([f[i["CUENTA"]], f[i["NOMBRE CUENTA"]], f[i["CODIGO CECO"]], f[i["DEBITO"]], f[i["CREDITO"]], descripcion])
    if meta is not None:
        meta.update(tabla.meta)
    return salida


def leer_tabla(contenido, nombre, hoja="", fila_encabezado=1, meta=None):
    """Devuelve (encabezados, filas, numero_fila_inicial). `meta` (dict) recibe los datos de un formato reconocido."""
    ext = "." + nombre.lower().rsplit(".", 1)[-1] if "." in nombre else ""
    if ext in (".xlsx", ".xlsm"):
        filas = _leer_xlsx(contenido, hoja)
    elif ext == ".xls":
        filas = _leer_xls(contenido, hoja)
    elif ext in (".csv", ".txt"):
        filas = _leer_csv(contenido)
    else:
        raise ErrorArchivo(f"Formato no admitido: {ext or 'sin extensión'}. Use .xlsx, .xls o .csv")
    reconocido = {}
    filas = adaptar_reporte_ctrp(filas, reconocido)
    if reconocido:
        fila_encabezado = 1
        if meta is not None:
            meta.update(reconocido)
    if fila_encabezado < 1 or fila_encabezado > len(filas):
        raise ErrorArchivo(f"El archivo no tiene la fila de encabezado {fila_encabezado} (tiene {len(filas)} filas).")
    encabezados = [texto(c) for c in filas[fila_encabezado - 1]]
    while encabezados and not encabezados[-1]:
        encabezados.pop()
    if not any(encabezados):
        raise ErrorArchivo(f"La fila {fila_encabezado} no contiene encabezados.")
    datos = filas[fila_encabezado:]
    return encabezados, datos, fila_encabezado + 1


def indice_columna(referencia, encabezados):
    """Ubica una columna por nombre (sin importar tildes/mayúsculas) o por letra de Excel con '#': #A, #AB."""
    if referencia is None or es_vacio(referencia):
        return None
    ref = str(referencia).strip()
    letra = re.fullmatch(r"#([A-Za-z]{1,3})", ref)
    if letra:
        n = 0
        for ch in letra.group(1).upper():
            n = n * 26 + (ord(ch) - 64)
        return n - 1
    claves = [clave(e) for e in encabezados]
    buscada = clave(ref)
    if buscada and buscada in claves:
        return claves.index(buscada)
    return None
