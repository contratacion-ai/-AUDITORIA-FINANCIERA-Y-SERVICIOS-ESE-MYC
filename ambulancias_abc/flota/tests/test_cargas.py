import datetime as dt
import io
import json
import zipfile
from decimal import Decimal

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from flota.importacion.rips import APH, TRASLADO, clasificar, leer_rips_archivos
from flota.models import Ambulancia, CodigoTraslado, Consumo, LoteImportacion, Sede, Traslado
from flota.servicios import importar
from flota.servicios.costeo import calcular_mes, facturacion_por_ambito, recorridos_por_sede, tabla_consumos
from flota.servicios.datos_iniciales import cargar

from .utiles import fila_his, xlsx_his

SEP = dt.date(2026, 9, 1)
CATALOGO = {"601T01": TRASLADO, "601T02": TRASLADO, "S31307": TRASLADO}

AF = "500000000101,ESE MORENO Y CLAVIJO,NI,900034131,FR100,15/09/2026,01/09/2026,30/09/2026,EPS037,NUEVA EPS,C1,,,0,0,0,5000000\n"
AT = (
    "FR100,500000000101,CC,1000001,AUT1,2,601T02,TRASLADO ASISTENCIAL BASICO TERRESTRE SECUNDARIO,101,17000,1717000\n"
    "FR100,500000000101,CC,1000002,APH,2,999X01,ATENCION PREHOSPITALARIA,1,350000,350000\n"
    "FR100,500000000101,CC,1000003,AUT3,2,888Z01,SERVICIO SIN NOMBRE CONOCIDO,1,100000,100000\n"
    "FR100,500000000101,CC,1000004,AUT4,1,MAT001,GUANTES,10,1000,10000\n"
)
AP = (
    "FR100,500000000101,CC,1000005,20/09/2026,AUT5,601T01,3,1,1,R509,,,,459300\n"
    "FR100,500000000101,CC,1000006,21/10/2026,AUT6,601T01,2,1,1,R509,,,,459300\n"
    "FR100,500000000101,CC,1000007,21/09/2026,AUT7,890201,1,1,1,R509,,,,40000\n"
)


def zip_rips():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("AF000001.txt", AF)
        z.writestr("AT000001.txt", AT.encode("latin-1"))
        z.writestr("AP000001.txt", AP)
        z.writestr("US000001.txt", "CC,1000001,EPS037,2,APELLIDO,,NOMBRE,,30,1,M,81,794,U\n")
    return buf.getvalue()


class ClasificacionRipsTests(TestCase):
    def test_reglas(self):
        self.assertEqual(clasificar("601T02", catalogo=CATALOGO)[0], TRASLADO)
        self.assertEqual(clasificar("X1", "ATENCIÓN PREHOSPITALARIA BÁSICA")[0], APH)
        self.assertEqual(clasificar("X1", "", "APH")[0], APH)
        self.assertEqual(clasificar("X1", "TRASLADO EN AMBULANCIA")[0], TRASLADO)
        self.assertEqual(clasificar("X1", "OTRO", tipo_servicio="2")[0], TRASLADO)
        self.assertIsNone(clasificar("890201", "CONSULTA")[0])

    def test_lee_zip_de_archivos_planos(self):
        r = leer_rips_archivos([("rips_sep.zip", zip_rips())], "sal", CATALOGO)
        self.assertEqual(r["facturas"], 1)
        self.assertEqual(r["lineas_revisadas"], 7)
        cats = sorted((x["codigo_cups"], x["categoria"]) for x in r["registros"])
        self.assertEqual(cats, [("601T01", TRASLADO), ("601T01", TRASLADO), ("601T02", TRASLADO), ("888Z01", TRASLADO), ("999X01", APH)])
        at = next(x for x in r["registros"] if x["codigo_cups"] == "601T02")
        self.assertEqual(at["fecha_servicio"], SEP)  # AT no trae fecha: inicio del periodo del AF
        self.assertIn("fecha tomada del periodo", at["observaciones"][0])
        ap = next(x for x in r["registros"] if x["codigo_cups"] == "601T01" and x["fecha_servicio"].month == 9)
        self.assertEqual(ap["ambito"], "Urgencias")
        self.assertNotIn("1000001", json.dumps([x["documento_hash"] for x in r["registros"]]))


class BaseDatos(TestCase):
    @classmethod
    def setUpTestData(cls):
        cargar()
        cls.tame = Sede.objects.get(codigo="tame")
        cls.tame.codigo_prestador = "500000000101"
        cls.tame.save()


class CargaRipsTests(BaseDatos):
    def test_carga_por_mes_y_hospital(self):
        lote = importar.importar_rips([("rips_sep.zip", zip_rips())], nombre="rips_sep.zip", mes=SEP, sede=self.tame)
        self.assertEqual(lote.nuevos, 5)
        self.assertEqual((lote.mes, lote.sede), (SEP, self.tame))
        self.assertEqual(lote.detalle["aph"]["lineas"], 1)
        self.assertEqual(lote.detalle["traslados"]["lineas"], 4)
        self.assertEqual(lote.detalle["fuera_del_mes"], 1)
        aph = Traslado.objects.get(categoria=APH)
        self.assertEqual((aph.sede, aph.lote_rips), (self.tame, lote))
        ruta = Traslado.objects.get(codigo_cups="601T02").ruta
        self.assertEqual((ruta.origen, ruta.destino), ("TAME", "SARAVENA"))
        # Recargar no duplica
        lote2 = importar.importar_rips([("rips_sep.zip", zip_rips())], nombre="rips_sep.zip", mes=SEP, sede=self.tame)
        self.assertEqual(lote2.nuevos, 0)
        self.assertEqual(Traslado.objects.count(), 5)

    def test_concilia_con_facturacion_his(self):
        importar.importar_facturacion(xlsx_his([fila_his("HOSPITAL SAN ANTONIO", "601T02", 101, 17000, 77, "FR100")]), nombre="f.xlsx")
        lote = importar.importar_rips([("rips_sep.zip", zip_rips())], nombre="rips_sep.zip", mes=SEP, sede=self.tame)
        self.assertEqual(lote.actualizados, 1)
        his = Traslado.objects.get(origen_registro=Traslado.HIS)
        self.assertTrue(his.conciliado_rips)
        self.assertEqual(his.lote_rips, lote)
        self.assertEqual(Traslado.objects.filter(codigo_cups="601T02").count(), 1)

    def test_codigo_aph_en_catalogo(self):
        CodigoTraslado.objects.create(codigo="888Z01", descripcion="APH avanzada", modalidad=CodigoTraslado.BASE, categoria=APH)
        importar.importar_rips([("rips_sep.zip", zip_rips())], nombre="rips_sep.zip", mes=SEP, sede=self.tame)
        self.assertEqual(Traslado.objects.filter(categoria=APH).count(), 2)


class ConsumosTests(BaseDatos):
    def _xlsx(self, filas, encabezado=("placa", "fecha", "descripcion", "cantidad", "unidad", "valor", "meses_prorrateo", "documento_soporte")):
        wb = Workbook()
        ws = wb.active
        ws.append(list(encabezado))
        for f in filas:
            ws.append(f)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def test_carga_excel_con_tipo_de_la_tabla_y_prorrateo(self):
        arch = self._xlsx([["OEU846", dt.datetime(2026, 3, 15), "SOAT 2026", 1, "und", 1200000, 12, "POL-1"]])
        lote = importar.importar_consumos(arch, nombre="seguros.xlsx", tipo_forzado=Consumo.SEGUROS_IMPUESTOS)
        self.assertEqual(lote.nuevos, 1)
        c = Consumo.objects.get()
        self.assertEqual((c.tipo, c.meses_prorrateo), (Consumo.SEGUROS_IMPUESTOS, 12))
        r = calcular_mes(SEP)
        self.assertAlmostEqual(r.ambulancias["OEU846"]["pool_otros"], 100000)
        self.assertEqual(calcular_mes(dt.date(2027, 3, 1)).ambulancias["OEU846"]["pool_otros"], 0)
        t = tabla_consumos("seguros", dt.date(2026, 2, 1), dt.date(2026, 4, 1))
        self.assertEqual(t["filas"][0]["valores"], [0.0, 100000.0, 100000.0])

    def test_sinonimos_de_tipo(self):
        arch = self._xlsx([["OKZ365", "2026-09-30", "Nómina", 1, "mes", 13754000, 1, "N-9", "PERSONAL"]],
                          ("placa", "fecha", "descripcion", "cantidad", "unidad", "valor", "meses_prorrateo", "documento_soporte", "tipo"))
        importar.importar_consumos(arch, nombre="personal.xlsx")
        self.assertEqual(Consumo.objects.get().tipo, Consumo.TRIPULACION)
        self.assertEqual(calcular_mes(SEP).fuente_pool["OKZ365"]["tripulacion"], "REAL")


class RecorridosSedeTests(BaseDatos):
    def test_totales_cuadran_con_facturacion(self):
        filas = [
            fila_his("HOSPITAL SAN ANTONIO", "601T02", 101, 17000, 1, "A1"),
            fila_his("HOSPITAL SAN ANTONIO", "601T02", 194, 17000, 2, "A2"),
            fila_his("HOSPITAL SAN ANTONIO", "601T02", 71, 17000, 3, "A3"),  # sin ruta en el anexo
            fila_his("HOSPITAL SAN FRANCISCO", "601T02", 45, 17000, 4, "B1"),
        ]
        importar.importar_facturacion(xlsx_his(filas), nombre="f.xlsx", mes=SEP)
        datos = recorridos_por_sede(SEP, SEP)
        total = sum(b["total"]["facturado"] for b in datos)
        self.assertAlmostEqual(total, (101 + 194 + 71 + 45) * 17000, delta=0.01)
        tame = next(b for b in datos if b["sede"].codigo == "tame")
        saravena = next(f for f in tame["filas"] if f["ruta"].destino == "SARAVENA")
        self.assertAlmostEqual(saravena["traslados"], 1, places=9)
        self.assertEqual(saravena["tarifa"], 1717000)
        self.assertGreater(saravena["costo_estandar"], 0)
        self.assertEqual(tame["otras"]["facturado"], 71 * 17000)
        self.assertEqual(tame["n_ambulancias"], 4)
        costo_total = sum(b["total"]["costo_real"] for b in datos)
        self.assertAlmostEqual(costo_total, calcular_mes(SEP).total_asignado, delta=1)

    def test_facturacion_por_ambito(self):
        importar.importar_facturacion(xlsx_his([fila_his("HOSPITAL SAN ANTONIO", "601T02", 101, 17000, 1, "A1", unidad="HOSPITALIZACIÓN")]), nombre="f.xlsx")
        filas = facturacion_por_ambito(SEP, SEP)
        self.assertEqual(filas[0]["unidad_funcional"], "HOSPITALIZACIÓN")
        self.assertEqual(filas[0]["facturado"], Decimal("1717000"))


class VistasCargasTests(BaseDatos):
    def setUp(self):
        self.client.force_login(User.objects.create_superuser("admin", "", "Clave-Prueba-1"))

    def test_boton_cargar_facturacion_por_hospital(self):
        filas = [fila_his("HOSPITAL SAN ANTONIO", "601T02", 101, 17000, 1, "A1"), fila_his("HOSPITAL SAN FRANCISCO", "601T02", 45, 17000, 2, "B1")]
        arch = SimpleUploadedFile("fact.xlsx", xlsx_his(filas).getvalue())
        r = self.client.post(reverse("flota:cargar_facturacion"), {"mes": "2026-09", "sede": self.tame.pk, "archivo": arch})
        self.assertEqual(r.status_code, 302)
        self.assertIn("facturacion", r["Location"])
        self.assertEqual(Traslado.objects.count(), 1)
        lote = LoteImportacion.objects.get(tipo="FACTURACION_HIS")
        self.assertEqual((lote.mes, lote.sede, lote.omitidos), (SEP, self.tame, 1))

    def test_boton_cargar_rips_varios_archivos(self):
        r = self.client.post(reverse("flota:cargar_rips"), {
            "mes": "2026-09", "sede": self.tame.pk,
            "archivos": [SimpleUploadedFile("AF1.txt", AF.encode()), SimpleUploadedFile("AT1.txt", AT.encode()), SimpleUploadedFile("AP1.txt", AP.encode())],
        })
        lote = LoteImportacion.objects.get(tipo="RIPS")
        self.assertRedirects(r, reverse("flota:rips_detalle", args=[lote.pk]))
        self.assertEqual(lote.nuevos, 5)
        r = self.client.get(reverse("flota:rips_detalle", args=[lote.pk]), {"formato": "xlsx"})
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws.max_row, 6)
        self.assertContains(self.client.get(reverse("flota:rips_detalle", args=[lote.pk]), {"categoria": "APH"}), "PREHOSPITALARIA")

    def test_paginas_y_descargas(self):
        importar.importar_facturacion(xlsx_his([fila_his("HOSPITAL SAN ANTONIO", "601T02", 101, 17000, 1, "A1")]), nombre="f.xlsx")
        for nombre, args, params in [
            ("flota:cargas", [], {}), ("flota:cargar_facturacion", [], {}), ("flota:cargar_rips", [], {}),
            ("flota:rips_cargas", [], {}), ("flota:recorridos_sede", [], {"desde": "2026-09", "hasta": "2026-09"}),
            ("flota:facturacion_ambito", [], {"desde": "2026-09", "hasta": "2026-09"}),
            ("flota:dashboard", [], {"categoria": "APH"}),
        ] + [("flota:consumos_tabla", [k], {}) for k in Consumo.TABLAS]:
            self.assertEqual(self.client.get(reverse(nombre, args=args), params).status_code, 200, nombre)
        for nombre, args in [("flota:recorridos_sede", []), ("flota:facturacion_ambito", []), ("flota:consumos_tabla", ["combustible"])]:
            r = self.client.get(reverse(nombre, args=args), {"desde": "2026-09", "hasta": "2026-09", "formato": "xlsx"})
            self.assertEqual(r["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", nombre)
        r = self.client.get(reverse("flota:plantilla_consumos_tabla", args=["seguros"]))
        fila = [c.value for c in load_workbook(io.BytesIO(r.content)).active[2]]
        self.assertEqual(fila[2], Consumo.SEGUROS_IMPUESTOS)

    def test_boton_cargar_tabla_combustible(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["placa", "fecha", "cantidad", "valor", "odometro", "documento_soporte"])
        ws.append(["OEU846", "2026-09-02", 10, 236500, 1000, "V1"])
        ws.append(["OEU846", "2026-09-20", 20, 473000, 1250, "V2"])
        buf = io.BytesIO()
        wb.save(buf)
        r = self.client.post(reverse("flota:cargar_consumos", args=["combustible"]), {"archivo": SimpleUploadedFile("comb.xlsx", buf.getvalue())})
        self.assertRedirects(r, reverse("flota:consumos_tabla", args=["combustible"]))
        self.assertEqual(Consumo.objects.filter(tipo=Consumo.COMBUSTIBLE, unidad="gal").count(), 2)
        t = tabla_consumos("combustible", SEP, SEP, "OEU846")
        self.assertEqual(t["filas"][0]["km_odometro"], 250)
        self.assertAlmostEqual(t["filas"][0]["rendimiento"], 12.5)
        self.assertEqual(Ambulancia.objects.get(placa="OEU846").consumos.count(), 2)


# ---------------------------------------------------------------------------
# RIPS de varios meses / consolidados y auditoría RIPS vs facturación
# ---------------------------------------------------------------------------

PREST_PR = "815910020605"


def _af(factura, inicio, fin):
    return f"{PREST_PR},ESE,NI,900034131,{factura},{fin},{inicio},{fin},EPS037,NUEVA EPS,C1,,,0,0,0,1000000\n"


def _ap(factura, doc, fecha, codigo, valor):
    return f"{factura},{PREST_PR},CC,{doc},{fecha},AUT,{codigo},3,1,1,R509,,,,{valor}\n"


def zip_hospital_varios_meses():
    """ZIP del hospital con carpetas por mes, una factura repetida en dos carpetas,
    una línea AP incompleta, un ZIP anidado con JSON de otro hospital y un XML."""
    ene = "Hospital Juan de Dios /01. Enero/"
    feb = "Hospital Juan de Dios /02. Febrero/"
    interno = io.BytesIO()
    with zipfile.ZipFile(interno, "w") as z:
        z.writestr("FHSL900/FHSL900.json", json.dumps({
            "numFactura": "FHSL900",
            "usuarios": [{"tipoDocumentoIdentificacion": "CC", "numDocumentoIdentificacion": "77",
                          "servicios": {"otrosServicios": [{"codTecnologiaSalud": "601T01", "cantidadOS": 1, "vrUnitOS": 459300,
                                                            "vrServicio": 459300, "fechaSuministroTecnologia": "2026-02-03 08:00",
                                                            "consecutivo": 1, "codPrestador": "810650020602"}]}}],
        }))
        z.writestr("FHSL900/FHSL900.xml", "<Invoice/>")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr(ene + "AF000001.txt", _af("FHJD100", "01/01/2026", "31/01/2026") + _af("FHJD101", "01/01/2026", "31/01/2026"))
        z.writestr(ene + "AP000001.txt",
                   _ap("FHJD100", "1", "10/01/2026", "601T02", "1700000") + _ap("FHJD101", "2", "20/01/2026", "601T01", "459300")
                   + "FHJD,812200020604,CC,9,07/01/2026,,907009,1,1,Z000,,,,0\n")  # 14 campos
        z.writestr(feb + "AF000002.txt", _af("FHJD101", "01/01/2026", "31/01/2026") + _af("FHJD200", "01/02/2026", "28/02/2026"))
        z.writestr(feb + "AP000002.txt",
                   _ap("FHJD101", "2", "20/01/2026", "601T01", "459300") + _ap("FHJD200", "3", "05/02/2026", "601T02", "3400000"))
        z.writestr("otros/consolidado.zip", interno.getvalue())
    return buf.getvalue()


class RipsVariosMesesTests(BaseDatos):
    def test_lector_meses_carpetas_repetidas_y_zip_anidado(self):
        from flota.importacion.rips import mes_de_carpeta, prefijo_factura

        self.assertEqual(mes_de_carpeta("x.zip/Hospital Juan de Dios /08. Agosto/AP1.txt"), 8)
        self.assertEqual(prefijo_factura("FHJD15037"), "FHJD")
        r = leer_rips_archivos([("hospital.zip", zip_hospital_varios_meses())], "sal", CATALOGO)
        self.assertEqual(sorted((x["numero_factura"], x["fecha_servicio"].month) for x in r["registros"]),
                         [("FHJD100", 1), ("FHJD101", 1), ("FHJD200", 2), ("FHSL900", 2)])
        self.assertEqual(r["lineas_repetidas"], 1)
        self.assertEqual(r["facturas_repetidas"], 1)
        self.assertEqual(r["ignorados"], 1)  # el XML
        incompletas = [h for h in r["hallazgos"] if "14 campos" in h["descripcion"]]
        self.assertEqual(len(incompletas), 1)
        self.assertEqual(r["cobertura"][("FHJD", PREST_PR, "2026-01")], 2)

    def _his(self, factura, fecha, codigo, cantidad, valor, doc):
        from flota.importacion.facturacion_his import hash_documento
        from django.conf import settings

        pr = Sede.objects.get(codigo="pto-rondon")
        return Traslado.objects.create(
            llave_externa=f"HIS|{factura}|{fecha}|{codigo}|{doc}", origen_registro=Traslado.HIS, sede=pr, fecha_servicio=fecha,
            codigo_cups=codigo, cantidad=cantidad, valor_facturado=valor, numero_factura=factura,
            documento_hash=hash_documento("CC", doc, settings.SAL_DOCUMENTO) if doc else "",
        )

    def test_carga_consolidada_detecta_hospital_y_mes_y_audita(self):
        from flota.models import Parametro
        from flota.servicios.auditoria_rips import auditoria_rips

        Parametro.objects.update_or_create(clave="rips_min_facturas_mes", defaults={"valor": 1})

        self._his("FHJD100", dt.date(2026, 1, 10), "601T02", 100, Decimal("1700000"), "1")   # coincide
        self._his("FHJD200", dt.date(2026, 2, 5), "601T02", 200, Decimal("3300000"), "3")    # valor distinto (RIPS 3.400.000)
        self._his("FHJD300", dt.date(2026, 2, 9), "601T01", 1, Decimal("459300"), "4")       # falta en RIPS
        self._his("", dt.date(2026, 2, 11), "601T01", 1, Decimal("459300"), "5")             # prefactura sin número
        lote = importar.importar_rips([("hospital.zip", zip_hospital_varios_meses())], nombre="hospital.zip")
        self.assertIsNone(lote.mes)
        self.assertIsNone(lote.sede)  # dos hospitales: Puerto Rondón (FHJD) y Arauquita (FHSL)
        self.assertEqual(lote.detalle["meses"], ["2026-01", "2026-02"])
        self.assertEqual((lote.actualizados, lote.nuevos), (2, 2))
        self.assertEqual(lote.detalle["con_diferencia_valor"], 1)
        self.assertEqual(Traslado.objects.get(numero_factura="FHSL900").sede.codigo, "arauquita")
        self.assertEqual(Sede.objects.get(codigo="pto-rondon").codigo_prestador, PREST_PR)
        self.assertEqual(Traslado.objects.get(numero_factura="FHJD200").valor_rips, Decimal("3400000"))
        # Recargar no duplica ni vuelve a conciliar
        lote2 = importar.importar_rips([("hospital.zip", zip_hospital_varios_meses())], nombre="hospital.zip")
        self.assertEqual((lote2.nuevos, lote2.actualizados), (0, 0))
        self.assertEqual(Traslado.objects.filter(numero_factura="FHJD101").count(), 1)

        a = auditoria_rips(dt.date(2026, 1, 1), dt.date(2026, 2, 1), "pto-rondon")
        t = a["total"]
        self.assertEqual((t["ambos"], t["diferencia"], t["solo_his"], t["solo_rips"], t["sin_factura"]), (1, 1, 1, 1, 1))
        self.assertAlmostEqual(t["dif_valor"], 100000)
        feb = next(d for d in a["por_mes"] if d["mes"].month == 2)
        self.assertEqual(feb["estado"], "Con diferencias")
        self.assertTrue(feb["cubierto"])
        self.assertEqual(a["n_detalle"], 3)

    def test_mes_sin_rips_cargado_no_es_diferencia(self):
        from flota.servicios.auditoria_rips import auditoria_rips

        self._his("FHJD900", dt.date(2026, 5, 2), "601T01", 1, Decimal("459300"), "8")
        a = auditoria_rips(dt.date(2026, 5, 1), dt.date(2026, 5, 1))
        self.assertEqual((a["total"]["solo_his"], a["total"]["rips_no_cargado"]), (0, 1))
        self.assertEqual(a["por_mes"][0]["estado"], "Sin RIPS cargado")


class AuditoriaRipsVistaTests(BaseDatos):
    def setUp(self):
        self.client.force_login(User.objects.create_superuser("aud", "", "Clave-Prueba-1"))

    def test_cargar_consolidado_sin_mes_ni_hospital_y_ver_auditoria(self):
        r = self.client.post(reverse("flota:cargar_rips"), {
            "mes": "", "sede": "",
            "archivos": [SimpleUploadedFile("hospital.zip", zip_hospital_varios_meses(), content_type="application/zip")],
        })
        lote = LoteImportacion.objects.get(tipo="RIPS")
        self.assertRedirects(r, reverse("flota:rips_detalle", args=[lote.pk]))
        self.assertContains(self.client.get(reverse("flota:rips_detalle", args=[lote.pk])), "2026-01 a 2026-02")
        self.assertContains(self.client.get(reverse("flota:rips_cargas")), "Varios hospitales")
        p = self.client.get(reverse("flota:auditoria_rips"), {"desde": "2026-01", "hasta": "2026-02"})
        self.assertContains(p, "Mes a mes por hospital")
        self.assertContains(p, "Hospital San Juan de Dios de Puerto Rondón")
        x = self.client.get(reverse("flota:auditoria_rips"), {"desde": "2026-01", "hasta": "2026-02", "formato": "xlsx"})
        wb = load_workbook(io.BytesIO(x.content))
        self.assertEqual(wb.sheetnames, ["Por mes y hospital", "Por recorrido", "Diferencias"])
