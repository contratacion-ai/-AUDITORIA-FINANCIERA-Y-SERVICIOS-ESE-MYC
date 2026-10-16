"""Archivos sintéticos para pruebas (sin datos reales de pacientes)."""

import datetime as dt
import io

from openpyxl import Workbook

ENCABEZADO_HIS = [
    "UNIDAD FUNCIONAL", "CONTRATO", "TIPO IDE.", "NUMERO IDE.", "PRIMER APELLIDO", "PRIMER NOMBRE", "SEDE",
    "COD. PROCEDIMIENTO", "ACTIVIDAD", "CANTIDAD", "FECHA DE FACTURACION", "PRE-FACTURA NRO", "FACTURA NRO",
    "VALOR ACTIVIDAD", "VALOR TOTAL", "REGIMEN DEL PACIENTE", "AMBITO", "FECHA DE SERVICIO", "NRO INGRESO", "NOMBRE EPS",
]


def fila_his(sede, cod, cantidad, unitario, ingreso, factura="F1", dia=5, unidad="URGENCIAS", doc="1000001"):
    fecha = dt.datetime(2026, 9, dia)
    return [
        unidad, "EPS EVENTO", "CC", doc, "APELLIDO", "NOMBRE", sede, cod, "TRASLADO", cantidad, fecha,
        f"P{ingreso}", factura, unitario, cantidad * unitario, "SUBSIDIADO", "Urgencias", fecha, ingreso, "EPS PRUEBA",
    ]


def xlsx_his(filas) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "TRASLADOS"
    ws.append(ENCABEZADO_HIS)
    for f in filas:
        ws.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "facturacion_prueba.xlsx"
    return buf


def xlsx_anexo() -> io.BytesIO:
    """Réplica reducida del Anexo 3 con los defectos reales del archivo."""
    wb = Workbook()
    ws = wb.active
    ws.title = "TRASLADOS Anexo 3"
    ws["A7"] = "VIGENCIA"
    for col, anio in zip("DLNP", (2013, 2016, 2018, 2026)):
        ws[f"{col}7"] = anio
    ws["A8"], ws["B8"], ws["C8"] = "ORIGEN", "DESTINO", "Km"
    ws["D8"], ws["L8"] = 2013, 2016
    ws["N8"] = "VALOR DEL KILOMETRO $13.000"
    ws["P8"] = 17000
    ws["A9"] = "TRASLADO BASICO ASISTENCIAL DE PACIENTE"
    ws["P9"] = 459300
    filas = [("TAME", "ARAUCA", 194), (None, "SARAVENA", 101), ("FORTUL", "SARAVENA", 45)]
    for i, (o, d, km) in enumerate(filas, start=10):
        ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"] = o, d, km
        ws[f"D{i}"] = km * 5000
        ws[f"L{i}"] = f"=ROUND(((0.26*#REF!*C{i})/100),0)*100"
        ws[f"P{i}"] = km * 17000
    ws["P12"] = 45 * 17000 + 1000  # tarifa que no cuadra con km × valor
    ws["A14"] = "NOTA DE ACTUALIZACIÓN: ruta alterna de 30 kilómetros hacia Saravena"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
