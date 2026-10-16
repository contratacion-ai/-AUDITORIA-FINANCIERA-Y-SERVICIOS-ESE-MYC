"""Descarga de tablas en Excel (.xlsx)."""

import io

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

FORMATO_PESOS = '"$" #,##0'
FORMATO_NUM = "#,##0.0"


def respuesta_xlsx(nombre: str, hojas: list[tuple[str, list[str], list[list], dict | None]]) -> HttpResponse:
    """hojas: [(título, encabezados, filas, {índice_columna: formato})]."""
    wb = Workbook()
    wb.remove(wb.active)
    for titulo, encabezados, filas, formatos in hojas:
        ws = wb.create_sheet(titulo[:31])
        ws.append(encabezados)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2E7D32")
        for fila in filas:
            ws.append(fila)
        for idx, fmt in (formatos or {}).items():
            for celda in ws.iter_rows(min_row=2, min_col=idx + 1, max_col=idx + 1):
                celda[0].number_format = fmt
        for i, enc in enumerate(encabezados, start=1):
            ancho = max([len(str(enc))] + [len(str(f[i - 1])) for f in filas[:200] if i - 1 < len(f) and f[i - 1] is not None])
            ws.column_dimensions[get_column_letter(i)].width = min(max(ancho + 2, 10), 60)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{nombre}.xlsx"'
    return resp
