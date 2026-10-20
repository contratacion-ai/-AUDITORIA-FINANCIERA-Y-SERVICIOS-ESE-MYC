"""Persistencia de las importaciones (idempotentes: reimportar no duplica)."""

from __future__ import annotations

import csv
import datetime as dt
import io
from decimal import Decimal, InvalidOperation

from django.conf import settings
from django.db import transaction
from django.utils.text import slugify
from openpyxl import load_workbook

from ..importacion import hallazgo, normalizar
from ..importacion.anexo_tarifas import leer_anexo_traslados
from ..importacion.costos_ceco import leer_costos_ambulancias
from ..importacion.facturacion_his import leer_facturacion
from ..importacion.rips import APH, TRASLADO, leer_rips_archivos
from ..models import (
    CodigoTraslado,
    Consumo,
    HallazgoCalidad,
    LoteImportacion,
    Parametro,
    PartidaCostoIndirecto,
    PerfilCostoEstudio,
    RegistroAuditoria,
    Sede,
    TarifaRuta,
    Traslado,
)
from .comunes import ResolutorCatalogos, guardar_hallazgos, parametro, parametros


def _cerrar_lote(lote, nuevos, actualizados, omitidos, hallazgos, resumen=""):
    lote.nuevos, lote.actualizados, lote.omitidos = nuevos, actualizados, omitidos
    lote.hallazgos = guardar_hallazgos(hallazgos, lote)
    lote.resumen = resumen
    lote.save()
    return lote


def codigo_ruta(origen, destino) -> str:
    return slugify(f"{normalizar(origen)}-{normalizar(destino)}")[:160]


# ---------------------------------------------------------------------------
# Tarifas (Anexo 3)
# ---------------------------------------------------------------------------


def guardar_tarifas(datos: dict, usuario=None, archivo="ANEXOS DE TARIFAS 2026.xlsx") -> LoteImportacion:
    lote = LoteImportacion.objects.create(tipo="TARIFAS", archivo=archivo, usuario=usuario)
    nuevos = actualizados = 0
    with transaction.atomic():
        for r in datos["rutas"]:
            _, creado = TarifaRuta.objects.update_or_create(
                vigencia=r["vigencia"],
                codigo=codigo_ruta(r["origen"], r["destino"]),
                defaults={
                    "origen": normalizar(r["origen"]),
                    "destino": normalizar(r["destino"]),
                    "km": r["km"],
                    "valor_km": r["valor_km"],
                    "tarifa": r["tarifa"],
                    "ambito_geografico": r["ambito_geografico"],
                    "fila_fuente": r["fila"],
                    "observacion": r.get("observacion", "")[:300],
                },
            )
            nuevos += int(creado)
            actualizados += int(not creado)
        fuente = f"{archivo}, hoja {datos.get('hoja', 'Anexo 3')}"
        if datos.get("valor_km"):
            _actualizar_parametro("valor_km_tarifa", datos["valor_km"], fuente)
        if datos.get("tarifa_base") and isinstance(datos["tarifa_base"].get("valor"), (int, float)):
            _actualizar_parametro("tarifa_traslado_basico", datos["tarifa_base"]["valor"], fuente)
    return _cerrar_lote(lote, nuevos, actualizados, 0, datos["hallazgos"], f"{len(datos['rutas'])} rutas vigencia {datos['vigencia']}")


def importar_tarifas(archivo, usuario=None, vigencia=2026, nombre="") -> LoteImportacion:
    datos = leer_anexo_traslados(archivo, vigencia, nombre or "ANEXOS DE TARIFAS")
    return guardar_tarifas(datos, usuario, nombre or getattr(archivo, "name", str(archivo)))


def _parametro_editado_a_mano(p: Parametro) -> bool:
    return RegistroAuditoria.objects.filter(modelo="Parametro", objeto_id=str(p.pk), usuario__isnull=False).exists()


def _actualizar_parametro(clave, valor, fuente, forzar=True) -> bool:
    """Actualiza un parámetro desde un archivo fuente.

    ``forzar=False`` (carga automática en una actualización del programa): no
    cambia parámetros que un usuario modificó a mano.
    """
    from .comunes import PARAMETROS_DEFECTO

    _, unidad, descripcion, _, supuesto = PARAMETROS_DEFECTO.get(clave, (None, "", clave, "", False))
    p, creado = Parametro.objects.get_or_create(
        clave=clave,
        defaults={"valor": valor, "descripcion": descripcion, "unidad": unidad, "fuente": fuente, "es_supuesto": supuesto},
    )
    if creado or p.valor == Decimal(str(valor)):
        return creado
    if not forzar and _parametro_editado_a_mano(p):
        return False
    anterior = p.valor
    p.valor = valor
    p.fuente = fuente
    p.version += 1
    p.save()
    RegistroAuditoria.objects.create(
        modelo="Parametro", objeto_id=str(p.pk), accion="actualizar_desde_archivo",
        cambios={"clave": clave, "antes": str(anterior), "despues": str(valor), "fuente": fuente},
    )
    return True


# ---------------------------------------------------------------------------
# Costos del centro de costos y modelo ABC corregido
# ---------------------------------------------------------------------------


def guardar_costos_ceco(datos: dict, usuario=None, archivo="COSTOS AMBULANCIAS.xlsx", forzar_parametros=True) -> LoteImportacion:
    """Guarda partidas, parámetros y correcciones de perfiles del libro de costos.

    Con el modelo ABC corregido (``datos['es_modelo']``) el archivo es la fuente
    de verdad: las partidas que ya no están se eliminan, las demás quedan
    incluidas en el ABC, y los hallazgos de versiones anteriores del mismo libro
    que ya no aparecen se marcan como resueltos.
    """
    lote = LoteImportacion.objects.create(tipo="COSTOS_CECO", archivo=archivo, usuario=usuario)
    nuevos = actualizados = 0
    hallazgos = list(datos["hallazgos"])
    cambios = []
    with transaction.atomic():
        claves = set()
        for p in datos["partidas"]:
            defaults = {"valor_mensual": Decimal(str(p["valor_mensual"])), "fuente": f"{archivo} {p['ubicacion']}"}
            if datos.get("es_modelo"):
                defaults["incluir_en_abc"] = p["incluir_en_abc"]
            obj, creado = PartidaCostoIndirecto.objects.update_or_create(
                categoria=p["categoria"], concepto=p["concepto"], defaults=defaults
            )
            claves.add(obj.pk)
            nuevos += int(creado)
            actualizados += int(not creado)
        if datos.get("es_modelo") and claves:
            borradas = PartidaCostoIndirecto.objects.exclude(pk__in=claves)
            if borradas.exists():
                cambios.append(f"{borradas.count()} partidas que ya no están en el modelo se retiraron")
                borradas.delete()

        fuente = f"{archivo} (modelo ABC corregido)" if datos.get("es_modelo") else archivo
        for clave, valor in datos.get("parametros", {}).items():
            if _actualizar_parametro(clave, valor, fuente, forzar_parametros):
                cambios.append(f"parámetro {clave} = {valor:g}")

        # Correcciones de la tabla por placa del modelo (p. ej. horas/día de OKZ364).
        for fila in datos.get("perfiles", []):
            perfil = PerfilCostoEstudio.objects.filter(ambulancia__placa=fila["placa"], vigente=True).order_by("-fecha_estudio").first()
            if perfil is None:
                continue
            diferencias = {}
            for campo in ("km_dia", "horas_dia", "combustible_dia", "mantenimiento_dia"):
                dec = 4 if campo == "horas_dia" else 2
                nuevo = Decimal(str(round(fila[campo], dec)))
                if abs(getattr(perfil, campo) - nuevo) > Decimal("0.5") / 10**dec:
                    diferencias[campo] = (getattr(perfil, campo), nuevo)
            if diferencias:
                for campo, (_, nuevo) in diferencias.items():
                    setattr(perfil, campo, nuevo)
                perfil.save()
                texto = ", ".join(f"{c} {a} → {n}" for c, (a, n) in diferencias.items())
                cambios.append(f"{fila['placa']}: {texto}")
                hallazgos.append(hallazgo(
                    "COSTOS AMBULANCIAS.xlsx", fila["ubicacion"], "SUPUESTO",
                    f"Corrección del modelo ABC aplicada al perfil de {fila['placa']}: {texto} (el informe técnico reporta el valor anterior).",
                    "BAJA",
                ))

    lote = _cerrar_lote(lote, nuevos, actualizados, 0, hallazgos,
                        f"{len(datos['partidas'])} partidas mensuales" + ("; " + "; ".join(cambios) if cambios else ""))
    if datos.get("es_modelo"):
        vigentes = {(h["ubicacion"][:200], h["tipo"], h["descripcion"]) for h in hallazgos}
        anteriores = HallazgoCalidad.objects.filter(fuente="COSTOS AMBULANCIAS.xlsx", resuelto=False)
        for h in anteriores:
            if (h.ubicacion, h.tipo, h.descripcion) not in vigentes:
                h.resuelto = True
                h.nota_resolucion = "No aparece en la versión corregida del libro (modelo ABC parametrizado)."
                h.save(update_fields=["resuelto", "nota_resolucion"])
    return lote


def importar_costos_ceco(archivo, usuario=None, nombre="") -> LoteImportacion:
    datos = leer_costos_ambulancias(archivo, float(parametro("tripulacion_dia_informe")), "COSTOS AMBULANCIAS.xlsx")
    return guardar_costos_ceco(datos, usuario, nombre or getattr(archivo, "name", str(archivo)))


# ---------------------------------------------------------------------------
# Traslados (facturación HIS y RIPS)
# ---------------------------------------------------------------------------


def es_por_km(cantidad: Decimal, valor_unitario, codigo: CodigoTraslado | None, pars) -> bool:
    """La línea se factura por kilómetro (y su cantidad son km) o con tarifa básica.

    Se decide por el valor unitario frente al manual tarifario; si no coincide
    con ninguno, por la modalidad del código; y en último caso por la cantidad.
    """
    if valor_unitario is not None:
        if pars["valor_km_tarifa"] and valor_unitario == pars["valor_km_tarifa"]:
            return True
        if pars["tarifa_traslado_basico"] and valor_unitario == pars["tarifa_traslado_basico"]:
            return False
    if codigo is not None:
        return codigo.modalidad == CodigoTraslado.POR_KM
    return cantidad > 1


def calcular_km(cantidad: Decimal, por_km: bool, pars) -> Decimal:
    """Km conducidos (driver ABC) de una línea facturada.

    * Por km: km facturados × factor ida y vuelta.
    * Tarifa básica: km_traslado_basico por cada unidad facturada.
    """
    if por_km:
        return (cantidad * pars["factor_ida_vuelta"]).quantize(Decimal("0.01"))
    return (pars["km_traslado_basico"] * max(cantidad, Decimal(1))).quantize(Decimal("0.01"))


def _aplicar_reglas_ingreso(traslados: list[Traslado]):
    """En un ingreso con tarifa básica y cobro por km, la línea básica no suma km
    (el recorrido ya está en la línea por km)."""
    por_ingreso: dict[str, list[Traslado]] = {}
    for t in traslados:
        if t.nro_ingreso:
            por_ingreso.setdefault(f"{t.sede_id}|{t.nro_ingreso}", []).append(t)
    for grupo in por_ingreso.values():
        if any(getattr(t, "_por_km", False) for t in grupo):
            for t in grupo:
                if not getattr(t, "_por_km", False) and t.km_recorridos:
                    t.km_recorridos = Decimal(0)
                    nota = "línea básica sin km: el recorrido se costea en la línea por km del mismo ingreso"
                    if nota not in t.observaciones_calidad:
                        t.observaciones_calidad = "; ".join(filter(None, [t.observaciones_calidad, nota]))


def importar_facturacion(archivo, usuario=None, nombre="", mes: dt.date | None = None, sede: Sede | None = None) -> LoteImportacion:
    """Importa el reporte de facturación de traslados del HIS.

    ``mes`` (opcional): mes declarado al cargar; se informa cuántas líneas
    tienen fecha de servicio fuera de él (se importan igual, en su mes real).
    ``sede`` (opcional): si se indica, solo se importan las líneas de ese hospital.
    """
    nombre = nombre or getattr(archivo, "name", str(archivo))
    codigos = {c.codigo: c for c in CodigoTraslado.objects.all()}
    base = {c for c, o in codigos.items() if o.modalidad == CodigoTraslado.BASE} or {"601T01"}
    datos = leer_facturacion(archivo, settings.SAL_DOCUMENTO, base, nombre)
    lote = LoteImportacion.objects.create(
        tipo="FACTURACION_HIS", archivo=nombre, usuario=usuario, mes=mes.replace(day=1) if mes else None, sede=sede
    )
    res = ResolutorCatalogos()
    pars = parametros()
    nuevos = actualizados = omitidos = otra_sede = fuera_mes = 0
    hallazgos = list(datos["hallazgos"])
    sin_ruta: dict[str, int] = {}
    sedes_desconocidas: dict[str, int] = {}
    tocados: list[Traslado] = []
    por_ambito: dict[str, dict] = {}

    with transaction.atomic():
        existentes = {t.llave_externa: t for t in Traslado.objects.filter(llave_externa__in=[r["llave_externa"] for r in datos["registros"]])}
        for r in datos["registros"]:
            sede_r = res.sede(r["sede"])
            if sede_r is None:
                sedes_desconocidas[r["sede"]] = sedes_desconocidas.get(r["sede"], 0) + 1
                omitidos += 1
                continue
            if sede is not None and sede_r.pk != sede.pk:
                otra_sede += 1
                omitidos += 1
                continue
            if mes and r["fecha_servicio"].replace(day=1) != mes.replace(day=1):
                fuera_mes += 1
            ruta, motivo = (None, "")
            por_km = es_por_km(r["cantidad"], r["valor_unitario"], codigos.get(r["codigo_cups"]), pars)
            if por_km:
                ruta, motivo = res.ruta_por_km(sede_r, r["cantidad"], pars["tolerancia_km_ruta"])
                if ruta is None:
                    sin_ruta[f"{sede_r.origen_tarifa} {r['cantidad']} km"] = sin_ruta.get(f"{sede_r.origen_tarifa} {r['cantidad']} km", 0) + 1
                    r["observaciones"].append(motivo)
            codigo = codigos.get(r["codigo_cups"])
            if codigo is None:
                r["observaciones"].append(f"código {r['codigo_cups']} no está en el catálogo de traslados")
            t = existentes.get(r["llave_externa"])
            creado = t is None
            if creado:
                t = Traslado(llave_externa=r["llave_externa"], origen_registro=Traslado.HIS, creado_por=usuario)
            elif t.origen_registro == Traslado.MANUAL:
                omitidos += 1
                continue
            # Campos que el usuario corrige a mano (ambulancia, ruta) no se pisan.
            if creado or t.ruta_id is None:
                t.ruta = ruta
                t.ruta_inferida = ruta is not None
            for campo in (
                "fecha_servicio", "codigo_cups", "codigo_cups_original", "descripcion", "cantidad", "valor_unitario",
                "valor_facturado", "ambito", "unidad_funcional", "contrato", "eps_nombre", "eps_codigo", "regimen",
                "numero_factura", "prefactura", "nro_ingreso", "autorizacion", "tipo_documento", "documento_hash",
            ):
                setattr(t, campo, r[campo])
            if codigo is not None and codigo.categoria == APH:
                t.categoria = APH
            t.sede = sede_r
            t.lote = lote
            t.km_recorridos = calcular_km(r["cantidad"], por_km, pars)
            t._por_km = por_km
            t.observaciones_calidad = "; ".join(r["observaciones"])[:2000]
            t.actualizado_por = usuario
            if not creado:
                t.version += 1
            tocados.append(t)
            nuevos += int(creado)
            actualizados += int(not creado)
            a = por_ambito.setdefault(f"{t.ambito} / {t.unidad_funcional}", {"lineas": 0, "valor": 0.0})
            a["lineas"] += 1
            a["valor"] += float(t.valor_facturado)
        _aplicar_reglas_ingreso(tocados)
        for t in tocados:
            t.save()

    for s_, n in sedes_desconocidas.items():
        hallazgos.append(hallazgo(nombre, f"SEDE '{s_}'", "FALTANTE", f"{n} líneas con sede no registrada; se omitieron. Agregue el alias en Sedes.", "ALTA"))
    for clave, n in sin_ruta.items():
        hallazgos.append(
            hallazgo(nombre, clave, "INCONSISTENCIA", f"{n} líneas facturan {clave} y no hay una ruta única con esos km en el Anexo 3", "MEDIA")
        )
    r = datos["resumen"]
    lote.detalle = {
        "por_ambito": por_ambito,
        "fuera_del_mes": fuera_mes,
        "de_otra_sede": otra_sede,
        "valor_importado": float(sum((t.valor_facturado for t in tocados), Decimal(0))),
    }
    texto = f"{r['lineas']} líneas en el archivo, ${r['valor_total']:,.0f}, mes principal {r['mes_principal']}"
    if mes and fuera_mes:
        texto += f"; {fuera_mes} líneas con fecha de servicio fuera de {mes:%Y-%m} (quedan en su mes real)"
    if otra_sede:
        texto += f"; {otra_sede} líneas de otros hospitales no se cargaron"
    return _cerrar_lote(lote, nuevos, actualizados, omitidos, hallazgos, texto)


def importar_rips(contenido, usuario=None, nombre="RIPS.json", mes: dt.date | None = None, sede: Sede | None = None) -> LoteImportacion:
    """Carga RIPS (JSON, TXT o ZIP), extrae traslados y APH y los concilia.

    ``contenido``: bytes de un archivo, o lista [(nombre, bytes)] de varios.
    ``mes`` y ``sede``: mes y hospital del envío. Las líneas con fecha fuera
    del mes o con prestador distinto al del hospital se cargan pero se marcan.

    Correspondencia con lo ya cargado: número de factura + código + valor
    (+ hash del documento cuando ambos lo tienen). Una línea que ya vino en
    la facturación del HIS no se duplica: se marca conciliado_rips.
    """
    if not isinstance(contenido, list):
        if not nombre.lower().endswith((".json", ".zip", ".txt")):
            nombre = f"{nombre}.json"
        contenido = [(nombre, contenido if isinstance(contenido, bytes) else contenido.encode("utf-8"))]
    archivos = contenido
    catalogo = dict(CodigoTraslado.objects.filter(activo=True).values_list("codigo", "categoria"))
    datos = leer_rips_archivos(archivos, settings.SAL_DOCUMENTO, catalogo)
    lote = LoteImportacion.objects.create(
        tipo="RIPS", archivo=nombre, usuario=usuario, mes=mes.replace(day=1) if mes else None, sede=sede
    )
    res = ResolutorCatalogos()
    pars = parametros()
    sedes_por_prestador = {s.codigo_prestador: s for s in res.sedes.values() if s.codigo_prestador}
    nuevos = conciliados = omitidos = fuera_mes = prestador_distinto = 0
    totales = {TRASLADO: {"lineas": 0, "valor": 0.0}, APH: {"lineas": 0, "valor": 0.0}}
    hallazgos = list(datos["hallazgos"])
    with transaction.atomic():
        for r in datos["registros"]:
            obs = list(r["observaciones"])
            if mes and r["fecha_servicio"].replace(day=1) != mes.replace(day=1):
                fuera_mes += 1
                obs.append(f"fecha {r['fecha_servicio']} fuera del mes cargado {mes:%Y-%m}")
            if sede is not None and sede.codigo_prestador and r["cod_prestador"] and r["cod_prestador"] != sede.codigo_prestador:
                prestador_distinto += 1
                obs.append(f"prestador {r['cod_prestador']} distinto al de {sede.nombre} ({sede.codigo_prestador})")
            totales[r["categoria"]]["lineas"] += 1
            totales[r["categoria"]]["valor"] += float(r["valor_facturado"])

            ya = Traslado.objects.filter(llave_externa=r["llave_externa"]).first()
            if ya:
                if ya.lote_rips_id is None:
                    ya.lote_rips = lote
                    ya.save(update_fields=["lote_rips"])
                omitidos += 1
                continue
            candidatos = Traslado.objects.filter(
                numero_factura=r["numero_factura"], codigo_cups=r["codigo_cups"], valor_facturado=r["valor_facturado"], conciliado_rips=False
            )
            if r["documento_hash"]:
                con_doc = candidatos.filter(documento_hash=r["documento_hash"])
                candidatos = con_doc if con_doc.exists() else candidatos
            existente = candidatos.order_by("fecha_servicio", "id").first()
            if existente:
                existente.conciliado_rips = True
                existente.rips_consecutivo = r["rips_consecutivo"]
                existente.lote_rips = lote
                campos = ["conciliado_rips", "rips_consecutivo", "lote_rips", "version", "actualizado"]
                if r["categoria"] == APH and existente.categoria != APH:
                    existente.categoria = APH
                    campos.append("categoria")
                existente.version += 1
                existente.save(update_fields=campos)
                conciliados += 1
                continue
            sede_r = sede or sedes_por_prestador.get(r["cod_prestador"])
            if sede_r is None:
                hallazgos.append(
                    hallazgo(nombre, f"factura {r['numero_factura']}", "FALTANTE",
                             f"Línea RIPS sin equivalente en la facturación HIS y prestador '{r['cod_prestador']}' sin sede asociada; "
                             "no se crea. Cargue el RIPS indicando el hospital o registre el código de prestador de la sede.",
                             "ALTA")
                )
                omitidos += 1
                continue
            por_km = es_por_km(r["cantidad"], r["valor_unitario"], None, pars) if r["categoria"] == TRASLADO else False
            ruta, _ = res.ruta_por_km(sede_r, r["cantidad"], pars["tolerancia_km_ruta"]) if por_km else (None, "")
            t = Traslado(
                llave_externa=r["llave_externa"],
                origen_registro=Traslado.RIPS,
                lote=lote,
                lote_rips=lote,
                categoria=r["categoria"],
                sede=sede_r,
                ruta=ruta,
                ruta_inferida=ruta is not None,
                conciliado_rips=True,
                creado_por=usuario,
                ambito=r["ambito"] or "SIN ÁMBITO",
                **{k: r[k] for k in ("fecha_servicio", "codigo_cups", "codigo_cups_original", "descripcion", "cantidad",
                                     "valor_unitario", "valor_facturado", "numero_factura", "autorizacion",
                                     "tipo_documento", "documento_hash", "rips_consecutivo")},
            )
            t.km_recorridos = calcular_km(t.cantidad, por_km, pars)
            obs.insert(0, f"creado desde RIPS ({r['regla']}): no estaba en la facturación HIS")
            t.observaciones_calidad = "; ".join(obs)[:2000]
            t.save()
            nuevos += 1
    if fuera_mes:
        hallazgos.append(hallazgo(nombre, f"RIPS {mes:%Y-%m}", "INCONSISTENCIA", f"{fuera_mes} líneas con fecha fuera del mes cargado", "MEDIA"))
    if prestador_distinto:
        hallazgos.append(hallazgo(nombre, f"RIPS {sede}", "INCONSISTENCIA", f"{prestador_distinto} líneas con código de prestador distinto al del hospital", "MEDIA"))
    lote.detalle = {
        "facturas": datos["facturas"],
        "lineas_revisadas": datos["lineas_revisadas"],
        "traslados": totales[TRASLADO],
        "aph": totales[APH],
        "fuera_del_mes": fuera_mes,
        "prestador_distinto": prestador_distinto,
    }
    return _cerrar_lote(
        lote, nuevos, conciliados, omitidos, hallazgos,
        f"{datos['facturas']} facturas, {datos['lineas_revisadas']} líneas revisadas: {totales[TRASLADO]['lineas']} traslados "
        f"y {totales[APH]['lineas']} APH extraídos; {conciliados} conciliados con la facturación HIS, {nuevos} nuevos",
    )


# ---------------------------------------------------------------------------
# Consumos (CSV o Excel)
# ---------------------------------------------------------------------------

COLUMNAS_CONSUMO = [
    "placa", "fecha", "tipo", "descripcion", "cantidad", "unidad", "valor", "odometro", "documento_soporte", "proveedor",
    "meses_prorrateo",
]


def _filas_consumo(archivo, nombre: str):
    """Filas como dict desde .xlsx o .csv (';' o ',')."""
    if nombre.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(archivo, data_only=True, read_only=True)
        ws = wb.worksheets[0]
        filas = ws.iter_rows(values_only=True)
        enc = [normalizar(c).lower() for c in next(filas, [])]
        for fila in filas:
            if fila and any(v not in (None, "") for v in fila):
                yield {enc[i]: v for i, v in enumerate(fila) if i < len(enc) and enc[i]}
        return
    texto = archivo.read()
    if isinstance(texto, bytes):
        texto = texto.decode("utf-8-sig", errors="replace")
    dialecto = csv.Sniffer().sniff(texto.splitlines()[0], delimiters=";,")
    for fila in csv.DictReader(io.StringIO(texto), dialect=dialecto):
        yield {normalizar(k).lower(): v for k, v in fila.items() if k}


def _a_texto(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def importar_consumos(archivo, usuario=None, nombre="consumos.csv", tipo_forzado: str | None = None) -> LoteImportacion:
    """Carga consumos por ambulancia desde CSV o Excel (columnas COLUMNAS_CONSUMO).

    ``tipo_forzado``: si se carga desde una tabla (p. ej. Combustible), las filas
    sin columna 'tipo' toman ese tipo. Llave: placa + fecha + tipo +
    documento_soporte (si el documento ya existe se actualiza, no se duplica).
    """
    lote = LoteImportacion.objects.create(tipo="CONSUMOS", archivo=nombre, usuario=usuario)
    res = ResolutorCatalogos()
    tipos = dict(Consumo.TIPOS)
    nuevos = actualizados = omitidos = 0
    hallazgos = []
    with transaction.atomic():
        for n, fila in enumerate(_filas_consumo(archivo, nombre), start=2):
            fila = {k: _a_texto(v) if not isinstance(v, (dt.date, dt.datetime)) else v for k, v in fila.items()}
            amb = res.ambulancia(fila.get("placa"))
            tipo = normalizar(fila.get("tipo")) or (tipo_forzado or "")
            tipo = Consumo.SINONIMOS.get(tipo, tipo)
            try:
                f = fila.get("fecha")
                fecha = f.date() if isinstance(f, dt.datetime) else f if isinstance(f, dt.date) else dt.date.fromisoformat(str(f)[:10])
                valor = Decimal(str(fila.get("valor", "")).replace(",", "."))
                cantidad = Decimal((str(fila.get("cantidad") or "0")).replace(",", "."))
                meses = int(fila.get("meses_prorrateo") or 1)
            except (ValueError, InvalidOperation):
                hallazgos.append(hallazgo(nombre, f"fila {n}", "FALTANTE", "Fecha (AAAA-MM-DD), valor, cantidad o meses no válidos; se omite", "MEDIA"))
                omitidos += 1
                continue
            if amb is None or tipo not in tipos or not 1 <= meses <= 36 or valor < 0:
                hallazgos.append(hallazgo(nombre, f"fila {n}", "FALTANTE",
                                          f"Placa '{fila.get('placa')}', tipo '{fila.get('tipo') or tipo_forzado}', valor o meses no válidos; se omite", "MEDIA"))
                omitidos += 1
                continue
            doc = fila.get("documento_soporte", "")
            existente = (
                Consumo.objects.filter(ambulancia=amb, fecha=fecha, tipo=tipo, documento_soporte=doc).first() if doc else None
            )
            c = existente or Consumo(ambulancia=amb, fecha=fecha, tipo=tipo, documento_soporte=doc, creado_por=usuario)
            c.descripcion = fila.get("descripcion", "")[:200]
            c.cantidad = cantidad
            c.unidad = fila.get("unidad", "")[:20] or ("gal" if tipo == Consumo.COMBUSTIBLE else "")
            c.valor = valor
            c.meses_prorrateo = meses
            c.odometro = int(fila["odometro"]) if str(fila.get("odometro", "")).isdigit() else None
            c.proveedor = fila.get("proveedor", "")[:120]
            c.lote = lote
            c.actualizado_por = usuario
            if existente:
                c.version += 1
            c.save()
            nuevos += int(existente is None)
            actualizados += int(existente is not None)
    return _cerrar_lote(lote, nuevos, actualizados, omitidos, hallazgos, f"tabla {tipo_forzado or 'según columna tipo'}")


def importar_consumos_csv(archivo, usuario=None, nombre="consumos.csv") -> LoteImportacion:
    """Compatibilidad con la versión anterior."""
    return importar_consumos(archivo, usuario, nombre)
