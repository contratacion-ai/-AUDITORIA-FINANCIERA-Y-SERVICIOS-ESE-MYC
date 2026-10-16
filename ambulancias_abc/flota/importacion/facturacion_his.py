"""Lector del reporte de facturación de traslados del sistema de información
hospitalario (formato 'TRASLADOS FACTURACION POR AMBITO <MES>.xlsx').

Protección de datos (Ley 1581 de 2012): no se leen nombres, teléfonos ni
direcciones. El documento del paciente se guarda solo como hash SHA-256 con
sal, suficiente para cruzar con RIPS sin almacenar el número.
"""

from __future__ import annotations

import datetime as dt
import hashlib
from collections import Counter, defaultdict
from decimal import Decimal

from openpyxl import load_workbook

from . import hallazgo, normalizar, normalizar_cups

FUENTE = "Facturación traslados (HIS)"

# nombre normalizado de columna -> campo
COLUMNAS = {
    "UNIDAD FUNCIONAL": "unidad_funcional",
    "CONTRATO": "contrato",
    "TIPO IDE.": "tipo_documento",
    "NUMERO IDE.": "numero_documento",
    "SEDE": "sede",
    "COD. PROCEDIMIENTO": "codigo_cups_original",
    "ACTIVIDAD": "descripcion",
    "CANTIDAD": "cantidad",
    "FECHA DE FACTURACION": "fecha_facturacion",
    "NRO. AUTORIZACION": "autorizacion",
    "MUN NOM": "municipio_paciente",
    "PRE-FACTURA NRO": "prefactura",
    "FACTURA NRO": "numero_factura",
    "VALOR ACTIVIDAD": "valor_unitario",
    "VALOR TOTAL": "valor_facturado",
    "REGIMEN DEL PACIENTE": "regimen",
    "AMBITO": "ambito",
    "FECHA DE SERVICIO": "fecha_servicio",
    "NRO INGRESO": "nro_ingreso",
    "NOMBRE EPS": "eps_nombre",
    "CODIGO EPS": "eps_codigo",
}
OBLIGATORIAS = {"sede", "codigo_cups_original", "cantidad", "valor_facturado", "fecha_servicio"}


def hash_documento(tipo, numero, sal: str) -> str:
    if numero in (None, ""):
        return ""
    base = f"{sal}|{normalizar(tipo)}|{str(numero).strip()}"
    return hashlib.sha256(base.encode("utf-8")).hexdigest()


def _fecha(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str) and v.strip():
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return dt.datetime.strptime(v.strip()[:10], fmt).date()
            except ValueError:
                pass
    return None


def _num(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    try:
        return Decimal(str(v).replace(",", "").strip())
    except Exception:
        return None


def _texto(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def leer_facturacion(archivo, sal: str, codigos_base: set[str] | None = None, nombre_fuente: str = FUENTE) -> dict:
    """Devuelve {'registros': [...], 'hallazgos': [...], 'resumen': {...}}.

    ``codigos_base``: códigos CUPS de tarifa fija (traslado básico), para
    detectar ingresos con cobro base + cobro por km.
    """
    codigos_base = codigos_base or {"601T01"}
    wb = load_workbook(archivo, data_only=True, read_only=True)
    ws = None
    for hoja in wb.worksheets:
        primera = [normalizar(c) for c in next(hoja.iter_rows(min_row=1, max_row=1, values_only=True), [])]
        if "COD. PROCEDIMIENTO" in primera and "VALOR TOTAL" in primera:
            ws = hoja
            encabezado = primera
            break
    if ws is None:
        raise ValueError("No se encontró una hoja con las columnas 'COD. PROCEDIMIENTO' y 'VALOR TOTAL'")

    idx = {}
    for i, nombre in enumerate(encabezado):
        campo = COLUMNAS.get(nombre)
        if campo and campo not in idx:
            idx[campo] = i
    faltan = OBLIGATORIAS - set(idx)
    if faltan:
        raise ValueError(f"Faltan columnas obligatorias: {', '.join(sorted(faltan))}")

    hallazgos: list[dict] = []
    registros: list[dict] = []
    ocurrencias = Counter()
    hoja = ws.title

    for n_fila, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not fila or not any(v not in (None, "") for v in fila):
            continue
        g = lambda campo: fila[idx[campo]] if campo in idx and idx[campo] < len(fila) else None  # noqa: E731
        ubic = f"{hoja}!fila {n_fila}"
        cod_orig = _texto(g("codigo_cups_original"))
        cod = normalizar_cups(cod_orig)
        cantidad = _num(g("cantidad"))
        unitario = _num(g("valor_unitario"))
        total = _num(g("valor_facturado"))
        fecha_serv = _fecha(g("fecha_servicio"))
        obs = []

        if cod != cod_orig:
            obs.append(f"código '{cod_orig}' normalizado a {cod}")
        if fecha_serv is None:
            hallazgos.append(hallazgo(nombre_fuente, ubic, "FALTANTE", "Fila sin fecha de servicio válida; se omite", "ALTA"))
            continue
        if total is None:
            hallazgos.append(hallazgo(nombre_fuente, ubic, "FALTANTE", "Fila sin valor total; se omite", "ALTA"))
            continue
        if cantidad is not None and unitario is not None and cantidad * unitario != total:
            obs.append("cantidad × valor unitario ≠ valor total")
        if total == 0:
            obs.append("valor facturado en cero")
        if not _texto(g("numero_factura")):
            obs.append("sin número de factura (solo prefactura)")

        reg = {
            "fila": n_fila,
            "sede": _texto(g("sede")),
            "unidad_funcional": _texto(g("unidad_funcional")),
            "ambito": _texto(g("ambito")) or "SIN ÁMBITO",
            "contrato": _texto(g("contrato")),
            "eps_nombre": _texto(g("eps_nombre")),
            "eps_codigo": _texto(g("eps_codigo")),
            "regimen": _texto(g("regimen")),
            "tipo_documento": _texto(g("tipo_documento")),
            "documento_hash": hash_documento(g("tipo_documento"), g("numero_documento"), sal),
            "codigo_cups": cod,
            "codigo_cups_original": cod_orig,
            "descripcion": _texto(g("descripcion"))[:250],
            "cantidad": cantidad or Decimal(0),
            "valor_unitario": unitario,
            "valor_facturado": total,
            "fecha_servicio": fecha_serv,
            "fecha_facturacion": _fecha(g("fecha_facturacion")),
            "numero_factura": _texto(g("numero_factura")),
            "prefactura": _texto(g("prefactura")),
            "nro_ingreso": _texto(g("nro_ingreso")),
            "autorizacion": _texto(g("autorizacion"))[:60],
            "observaciones": obs,
        }
        base_llave = "|".join(
            [reg["prefactura"], reg["nro_ingreso"], cod, fecha_serv.isoformat(), str(reg["cantidad"]), str(total)]
        )
        ocurrencias[base_llave] += 1
        reg["llave_externa"] = f"HIS|{base_llave}|{ocurrencias[base_llave]}"
        registros.append(reg)

    # Validaciones entre filas.
    meses = Counter((r["fecha_servicio"].year, r["fecha_servicio"].month) for r in registros)
    mes_principal = meses.most_common(1)[0][0] if meses else None
    por_ingreso = defaultdict(list)
    for r in registros:
        if r["nro_ingreso"]:
            por_ingreso[r["nro_ingreso"]].append(r)
        if mes_principal and (r["fecha_servicio"].year, r["fecha_servicio"].month) != mes_principal:
            r["observaciones"].append(f"fecha de servicio {r['fecha_servicio']} fuera del mes principal del archivo")

    for ingreso, filas in por_ingreso.items():
        firmas = Counter((f["codigo_cups"], f["cantidad"], f["valor_facturado"], f["fecha_servicio"]) for f in filas)
        for firma, n in firmas.items():
            if n > 1:
                for f in filas:
                    if (f["codigo_cups"], f["cantidad"], f["valor_facturado"], f["fecha_servicio"]) == firma:
                        f["observaciones"].append(f"posible duplicado: {n} líneas iguales en el ingreso {ingreso}")
        tiene_base = any(f["codigo_cups"] in codigos_base and f["cantidad"] == 1 for f in filas)
        tiene_km = any(f["cantidad"] > 1 for f in filas)
        if tiene_base and tiene_km:
            for f in filas:
                f["observaciones"].append(
                    f"ingreso {ingreso} con tarifa básica y cobro por km: verificar si procede el cobro conjunto"
                )

    conteo_obs = Counter()
    for r in registros:
        for o in r["observaciones"]:
            clave = o.split(":")[0] if o.startswith(("posible duplicado", "ingreso")) else o.split(" (")[0]
            if o.startswith("código"):
                clave = "código CUPS con caracteres extra"
            if o.startswith("fecha de servicio"):
                clave = "fecha de servicio fuera del mes principal"
            if o.startswith("ingreso"):
                clave = "tarifa básica y cobro por km en el mismo ingreso"
            if o.startswith("posible duplicado"):
                clave = "posible línea duplicada"
            conteo_obs[clave] += 1
    for clave, n in conteo_obs.items():
        sev = "ALTA" if clave in ("posible línea duplicada", "tarifa básica y cobro por km en el mismo ingreso") else "MEDIA"
        hallazgos.append(hallazgo(nombre_fuente, hoja, "INCONSISTENCIA", f"{n} líneas: {clave}", sev))

    resumen = {
        "lineas": len(registros),
        "valor_total": sum((r["valor_facturado"] for r in registros), Decimal(0)),
        "ingresos": len(por_ingreso),
        "mes_principal": f"{mes_principal[0]}-{mes_principal[1]:02d}" if mes_principal else None,
    }
    return {"registros": registros, "hallazgos": hallazgos, "resumen": resumen}
