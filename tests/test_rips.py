"""Pruebas de rips_auditoria con RIPS sintéticos.

Cubren, como mínimo: un archivo válido, un archivo con errores de estructura
y un ZIP con varios hospitales; además homologación, plantilla, exportación
y CLI.
"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from rips_auditoria import cargar_configuracion, procesar
from rips_auditoria.cli import main
from rips_auditoria.especificacion import ESPECIFICACIONES
from rips_auditoria.excel import generar_plantilla
from rips_auditoria.exportacion import exportar_todo, exportar_unidad
from rips_auditoria.homologacion import ErrorMapeo, cargar_mapeo, homologar, validar_mapeo
from rips_auditoria.lectura import detectar_mes

from . import fabrica


def reglas(res, severidad=None) -> set[str]:
    inc = res.incidencias
    if severidad:
        inc = inc[inc["severidad"] == severidad]
    return set(inc["regla"])


# ---------------------------------------------------------------------------
# Especificación
# ---------------------------------------------------------------------------

def test_cantidad_de_campos_segun_anexo_tecnico():
    esperado = {"CT": 4, "AF": 17, "US": 14, "AD": 6, "AC": 17, "AP": 15,
                "AU": 17, "AH": 19, "AN": 14, "AM": 14, "AT": 11}
    assert {t: e.num_campos for t, e in ESPECIFICACIONES.items()} == esperado


@pytest.mark.parametrize("carpeta,esperado", [
    ("05. Mayo", (5, None)), ("12. Diciembre ", (12, None)), ("SEPTIEMBRE 2025", (9, 2025)),
    ("2026-03", (3, 2026)), ("10260__FIDUPREVISORA", (None, None)), ("Hospital San Antonio", (None, None)),
])
def test_detectar_mes(carpeta, esperado):
    assert detectar_mes(carpeta) == esperado


# ---------------------------------------------------------------------------
# Archivo válido
# ---------------------------------------------------------------------------

def test_remision_valida_sin_errores(tmp_path):
    carpeta = fabrica.escribir(tmp_path / "Hospital San Antonio" / "05. Mayo" / "123__SANITAS",
                               fabrica.envio_valido())
    res = procesar([tmp_path / "Hospital San Antonio"])

    assert res.unidades() == [("Hospital San Antonio", "2026-05")]
    assert reglas(res, "ERROR") == set()
    assert reglas(res, "ADVERTENCIA") == set()
    u = res.unidad("Hospital San Antonio", "2026-05")
    conteos = u.resumen()["registros_por_archivo"]
    assert conteos == {"CT": 9, "AF": 1, "US": 2, "AD": 0, "AC": 1, "AP": 1, "AU": 1, "AH": 1, "AN": 1,
                       "AM": 1, "AT": 1}
    assert list(u.conciliacion["estado"]) == ["CONCILIADA"]
    assert u.conciliacion["total_servicios"].iloc[0] == 103000
    ac = u.datos["AC"].iloc[0]
    assert ac["num_factura"] == "FV1001" and ac["finalidad"] == "10"
    assert ac["archivo"].endswith("AC000123.txt") and ac["linea"] == 1
    assert carpeta.exists()


# ---------------------------------------------------------------------------
# Archivo con errores de estructura
# ---------------------------------------------------------------------------

def test_remision_con_errores_reporta_y_no_se_detiene(tmp_path):
    fabrica.escribir(tmp_path / "Hospital X" / "05. Mayo" / "124__EPS", fabrica.envio_con_errores())
    res = procesar([tmp_path / "Hospital X"])
    r = reglas(res)

    for esperada in [
        "CANTIDAD_CAMPOS", "FORMATO_FECHA", "VALOR_NO_PERMITIDO", "VALOR_CATALOGO_EXTENDIDO",
        "FACTURA_NO_EXISTE_EN_AF", "USUARIO_NO_EXISTE_EN_US", "CAMPO_OBLIGATORIO", "FORMATO_NUMERICO",
        "COMILLAS", "SEPARADOR_FINAL", "LINEA_VACIA", "INCONSISTENCIA_VALOR", "CT_TOTAL_REGISTROS",
        "CT_ARCHIVO_NO_ENCONTRADO", "LONGITUD_EXCEDIDA", "CODIFICACION_NO_UTF8", "TEXTO_REPARADO",
        "DIFERENCIA_VALOR_FACTURA",
    ]:
        assert esperada in r, f"falta la regla {esperada}"

    u = res.unidad("Hospital X", "2026-05")
    # 12 líneas de AC con datos; sólo la de 16 campos se excluye
    assert len(u.datos["AC"]) == 11
    rechazo = res.incidencias[res.incidencias["regla"] == "CANTIDAD_CAMPOS"].iloc[0]
    assert rechazo["linea"] == 2 and "AC000124" in rechazo["archivo"]
    # La línea vacía es la 12; la siguiente conserva su número real (13)
    assert 13 in set(u.datos["AC"]["linea"])
    # Reparación de doble codificación y lectura de cp1252
    assert "CASTAÑEDA" in set(u.datos["US"]["primer_apellido"])
    assert u.datos["AT"]["nombre_servicio"].iloc[0] == "INTERNACIÓN ADULTOS"
    # Comillas retiradas
    assert (u.datos["AC"]["num_factura"] == "FV2001").sum() == 10
    # Severidades: catálogo 2275 es advertencia, fuera de catálogo es error
    inc = res.incidencias
    assert set(inc.loc[inc["regla"] == "VALOR_CATALOGO_EXTENDIDO", "severidad"]) == {"ADVERTENCIA"}
    fuera = inc[inc["regla"] == "VALOR_NO_PERMITIDO"]
    assert {("AC", "finalidad", "99"), ("US", "tipo_usuario", "9"), ("US", "sexo", "X")} <= set(
        map(tuple, fuera[["tipo_archivo", "campo", "valor"]].to_numpy()))


def test_severidad_configurable(tmp_path):
    fabrica.escribir(tmp_path / "H" / "05. Mayo" / "r", fabrica.envio_con_errores())
    cfg = cargar_configuracion()
    cfg["validacion"]["severidades"] = {"LONGITUD_EXCEDIDA": "OMITIR", "VALOR_CATALOGO_EXTENDIDO": "INFO"}
    res = procesar([tmp_path / "H"], cfg)
    assert "LONGITUD_EXCEDIDA" not in reglas(res)
    assert "VALOR_CATALOGO_EXTENDIDO" in reglas(res, "INFO")


def test_remision_sin_archivos_obligatorios(tmp_path):
    archivos = fabrica.envio_valido()
    del archivos["US000123.txt"]
    fabrica.escribir(tmp_path / "H" / "05. Mayo" / "r", archivos)
    res = procesar([tmp_path / "H"])
    assert "ARCHIVO_OBLIGATORIO_FALTANTE" in reglas(res, "ERROR")
    assert "CT_ARCHIVO_NO_ENCONTRADO" in reglas(res)


def test_prefiere_archivo_corregido(tmp_path):
    archivos = fabrica.envio_valido()
    archivos["AC000123_AC_CORREGIDO.txt"] = archivos["AC000123.txt"].replace(",10,13,", ",03,13,")
    archivos["AC000123_AC_CORREGIDO_LOG_CORRECCIONES.txt"] = "log"
    fabrica.escribir(tmp_path / "H" / "05. Mayo" / "r", archivos)
    res = procesar([tmp_path / "H"])
    assert res.datos["AC"]["finalidad"].tolist() == ["03"]
    assert "ARCHIVO_CORREGIDO_USADO" in reglas(res)
    assert "ARCHIVO_NO_RECONOCIDO" not in reglas(res)  # el LOG se ignora en silencio


# ---------------------------------------------------------------------------
# ZIP con varios hospitales
# ---------------------------------------------------------------------------

@pytest.fixture
def zip_lote(tmp_path) -> Path:
    return fabrica.zip_varios_hospitales(tmp_path / "Lote_Hospitales-20261014T130903Z-1-001.zip")


def test_zip_varios_hospitales_organiza_por_hospital_y_mes(zip_lote):
    res = procesar([zip_lote])
    assert res.unidades() == [
        ("Hospital San Antonio", "2026-05"), ("Hospital San Antonio", "2026-06"),
        ("Hospital San Jose", "2026-05"), ("Hospital San Jose", "2026-06"),
    ]
    assert res.hospitales() == ["Hospital San Antonio", "Hospital San Jose"]
    # ZIP anidado procesado, ZIP corrupto reportado sin detener el lote
    sj_jun = res.unidad("Hospital San Jose", "2026-06")
    assert sj_jun.datos["AF"]["num_factura"].tolist() == ["FB5101"]
    corrupto = res.incidencias[res.incidencias["regla"] == "ZIP_CORRUPTO"]
    assert len(corrupto) == 1 and corrupto.iloc[0]["hospital"] == "Hospital San Jose"
    # Cada unidad contiene sólo lo suyo
    sa_may = res.unidad("Hospital San Antonio", "2026-05")
    assert set(sa_may.datos["AP"]["num_factura"]) == {"FV1001"}
    # El JSON de la Res. 2275 se inventaría como ignorado
    ign = res.archivos[res.archivos["estado"] == "ignorado"]
    assert ign["archivo"].str.endswith("FV1001.json").any()
    assert reglas(res, "ERROR") == {"ZIP_CORRUPTO"}


def test_exportacion_excel_json_por_hospital_mes(zip_lote, tmp_path):
    res = procesar([zip_lote])
    salida = tmp_path / "salida"
    archivos = exportar_todo(res, salida)
    nombres = {p.relative_to(salida).as_posix() for p in archivos}
    for h, p in res.unidades():
        for pref, ext in [("RIPS", "xlsx"), ("RIPS", "json"), ("AUDITORIA", "xlsx"), ("AUDITORIA", "json")]:
            assert f"{h}/{p}/{pref}_{h}_{p}.{ext}" in nombres
    assert {"RIPS_consolidado.xlsx", "RIPS_consolidado.json"} <= nombres

    # Excel de una unidad: hojas del Anexo con encabezados del Anexo
    wb = openpyxl.load_workbook(salida / "Hospital San Jose/2026-05/RIPS_Hospital San Jose_2026-05.xlsx",
                                read_only=True)
    for t, spec in ESPECIFICACIONES.items():
        enc = next(wb[t].iter_rows(max_row=1, values_only=True))
        assert list(enc[: spec.num_campos]) == [c.etiqueta for c in spec.campos]
    fila_af = list(wb["AF"].iter_rows(min_row=2, max_row=2, values_only=True))[0]
    assert fila_af[4] == "FB5001" and fila_af[16] == 103000

    # JSON consolidado: hospital -> meses -> registros por archivo
    doc = json.loads((salida / "RIPS_consolidado.json").read_text(encoding="utf-8"))
    assert set(doc["hospitales"]) == {"Hospital San Antonio", "Hospital San Jose"}
    mayo = doc["hospitales"]["Hospital San Jose"]["meses"]["2026-05"]
    assert mayo["registros"]["AF"][0]["valor_neto"] == 103000
    assert mayo["registros"]["AF"][0]["fecha_expedicion"] == "30/05/2026"
    assert set(mayo["registros"]) == set(ESPECIFICACIONES)

    unidad = json.loads((salida / "Hospital San Antonio/2026-06/RIPS_Hospital San Antonio_2026-06.json")
                        .read_text(encoding="utf-8"))
    assert unidad["hospital"] == "Hospital San Antonio" and unidad["periodo"] == "2026-06"
    assert unidad["resumen"]["facturas"] == 1


def test_exportar_una_sola_unidad(zip_lote, tmp_path):
    res = procesar([zip_lote])
    archivos = exportar_unidad(res, "Hospital San Jose", "2026-06", tmp_path)
    assert len(archivos) == 4 and all(p.exists() for p in archivos)
    with pytest.raises(KeyError):
        res.unidad("Hospital San Jose", "2020-01")


# ---------------------------------------------------------------------------
# TXT sueltos
# ---------------------------------------------------------------------------

def test_txt_sueltos_con_hospital_forzado_y_periodo_por_factura(tmp_path):
    carpeta = fabrica.escribir(tmp_path / "sueltos", fabrica.envio_valido(mes="03"))
    archivos = sorted(carpeta.glob("*.txt"))
    res = procesar(archivos, hospital="Hospital Tame")
    assert res.unidades() == [("Hospital Tame", "2026-03")]
    assert res.envios["fuente_periodo"].tolist() == ["factura"]


def test_cargar_carpeta_del_mes_conserva_hospital(tmp_path):
    fabrica.escribir(tmp_path / "Hospital Tame" / "05. Mayo" / "123__SANITAS" / "Rips", fabrica.envio_valido())
    res = procesar([tmp_path / "Hospital Tame" / "05. Mayo"])
    assert res.unidades() == [("Hospital Tame", "2026-05")]
    assert res.envios["carpeta_envio"].tolist() == ["123__SANITAS"]


def test_hospital_por_codigo_de_prestador(tmp_path):
    carpeta = fabrica.escribir(tmp_path / "sueltos", fabrica.envio_valido())
    cfg = cargar_configuracion()
    cfg["organizacion"]["hospitales_por_prestador"] = {fabrica.PRESTADOR: "Hospital Configurado"}
    res = procesar(sorted(carpeta.glob("*.txt")), cfg)
    assert res.hospitales() == ["Hospital Configurado"]


# ---------------------------------------------------------------------------
# Homologación
# ---------------------------------------------------------------------------

def test_homologacion_por_defecto(tmp_path):
    fabrica.escribir(tmp_path / "H" / "05. Mayo" / "r", fabrica.envio_valido())
    res = procesar([tmp_path / "H"])
    u = res.unidad("H", "2026-05")
    hom = homologar(u, cargar_mapeo(cargar_configuracion()["homologacion"]["archivo_mapeo"]))

    us = hom.hojas["Usuarios"].set_index("numDocumentoIdentificacion")
    assert us.loc["1001", "tipoUsuario"] == "04"  # subsidiado
    assert us.loc["1001", "codMunicipioResidencia (5 dig)"] == "81794"
    assert us.loc["1001", "codZonaTerritorialResidencia"] == "02"  # urbana
    assert us.loc["2002", "codZonaTerritorialResidencia"] == "01"  # rural
    ac = hom.hojas["Consultas"].iloc[0]
    assert ac["fechaInicioAtencion (yyyy-MM-dd HH:mm)"] == "2026-05-05 00:00"
    assert ac["causaMotivoAtencion"] == "38"  # 13 enfermedad general
    assert ac["finalidadTecnologiaSalud"] == "44"
    assert ac["conceptoRecaudo"] == "02" and ac["valorPagoModerador"] == 5000
    assert ac["numFactura"] == "FV1001" and ac["_linea"] == 1
    au = hom.hojas["Urgencias"].iloc[0]
    assert au["fechaEgreso (yyyy-MM-dd HH:mm)"] == "2026-05-07 14:30"
    assert au["condicionDestinoUsuarioEgreso"] == "01"
    assert hom.hojas["Medicamentos"].iloc[0]["fechaDispensAdmon (yyyy-MM-dd HH:mm)"] == "2026-05-01 00:00"
    assert hom.hojas["Transacciones"].iloc[0]["valorNetoPagar"] == 103000
    assert hom.incidencias.empty


def test_homologacion_conserva_codigos_destino_y_reporta_sin_equivalencia(tmp_path):
    fabrica.escribir(tmp_path / "H" / "05. Mayo" / "r", fabrica.envio_con_errores())
    res = procesar([tmp_path / "H"])
    hom = homologar(res.unidad("H", "2026-05"), cargar_mapeo(cargar_configuracion()["homologacion"]["archivo_mapeo"]))
    ac = hom.hojas["Consultas"]
    assert "38" in set(ac["causaMotivoAtencion"])  # ya venía en Res. 2275
    sin = hom.incidencias
    assert set(sin["regla"]) == {"HOMOLOGACION_SIN_EQUIVALENCIA"}
    assert {"99", "9"} <= set(sin["valor"])


def test_mapeo_personalizado_y_validacion_de_mapeo(tmp_path):
    fabrica.escribir(tmp_path / "H" / "05. Mayo" / "r", fabrica.envio_valido())
    res = procesar([tmp_path / "H"])
    mapeo = {
        "perfil": "Herramienta X",
        "incluir_trazabilidad": False,
        "tablas": {"sexo": {"valores": {"F": "MUJER", "M": "HOMBRE"}}},
        "hojas": [{"nombre": "Pacientes", "origen": "US", "columnas": [
            {"destino": "documento", "plantilla": "{tipo_id_usuario}-{num_id_usuario}"},
            {"destino": "genero", "origen": "sexo", "tabla": "sexo"},
            {"destino": "fuente", "constante": "RIPS3374"},
        ]}],
    }
    validar_mapeo(mapeo)
    hom = homologar(res.unidad("H", "2026-05"), mapeo)
    assert hom.hojas["Pacientes"].to_dict("records")[0] == {"documento": "CC-1001", "genero": "MUJER",
                                                            "fuente": "RIPS3374"}
    malo = {"hojas": [{"nombre": "X", "origen": "US", "columnas": [{"destino": "a", "origen": "no_existe"}]}]}
    with pytest.raises(ErrorMapeo, match="no_existe"):
        validar_mapeo(malo)
    with pytest.raises(ErrorMapeo, match="tabla"):
        validar_mapeo({"hojas": [{"nombre": "X", "origen": "US",
                                  "columnas": [{"destino": "a", "origen": "sexo", "tabla": "nada"}]}]})


# ---------------------------------------------------------------------------
# Plantilla y CLI
# ---------------------------------------------------------------------------

def test_plantilla_parametrizada(tmp_path):
    ruta = generar_plantilla(tmp_path / "plantilla.xlsx")
    wb = openpyxl.load_workbook(ruta)
    assert set(ESPECIFICACIONES) <= set(wb.sheetnames)
    assert {"Instrucciones", "Diccionario", "Catalogos", "Reglas_Validacion"} <= set(wb.sheetnames)
    ws = wb["US"]
    assert [c.value for c in ws[1]] == [c.etiqueta for c in ESPECIFICACIONES["US"].campos]
    assert ws.cell(1, 1).comment is not None and "Longitud" in ws.cell(1, 1).comment.text
    assert len(ws.data_validations.dataValidation) >= 1
    dic = pd.read_excel(ruta, sheet_name="Diccionario")
    assert len(dic) == sum(e.num_campos for e in ESPECIFICACIONES.values())


def test_cli_procesar_listar_y_sin_archivos(zip_lote, tmp_path, capsys):
    assert main(["listar", str(zip_lote)]) == 0
    assert "Hospital San Jose" in capsys.readouterr().out
    salida = tmp_path / "out"
    assert main(["procesar", str(zip_lote), "-s", str(salida), "--sin-consolidado",
                 "--filtro-hospital", "Hospital San Jose", "--filtro-mes", "2026-05"]) == 0
    generados = sorted(p.name for p in salida.rglob("*.*"))
    assert generados == sorted([
        "AUDITORIA_Hospital San Jose_2026-05.json", "AUDITORIA_Hospital San Jose_2026-05.xlsx",
        "RIPS_Hospital San Jose_2026-05.json", "RIPS_Hospital San Jose_2026-05.xlsx"])
    vacio = tmp_path / "vacio"
    vacio.mkdir()
    assert main(["listar", str(vacio)]) == 2
    assert main(["procesar", str(zip_lote), "-s", str(salida), "--filtro-mes", "1999-01"]) == 1
    assert main(["procesar", str(zip_lote), "-s", str(salida), "-m", str(tmp_path / "no.yaml")]) == 1


# ---------------------------------------------------------------------------
# Lanzador del cargue (iniciar_cargue.bat / .sh)
# ---------------------------------------------------------------------------

def test_cargue_guiado_con_rutas(zip_lote, tmp_path, monkeypatch):
    from rips_auditoria import cargue

    abiertas = []
    monkeypatch.setattr(cargue, "abrir_carpeta", abiertas.append)
    salida = tmp_path / "out"
    assert cargue.main([str(zip_lote), "-s", str(salida), "--sin-consolidado"]) == 0
    assert (salida / "Hospital San Jose" / "2026-05" / "RIPS_Hospital San Jose_2026-05.xlsx").exists()
    assert abiertas == [salida]


def test_cargue_guiado_pregunta_y_valida_rutas(tmp_path, monkeypatch):
    from rips_auditoria import cargue

    monkeypatch.setattr(cargue, "elegir_con_ventana", lambda: None)  # sin interfaz gráfica
    monkeypatch.setattr("builtins.input", lambda _: str(tmp_path / "no_existe.zip"))
    assert cargue.main(["--no-abrir"]) == 1
    monkeypatch.setattr(cargue, "elegir_con_ventana", lambda: [])  # el usuario canceló
    assert cargue.main(["--no-abrir"]) == 1


def test_lote_solo_con_incidencias_de_campo(tmp_path):
    # Regresión: con sólo incidencias vectorizadas (sin hospital/período), la
    # ubicación fallaba con "Invalid value ... for dtype 'float64'".
    archivos = fabrica.envio_valido()
    archivos["AF000123.txt"] = archivos["AF000123.txt"].replace("SANITAS EPS", "SANITAS EPS NOMBRE MUY LARGO QUE EXCEDE TREINTA")
    fabrica.escribir(tmp_path / "Hospital San Lorenzo" / "05. Mayo" / "r", archivos)
    res = procesar([tmp_path / "Hospital San Lorenzo"])
    assert set(res.incidencias["regla"]) == {"LONGITUD_EXCEDIDA"}
    assert set(res.incidencias["hospital"]) == {"Hospital San Lorenzo"}
    assert set(res.incidencias["periodo"]) == {"2026-05"}
