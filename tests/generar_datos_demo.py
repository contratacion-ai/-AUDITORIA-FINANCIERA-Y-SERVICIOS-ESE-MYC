"""Genera archivos de DEMOSTRACIÓN (datos ficticios) con el mismo formato de Software 360 y de las plantillas:

    python tests/generar_datos_demo.py [carpeta_salida]

- facturacion_ambito_demo.xlsx  · facturación por ámbito (hojas ENERO 2026, FEBRERO 2026, PARAM CONTRATO, PARAM SERVICIO)
- facturas_generadas_demo.xlsx  · informe de facturas generadas (una fila por factura, títulos arriba como en 360)
- farmacia_demo.xlsx            · dispensación / consumo de farmacia por sede

No contiene datos reales de pacientes: documentos y nombres son inventados."""
import os, random, sys
from datetime import date, timedelta
from openpyxl import Workbook

SEDES = ["HOSPITAL SAN ANTONIO", "HOSPITAL SAN FRANCISCO", "HOSPITAL SAN LORENZO", "CENTRO DE SALUD JESUS CORONEL"]
EPS = [("EPS001", "NUEVA EMPRESA PROMOTORA DE SALUD", "900156264"), ("EPS002", "ENTIDAD PROMOTORA DE SALUD SANITAS S A S", "800251440"),
       ("EPS003", "COOSALUD ENTIDAD PROMOTORA DE SALUD S A", "900226715")]
CONTRATOS = {"NUEVA EMPRESA PROMOTORA DE SALUD": [("NUEVA EPS EVENTO SUBSIDIADO", "EVENTO"), ("NUEVA EPS SUBSIDIADO PYM GLOBAL", "CAPITA")],
             "ENTIDAD PROMOTORA DE SALUD SANITAS S A S": [("SANITAS REGIMEN SUBSIDIADO EVENTO", "EVENTO")],
             "COOSALUD ENTIDAD PROMOTORA DE SALUD S A": [("COOSALUD CAPITA SUBSIDIADO", "CAPITA")]}
PROC = [("CONSULTAS", "890201", "CONSULTA DE PRIMERA VEZ POR MEDICINA GENERAL", 32200), ("LABORATORIO CLINICO", "903841", "GLUCOSA EN SUERO", 9800),
        ("ODONTOLOGIA", "997310", "CONTROL DE PLACA DENTAL", 21400), ("IMAGENOLOGIA", "871121", "RADIOGRAFIA DE TORAX", 46860)]
MEDS = [("MEDICAMENTOS", "G-0052", "LACTATO DE RINGER SOLUCION INYECTABLE 500 ML", 6900), ("MEDICAMENTOS", "G-0101", "ACETAMINOFEN 500 MG TABLETA", 120),
        ("MEDICAMENTOS", "G-0233", "DICLOFENACO 75 MG/3 ML AMPOLLA", 1450), ("MEDICAMENTOS", "G-0410", "AMOXICILINA 500 MG CAPSULA", 380),
        ("MATERIALES E INSUMOS MEDICOS", "D-0015", "JERINGA DESECHABLE 5 ML", 650), ("MATERIALES E INSUMOS MEDICOS", "D-0077", "GUANTES DE EXAMEN TALLA M (PAR)", 900),
        ("MATERIALES E INSUMOS MEDICOS", "D-0120", "CATETER INTRAVENOSO No 20", 4200), ("SUPLEMENTOS DIETARIOS", "S-0003", "SUPLEMENTO NUTRICIONAL EN POLVO 400 G", 38000)]
HDR = ["CONTRATO", "SEDE", "SERVICIO", "AMBITO", "REGIMEN DEL PACIENTE", "COD. PROCEDIMIENTO", "COD. MEDICAMENTO", "CANTIDAD", "VALOR ACTIVIDAD",
       "VALOR TOTAL", "FACTURA NRO", "FECHA DE FACTURACION", "CODIGO EPS", "NOMBRE EPS", "NUMERO IDE.", "ESPECIALIDAD", "AGENDA", "PROFESIONAL",
       "REGULADO", "NOMBRE DIAGNÓSTICO", "ACTIVIDAD", "TIPO DE PACIENTE"]

def facturacion(path, facturas_por_mes=120, seed=7):
    rnd = random.Random(seed); wb = Workbook(write_only=True); fact = []
    pc = wb.create_sheet("PARAM CONTRATO"); pc.append([]); pc.append([]); pc.append([None, "CONTRATO", "MODALIDAD"])
    for cs in CONTRATOS.values():
        for c, m in cs: pc.append([None, c, m])
    ps = wb.create_sheet("PARAM SERVICIO"); ps.append([]); ps.append([]); ps.append([None, "SERVICIO"])
    for s in sorted({p[0] for p in PROC + MEDS}): ps.append([None, s])
    for mes, nombre in ((1, "ENERO 2026"), (2, "FEBRERO 2026")):
        ws = wb.create_sheet(nombre); ws.append(HDR)
        for n in range(facturas_por_mes):
            cod, eps_n, _ = rnd.choice(EPS); contrato, modal = rnd.choice(CONTRATOS[eps_n]); sede = rnd.choice(SEDES)
            fn = f"FE{mes}{n:05d}"; dia = min(28, int(rnd.triangular(1, 28, 27))); f = date(2026, mes, dia)
            amb = rnd.choice(["AMBULATORIO", "URGENCIAS", "HOSPITALIZACION"]); ide = str(10_000_000 + rnd.randint(1, 400)); total = 0.0
            items = [rnd.choice(PROC)] + rnd.sample(MEDS, rnd.randint(0, 3))
            for serv, code, act, precio in items:
                med = serv != PROC[0][0] and code[0] in "GDS"
                pr = precio * (1.25 if (med and sede == "HOSPITAL SAN LORENZO" and code == "G-0233") else 1)   # precio no homogéneo a propósito
                cant = rnd.randint(1, 3) if not med else rnd.randint(1, 10)
                vt = 0 if modal == "CAPITA" and med and rnd.random() < .3 else pr * cant; total += vt
                ws.append([contrato, sede, serv, amb, "SUBSIDIADO", "" if med else code, code if med else "", cant, pr, vt, fn, f.strftime("%d/%m/%Y"),
                           cod, eps_n, ide, "MEDICINA GENERAL", "CONSULTA EXTERNA", "PROFESIONAL DEMO", "NO", "DIAGNOSTICO DEMO", act, "SUBSIDIADO"])
            fact.append((mes, fn, f, sede, eps_n, contrato, modal, amb, total, ide))
    wb.save(path); return fact

def facturas_generadas(path, fact, seed=11):
    rnd = random.Random(seed); wb = Workbook(write_only=True); ws = wb.create_sheet("FACTU GENERADA ENERO")
    for t in ("ESE DEPARTAMENTAL DE PRIMER NIVEL MORENO Y CLAVIJO", "INFORME DE FACTURAS GENERADAS", "FECHA DE CONSULTA DESDE: 01/01/2026,HASTA 31/01/2026",
              "SEDE: Todas", "TIPO FACTURA: Factura", None): ws.append([t])
    ws.append(["NRO. INGRESO", "GENERADA EN", "NUMERO FACTURA", "ESTADO FACTURA", "FECHA DE FACTURA", "SUBTOTAL", "VALOR CUOTA", "VALOR COPAGO", "VALOR BONO", "IVA",
               "TOTAL", "VALOR DESCUENTO", "PARTICULAR", "REGISTRADA POR", "ESTADO AUDITORIA", "FECHA DE ANULACIÓN", "TIPO IDE.", "NUMERO IDE.", "PACIENTE",
               "FECHA INGRESO", "FECHA EGRESO", "RECIBO CAJA NRO", "VALOR", "TIPO CONTRATO", "TERCERO", "NIT", "CONTRATO", "REGIMEN DEL CONTRATO",
               "FECHA RADICADO", "NRO. RADICADO EPS", "RADICADO ACTUAL", "VALOR EXCENDENTE", "SEDE", "ESTADO ENVIO DIAN", "CUFE", "CLASE CONTRATO", "CLASE DE ATENCIÓN"])
    nit = {e[1]: e[2] for e in EPS}
    for i, (mes, fn, f, sede, eps_n, contrato, modal, amb, total, ide) in enumerate(x for x in fact if x[0] == 1):
        est = rnd.choices(["Factura Radicada", "Glosada", "Pendiente Por Radicar"], [80, 12, 8])[0]
        rad = "" if est.startswith("Pendiente") else f + timedelta(days=rnd.choice([5, 10, 17, 24, 40]))
        cuota = rnd.choice([0, 0, 0, 4700]); ing = f - timedelta(days=rnd.choice([0, 1, 3, 45]))
        valor = round(total) if i % 37 else round(total) + 1000                                   # encabezado ≠ detalle a propósito cada 37
        ws.append([str(2_000_000 + i), amb, fn, est, f, valor + cuota + (500 if i % 41 == 0 else 0), cuota, 0, 0, 0, valor,   # F06 sembrado cada 41
                   0, "NO", rnd.choice(["FACTURADOR UNO", "FACTURADOR DOS", "FACTURADOR TRES"]),
                   "RADICADA" if rad else "LISTA PARA RADICAR", None, "CC", ide, "PACIENTE DEMO", ing, ing + timedelta(days=1) if amb != "AMBULATORIO" else None,
                   "0", "0.00", "EVENTO AGRUPADO" if modal == "EVENTO" else "CAPITADA", eps_n, nit[eps_n], contrato, "SUBSIDIADO", rad or None,
                   "LOTE DEMO" if rad else None, "1" if rad else "0", "0.00", sede, "ENVIADA" if i % 50 else "ENVIADA Y NO RECIBIDA", "" if i % 50 == 0 else f"cufe{i:08d}",
                   "RÉGIMEN SUBSIDIADO", "REGIMEN SUBSIDIADO"])
    wb.save(path)

def farmacia(path, seed=13):
    rnd = random.Random(seed); wb = Workbook(write_only=True); ws = wb.create_sheet("DISPENSACION")
    ws.append(["MES", "SEDE", "CODIGO", "DESCRIPCION", "TIPO", "CONCEPTO", "CANTIDAD", "COSTO_UNITARIO", "VALOR_COSTO"])
    for mes in ("2026-01", "2026-02"):
        for sede in SEDES:
            for serv, code, desc, precio in MEDS:
                tipo = "MEDICAMENTO" if serv == "MEDICAMENTOS" else "INSUMO" if "INSUMO" in serv else "SUPLEMENTO"
                ws.append([mes, sede, code, desc, tipo, "DISPENSACION", rnd.randint(5, 120), round(precio * rnd.uniform(.55, 1.05), 2), None])
    wb.save(path)

def generar(carpeta):
    os.makedirs(carpeta, exist_ok=True)
    rutas = {k: os.path.join(carpeta, k) for k in ("facturacion_ambito_demo.xlsx", "facturas_generadas_demo.xlsx", "farmacia_demo.xlsx")}
    fact = facturacion(rutas["facturacion_ambito_demo.xlsx"]); facturas_generadas(rutas["facturas_generadas_demo.xlsx"], fact); farmacia(rutas["farmacia_demo.xlsx"])
    return rutas

if __name__ == "__main__":
    for r in generar(sys.argv[1] if len(sys.argv) > 1 else "datos_demo").values(): print("✔", r)
