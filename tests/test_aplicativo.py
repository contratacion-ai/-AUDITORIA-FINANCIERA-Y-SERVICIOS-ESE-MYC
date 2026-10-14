"""Integración con el aplicativo de Auditoría ESE MyC v5 (plantilla RIPS y API)."""

from __future__ import annotations

import json
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import openpyxl
import pytest

from rips_auditoria import procesar
from rips_auditoria.aplicativo import (
    COLUMNAS,
    HOJA_DATOS,
    ClienteAplicativo,
    ErrorAplicativo,
    construir_filas,
    enviar_meses,
    exportar_para_aplicativo,
    limpiar_contrato,
)
from rips_auditoria.cli import main

from . import fabrica


def _lote(tmp_path, hospital, carpeta="10260__FIDUPREVISORA_EVENTO", mes="05", factura="FV1001"):
    raiz = tmp_path / hospital
    fabrica.escribir(raiz / f"{mes}. Mayo" / carpeta / "Rips", fabrica.envio_valido(factura=factura, mes=mes))
    return procesar([raiz])


@pytest.mark.parametrize("carpeta,esperado", [
    ("10204   SANITAS SUBSIDIADO ASITENCIAL _TXT", "SANITAS SUBSIDIADO ASITENCIAL"),
    ("10260__FIDUPREVISORA", "FIDUPREVISORA"),
    ("Rips010635_SANTAS EVENTO CONTRIBUTIVO TAME", "SANTAS EVENTO CONTRIBUTIVO TAME"),
    ("particular", "PARTICULAR"),
])
def test_limpiar_contrato(carpeta, esperado):
    assert limpiar_contrato(carpeta) == esperado


def test_filas_en_plantilla_del_aplicativo(tmp_path):
    res = _lote(tmp_path, "Hospital San Antonio")
    d = construir_filas(res)
    assert list(d.columns) == COLUMNAS
    assert set(d["MES"]) == {"2026-05"} and set(d["SEDE"]) == {"HOSPITAL SAN ANTONIO"}
    assert set(d["TIPO_SERVICIO"]) == {"Consultas", "Procedimientos", "Medicamentos", "Otros servicios"}
    assert set(d["EPS"]) == {"SANITAS EPS"} and set(d["CUCON"]) == {"CT-2026-01"}
    assert set(d["CONTRATO"]) == {"FIDUPREVISORA EVENTO"} and set(d["MODALIDAD"]) == {"EVENTO"}
    assert d["VALOR_SERVICIO"].sum() == 35000 + 50000 + 3000 + 20000
    assert set(d["VALOR_FACTURA_FEV"]) == {103000}
    am = d[d["TIPO_SERVICIO"] == "Medicamentos"].iloc[0]
    assert am["CODIGO_CUPS_CUMS"] == "19942561-6" and am["CANTIDAD"] == 2


def test_modalidad_inferida_sin_palabra_en_contrato(tmp_path):
    res = _lote(tmp_path, "H", carpeta="10205   SANITAS SUBSIDIADO PYM_TXT")
    d = construir_filas(res)
    assert set(d["MODALIDAD"]) == {"EVENTO"}  # la factura concilia
    assert d["VALIDACION"].str.contains("MODALIDAD INFERIDA").all()


def test_acumula_sedes_del_mismo_mes(tmp_path):
    salida = tmp_path / "salida"
    acum = tmp_path / "acumulado"
    exportar_para_aplicativo(_lote(tmp_path / "a", "Hospital San Antonio"), salida, carpeta_acumulado=acum)
    [ruta] = exportar_para_aplicativo(_lote(tmp_path / "b", "Hospital San Jose", factura="FB5001"), salida,
                                      carpeta_acumulado=acum)
    assert ruta.name == "RIPS_AUDITORIA_2026-05.xlsx"
    wb = openpyxl.load_workbook(ruta, read_only=True)
    assert wb.sheetnames[0] == HOJA_DATOS
    filas = list(wb[HOJA_DATOS].iter_rows(values_only=True))
    assert list(filas[3]) == COLUMNAS  # encabezado en la fila 4, como la plantilla del aplicativo
    sedes = {f[COLUMNAS.index("SEDE")] for f in filas[4:]}
    assert sedes == {"HOSPITAL SAN ANTONIO", "HOSPITAL SAN JOSE"}
    # Reprocesar una sede reemplaza sólo esa sede
    exportar_para_aplicativo(_lote(tmp_path / "c", "Hospital San Jose", factura="FB5002"), salida, carpeta_acumulado=acum)
    filas = list(openpyxl.load_workbook(ruta, read_only=True)[HOJA_DATOS].iter_rows(values_only=True))[4:]
    facturas = {f[COLUMNAS.index("FACTURA_FEV")] for f in filas}
    assert facturas == {"FV1001", "FB5002"}


class _AplicativoFalso(BaseHTTPRequestHandler):
    cargas: dict = {}

    def log_message(self, *a):
        pass

    def _json(self, datos, codigo=200):
        cuerpo = json.dumps(datos).encode()
        self.send_response(codigo)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(cuerpo)))
        self.end_headers()
        self.wfile.write(cuerpo)

    def do_GET(self):
        if self.path == "/api/health":
            return self._json({"ok": True})
        cid = int(self.path.rsplit("/", 1)[-1])
        return self._json({"id": cid, "estado": "ok", "mensaje": "1 facturas", "filas": 4})

    def do_POST(self):
        cuerpo = self.rfile.read(int(self.headers["Content-Length"]))
        assert self.path == "/api/cargas/rips"
        assert b'name="mes"' in cuerpo and b"2026-05" in cuerpo and b"RIPS_AUDITORIA_2026-05.xlsx" in cuerpo
        cid = len(self.cargas) + 1
        self.cargas[cid] = cuerpo
        return self._json({"cargas": [{"id": cid, "mes": "2026-05"}]})


@pytest.fixture
def aplicativo_falso():
    srv = ThreadingHTTPServer(("127.0.0.1", 0), _AplicativoFalso)
    hilo = threading.Thread(target=srv.serve_forever, daemon=True)
    hilo.start()
    yield f"http://127.0.0.1:{srv.server_address[1]}"
    srv.shutdown()


def test_envio_al_aplicativo(tmp_path, aplicativo_falso):
    archivos = exportar_para_aplicativo(_lote(tmp_path, "H"), tmp_path / "s", carpeta_acumulado=tmp_path / "ac")
    r = enviar_meses(archivos, {"url": aplicativo_falso, "espera_maxima_segundos": 5})
    assert r == [{"mes": "2026-05", "archivo": str(archivos[0]), "carga_id": 1, "estado": "ok",
                  "mensaje": "1 facturas", "filas": 4}]


def test_aplicativo_cerrado():
    assert ClienteAplicativo(puertos=(1, 1), timeout=0.5).detectar() is None
    with pytest.raises(ErrorAplicativo, match="no está abierto"):
        enviar_meses([], {"url": "http://127.0.0.1:1"})


def test_cli_envia_al_aplicativo(tmp_path, aplicativo_falso, capsys, monkeypatch):
    raiz = tmp_path / "Hospital San Antonio"
    fabrica.escribir(raiz / "05. Mayo" / "123__SANITAS" / "Rips", fabrica.envio_valido())
    cfg = tmp_path / "cfg.yaml"
    cfg.write_text(f"aplicativo_auditoria:\n  url: {aplicativo_falso}\n", encoding="utf-8")
    assert main(["procesar", str(raiz), "-s", str(tmp_path / "s"), "-c", str(cfg), "--sin-consolidado",
                 "--enviar-aplicativo"]) == 0
    assert "2026-05: ok" in capsys.readouterr().out


def test_emparejar_con_param_contrato(tmp_path):
    from rips_auditoria.aplicativo import emparejar_contrato

    catalogo = [("NUEVA EPS SUBSIDIADO PYM GLOBAL", "CAPITA"), ("FIDUPREVISORA EVENTO MAGISTERIO", "EVENTO")]
    assert emparejar_contrato("NUEVA EPS SUBSIDIADO PYM GLOBAL", catalogo) == catalogo[0]
    assert emparejar_contrato("NUEVA EPS SUBISDIADO PYM GLOBAL", catalogo) == catalogo[0]  # error de digitación
    assert emparejar_contrato("COOSALUD EVENTO", catalogo) is None
    res = _lote(tmp_path, "H", carpeta="9906  FIDUPREVISORA EVENTO MAGISTERIO_TXT")
    d = construir_filas(res, contratos_app=catalogo)
    assert set(d["CONTRATO"]) == {"FIDUPREVISORA EVENTO MAGISTERIO"} and set(d["MODALIDAD"]) == {"EVENTO"}
    assert d["VALIDACION"].str.contains("PARAM CONTRATO").all()
