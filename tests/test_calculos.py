import io

import openpyxl
import pandas as pd
import pytest

from rips_auditoria import maestros
from rips_auditoria.analisis import Resultado, analizar
from rips_auditoria.calculos import capita, cruces
from rips_auditoria.clasificacion import clasificar
from rips_auditoria.constantes import CONTROL
from rips_auditoria.ingesta.lote import procesar
from rips_auditoria.referencias import oficial
from rips_auditoria.servicios import construir


def _servicios_control() -> tuple[pd.DataFrame, pd.DataFrame]:
    """Servicios sintéticos con las frecuencias CUPS/CUMS, pacientes y asistencias del control OFICIAL
    (libro vigente) y una NT cuyo costo unitario × frecuencia da el costo mensual oficial: los CUPS
    reparten el costo no farmacéutico por grupo y los CUMS llevan el grupo MEDICAMENTOS."""
    o = oficial.cargar()
    grupos_nm = {g: v for g, v in o["costo_grupo_total"].items() if g != "MEDICAMENTOS"}
    tot_nm = sum(grupos_nm.values())
    gl = list(grupos_nm)
    filas, nt = [], []
    for m, r in enumerate(oficial.detalle_mensual(), start=1):
        per = r["mes"]
        P, A, Q, C = r["pacientes"], r["asistencias"], r["cups"], r["cums"]
        base, resto = divmod(Q, A)
        cant_grupo = [0] * len(gl)
        comun = {"num_factura": f"F{m}", "mes": per, "periodo": per, "municipio": "TAME", "hospital": "H",
                 "descripcion": None, "valor_unitario": 0.0, "valor_total": 0.0, "finalidad": "1",
                 "tipo_servicio_at": None, "es_episodio": False, "modalidad_declarada": "CAPITA", "eps_objetivo": True}
        for j in range(A):
            g = j % len(gl)
            q = base + (1 if j < resto else 0)
            cant_grupo[g] += q
            filas.append({**comun, "tipo_archivo": "AP", "paciente": f"P{m}-{j % P}",
                          "fecha": f"{1 + j // P:02d}/{m:02d}/2026", "codigo": f"T{m:02d}{g:02d}",
                          "tipo_codigo": "CUPS", "cantidad": float(q)})
        bm, rm = divmod(C, P)
        for j in range(P):   # medicamentos en los mismos contactos (no agregan asistencias)
            filas.append({**comun, "tipo_archivo": "AM", "paciente": f"P{m}-{j}", "fecha": f"01/{m:02d}/2026",
                          "codigo": f"M{m:02d}", "tipo_codigo": "CUMS", "cantidad": float(bm + (1 if j < rm else 0))})
        no_med = r["costo"] - r["medicamentos"]
        for g, grupo in enumerate(gl):
            nt.append({"codigo": f"T{m:02d}{g:02d}", "descripcion": grupo, "grupo": grupo, "tipo_codigo": "CUPS",
                       "municipio": "TODOS", "mes": None, "frecuencia_prevista": None,
                       "costo_unitario": no_med * grupos_nm[grupo] / tot_nm / cant_grupo[g]})
        nt.append({"codigo": f"M{m:02d}", "descripcion": "MEDICAMENTOS", "grupo": "MEDICAMENTOS", "tipo_codigo": "CUMS",
                   "municipio": "TODOS", "mes": None, "frecuencia_prevista": None,
                   "costo_unitario": r["medicamentos"] / C})
    return pd.DataFrame(filas), pd.DataFrame(nt)


def test_motor_capita_reproduce_totales_de_control():
    """Pipeline completo (clasificación → costeo NT × RIPS → grupos → detalle → KPI) = control oficial."""
    s, nt = _servicios_control()
    s = clasificar(s, None, nt)
    assert (s["modalidad"] == "CAPITA").all()
    c = capita.costear(s, nt)
    meses = sorted(c["mes"].unique())
    grupos = capita.costos_grupo_mes(c, meses)
    det = capita.detalle_mensual(c, grupos[meses].sum(axis=0), meses)
    assert abs(det["costo"].sum() - CONTROL["capita_total"]) <= 1
    assert abs(grupos.loc["MEDICAMENTOS", "acumulado"] - CONTROL["medicamentos"]) <= 1
    assert det["cups"].sum() == CONTROL["cups"]
    assert det["cums"].sum() == CONTROL["cums"]
    assert det["pacientes"].sum() == CONTROL["pacientes_unicos"]
    assert det["asistencias"].sum() == CONTROL["asistencias"]
    k = capita.kpis(det["costo"].tolist(), det["pacientes"].tolist(), det["asistencias"].tolist())
    assert round(k.promedio_mes) == CONTROL["capita_promedio_mes"]
    assert abs(k.costo_usuario_prom - CONTROL["costo_usuario_mes"]) <= 1


def test_clasificacion_precedencia_y_trazabilidad():
    s = pd.DataFrame([
        {"codigo": "A1", "modalidad_declarada": "EVENTO"},     # OFERTADO en catálogo -> contractual CAPITA
        {"codigo": "B1", "modalidad_declarada": "CAPITA"},     # NO OFERTADO -> contractual EVENTO
        {"codigo": "C1", "modalidad_declarada": None},         # solo en NT -> CAPITA
        {"codigo": "D1", "modalidad_declarada": "PGP"},        # sin marca -> declarada
        {"codigo": "E1", "modalidad_declarada": None},         # sin nada -> EVENTO por defecto
    ]).assign(tipo_archivo="AP", descripcion=None, finalidad="1", tipo_servicio_at=None, es_episodio=False)
    cat = pd.DataFrame([{"codigo": "A1", "grupo": "LABORATORIO", "estado": "OFERTADO"},
                        {"codigo": "B1", "grupo": "PROCEDIMIENTOS", "estado": "NO OFERTADO"}])
    nt = pd.DataFrame([{"codigo": "C1", "grupo": "IMAGENOLOGIA", "costo_unitario": 10.0}])
    r = clasificar(s, cat, nt).set_index("codigo")
    assert r.loc["A1", "modalidad_catalogo"] == "CAPITA" and r.loc["A1", "regla_modalidad"] == "CATALOGO_OFERTADO"
    assert r.loc["A1", "modalidad"] == "EVENTO" and r.loc["A1", "discrepancia_modalidad"]
    assert r.loc["B1", "modalidad_catalogo"] == "EVENTO" and r.loc["B1", "discrepancia_modalidad"]
    assert r.loc["C1", "regla_modalidad"] == "NOTA_TECNICA" and r.loc["C1", "modalidad"] == "CAPITA"
    assert r.loc["C1", "grupo"] == "IMAGENOLOGIA"
    assert r.loc["D1", "regla_modalidad"] == "MODALIDAD_DECLARADA" and r.loc["D1", "modalidad"] == "PGP"
    assert r.loc["E1", "regla_modalidad"] == "POR_DEFECTO" and r.loc["E1", "modalidad"] == "EVENTO"
    assert r.loc["E1", "regla_grupo"] in ("HEURISTICA", "REGLA_DASHBOARD")


def test_cruces_nt_tarifario_y_facturacion(carpeta_rips):
    l = procesar([carpeta_rips])
    s = construir(l)
    s["modalidad_declarada"] = None          # sin modalidad en la carpeta: decide catálogo/NT
    nt = pd.DataFrame([{"codigo": "902207", "grupo": "LABORATORIO", "municipio": "TAME", "mes": None,
                        "frecuencia_prevista": 0.5, "costo_unitario": 9000.0, "descripcion": "HEMOGRAMA"}])
    s = clasificar(s, None, nt)
    x = cruces.rips_vs_nota_tecnica(s, nt, 10)
    fila = x[x.codigo == "902207"].iloc[0]
    assert fila.estado == "SOBREPASO" and fila.desviacion == 0.5 and fila.desviacion_pct == 100
    tar = pd.DataFrame([{"codigo": "906249", "tarifa": 60000, "vigencia_desde": "01/01/2026", "vigencia_hasta": None},
                        {"codigo": "890201", "tarifa": 30000, "vigencia_desde": "01/01/2026", "vigencia_hasta": None}])
    t = cruces.evento_vs_tarifario(s, tar, 1).set_index("codigo")
    assert t.loc["906249", "estado"] == "SOBRE_TARIFA" and t.loc["906249", "dif_valor_total"] == 10000
    assert t.loc["890201", "estado"] == "CONFORME"
    assert t.loc["S31307", "estado"] == "SIN_TARIFA"
    fac = pd.DataFrame([{"periodo": "2026-08", "municipio": None, "ambito": None, "codigo": "906249", "cantidad": 2,
                         "valor": 130000, "descripcion": "VIH"},
                        {"periodo": "2026-08", "municipio": None, "ambito": None, "codigo": "999999", "cantidad": 1,
                         "valor": 5000, "descripcion": "SIN RIPS"}])
    f = cruces.rips_vs_facturacion(s, fac).set_index("codigo")
    assert f.loc["906249", "estado"] == "CONCILIADO"
    assert f.loc["999999", "estado"] == "FACTURADO_SIN_RIPS" and f.loc["999999", "glosa_potencial"] == 5000
    assert f.loc["890201", "estado"] == "RIPS_SIN_FACTURA"
    cf = cruces.conciliacion_facturas(l.tablas)
    assert set(cf["estado"]) == {"CONCILIADA"}


def test_costo_por_municipio_tiene_prioridad():
    s = pd.DataFrame([{"codigo": "X", "municipio": "TAME", "modalidad": "CAPITA", "es_episodio": False, "cantidad": 2.0},
                      {"codigo": "X", "municipio": "ARAUCA", "modalidad": "CAPITA", "es_episodio": False, "cantidad": 2.0}])
    nt = pd.DataFrame([{"codigo": "X", "municipio": "TAME", "costo_unitario": 100.0},
                       {"codigo": "X", "municipio": "TODOS", "costo_unitario": 50.0}])
    c = capita.costear(s, nt)
    assert c["costo_capita"].tolist() == [200.0, 100.0]


def test_maestro_informe_nt_costos_y_marcacion():
    """Libro mínimo con el diseño real del INFORME DE SERVICIOS 2026 (hojas MARCACION, NT COSTOS y
    NOTA TECNICA CAPITA NEPS)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MARCACION CUPS CAPITA"
    ws.append([None, "CodigoCUP", "Homologo", "DescripcionCUP", "Grupo", "Subgrupo", "Ruta", "OFERTA 03/09/2026",
               "MARCACION", "UVB"])
    ws.append([None, 902207, 902207, "HEMOGRAMA I", "LABORATORIO", "HEMATOLOGIA", "MORBILIDAD", "OFERTADO", "X", 43600])
    ws.append([None, "MD000013", "225064-7", "HIOSCINA", "MEDICAMENTOS", "MEDICAMENTOS", "MORBILIDAD", "OFERTADO", "X", None])
    ws.append([None, "S31307", None, "TRASLADO", "TRASLADOS", None, None, "NO OFERTADO", None, None])
    nt = wb.create_sheet("NOTA TECNICA CAPITA NEPS")
    nt.append([None, "SERVICIOS CONTRATO CAPITA", None, None, None, "FRECUENCIAS MENSUALES"])
    nt.append([None, "Homologo", "CodigoCUPS", "DescripcionCUPS", "Grupo", "TAME", "ARAUQUITA", "FORTUL",
               "CRAVO NORTE", "PUERTO RONDON"])
    nt.append([None, 902207, "902207", "HEMOGRAMA I", "LABORATORIO", 10, 12, 6, 0.5, 0.5])
    nt.append([None, "225064-7", "MD000013", "HIOSCINA", "MEDICAMENTOS", 3, 3, 1, 0, 0])
    nt.append([None, None, None, None, "", 13, 15, 7, 0.5, 0.5])
    co = wb.create_sheet("NT COSTOS")
    co.append(["titulo"])
    co.append([None, "Homologo", "CodigoCUPS", "DescripcionCUPS", "Grupo", "TOTAL SERVICIOS NOTA TECNICA", "ENERO",
               "FEBRERO", "TOTAL", "COSTO CUPS ", "COSTO ENERO A AGOSTO", "ENERO", "FEBRERO", "TOTAL"])
    co.append([None, 902207, 902207, "HEMOGRAMA I", "LABORATORIO", 29, 10, 30, 40, 10000, 400000, 100000, 300000, 400000])
    co.append([None, "225064-7", "MD000013", "HIOSCINA", "MEDICAMENTOS", 7, 4, 0, 4, 1255, 5020, 5020, 0, 5020])
    co.append([None, None, None, None, "", 36, 14, 30, 44, None, 405020, 105020, 300000, 405020])
    buf = io.BytesIO()
    wb.save(buf)
    res = maestros.cargar("informe", buf.getvalue(), "INFORME.xlsx")
    assert res["catalogo"]["filas"] == 3
    cat = maestros.leer("catalogo").set_index("codigo")
    assert cat.loc["902207", "estado"] == "OFERTADO" and cat.loc["S31307", "estado"] == "NO OFERTADO"
    assert cat.loc["MD000013", "homologo"] == "225064-7"
    ntd = maestros.leer("nota_tecnica")
    assert set(ntd["municipio"]) == {"TAME", "ARAUQUITA", "FORTUL", "CRAVO NORTE", "PUERTO RONDON"}
    assert (ntd.set_index("codigo").loc["902207", "costo_unitario"] == 10000).all()
    # El RIPS AM trae el CUM (homólogo): se clasifica y costea con la NT
    s = pd.DataFrame([{"tipo_archivo": "AM", "codigo": "225064-7", "cantidad": 2.0, "mes": "2026-01",
                       "municipio": "TAME", "descripcion": None, "finalidad": None, "tipo_servicio_at": None,
                       "es_episodio": False, "modalidad_declarada": None}])
    c = capita.costear(clasificar(s, maestros.leer("catalogo"), ntd), ntd)
    assert c.loc[0, "modalidad"] == "CAPITA" and c.loc[0, "codigo_nt"] == "MD000013"
    assert c.loc[0, "costo_capita"] == 2510
    assert maestros.leer_resumen_informe()["capita_mensual"] == [105020, 300000]


def test_maestro_tarifario_plantilla_y_pendiente():
    csv = "codigo;descripcion;tarifa;vigencia_desde;vigencia_hasta\n906249;VIH;135.149,09;01/01/2026;\n"
    r = maestros.cargar("tarifario", csv.encode(), "tarifario.csv")
    assert r["tarifario"]["notas"]          # marcado pendiente de confirmación
    t = maestros.leer("tarifario")
    assert t.iloc[0]["tarifa"] == pytest.approx(135149.09)
    est = {e["nombre"]: e for e in maestros.estado()}
    assert est["tarifario"]["pendiente_confirmacion"]


def test_analisis_completo_informes_y_json(carpeta_rips, tmp_path):
    from rips_auditoria.informes import excel, pdf
    l = procesar([carpeta_rips])
    r = analizar([l])
    assert r.indicadores["evento_total"] > 0
    assert not r.tablas["top20_evento"].empty
    assert (r.tablas["riesgo"]["codigo"] == "S31307").any()
    x = excel.escribir(r, tmp_path / "i.xlsx")
    hojas = openpyxl.load_workbook(x).sheetnames
    for h in ("Resumen", "CAPITA Grupos", "Detalle Mensual CAPITA", "TOP 20 EVENTO", "Simulación E→C", "Escenarios",
              "Riesgo", "Hallazgos", "NT vs RIPS", "EVENTO vs Tarifario", "RIPS vs Facturación", "Trazabilidad"):
        assert h in hojas
    # Anonimización: ningún documento ni nombre del paciente en el informe
    contenido = " ".join(str(c.value) for ws in openpyxl.load_workbook(x).worksheets for row in ws.iter_rows()
                         for c in row if c.value is not None)
    for dato in ("1000001", "2000002", "PEREZ", "LUIS"):
        assert dato not in contenido
    p = pdf.escribir(r, tmp_path / "i.pdf")
    assert p.read_bytes()[:4] == b"%PDF"
    r2 = Resultado.desde_json(r.a_json())
    assert r2.indicadores == r.indicadores
    assert r2.tablas["escenarios"].shape == r.tablas["escenarios"].shape


def test_dos_lotes_con_los_mismos_archivos_no_duplican(carpeta_rips):
    l1, l2 = procesar([carpeta_rips]), procesar([carpeta_rips])
    uno, dos = analizar([l1]), analizar([l1, l2])
    assert dos.indicadores["evento_total"] == uno.indicadores["evento_total"]
    assert any("más de un lote" in a for a in dos.advertencias)
