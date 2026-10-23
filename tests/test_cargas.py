"""Cargas: validación de estructura, limpieza, resolución de códigos, duplicados, anulación y mapeo por institución."""
import io
from decimal import Decimal as D

import pytest
from openpyxl import Workbook

from cargas.models import Carga, ErrorCarga, Facturacion, MovimientoContable, PlantillaMapeo, Rips
from cargas.plantilla import generar_plantilla
from cargas.procesador import ErrorCargaArchivo, anular_carga, procesar_carga
from modelo.models import CuentaRecurso, DistribucionManoObra, HomologacionCentro, ValorInductor


def xlsx(encabezados, filas, hoja="Hoja1", filas_previas=0):
    wb = Workbook()
    ws = wb.active
    ws.title = hoja
    for _ in range(filas_previas):
        ws.append(["Reporte generado por el sistema contable"])
    ws.append(encabezados)
    for f in filas:
        ws.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def cuentas(caso):
    CuentaRecurso.objects.create(institucion=caso.inst, prefijo="5101", recurso=caso.R1)
    CuentaRecurso.objects.create(institucion=caso.inst, prefijo="511117", recurso=caso.R14)
    CuentaRecurso.objects.create(institucion=caso.inst, prefijo="5815", excluir=True)
    HomologacionCentro.objects.create(institucion=caso.inst, codigo_origen="ESE01-TAME-URGENCIAS", centro=caso.F2)
    return caso


def cargar(caso, tipo, contenido, nombre="archivo.xlsx", **kw):
    return procesar_carga(institucion=caso.inst, usuario=None, tipo_archivo=tipo, contenido=contenido, nombre=nombre,
                          periodo=kw.pop("periodo", caso.periodo), **kw)


def test_estado_financiero_estandar_limpieza_y_resolucion(cuentas):
    c = cuentas
    contenido = xlsx(["CUENTA", "NOMBRE CUENTA", "CENTRO DE COSTOS", "VALOR"], [
        ["510101", "SUELDOS", "F1", "1.234.567,89"],
        ["5101.02", "HORAS EXTRA", "ese01-tame-urgencias", 1000],   # homologado, sin importar mayúsculas
        ["51111701", "ENERGIA", None, "$ 2.000"],
        ["581501", "VIGENCIAS ANTERIORES", None, 999],              # excluida
        ["510101", "SUELDOS", "F1", 0],                             # valor cero → omitida
        [None, None, None, None],                                   # vacía
        ["510101", "SUELDOS", "F1", "1.234.567,89"],                # duplicada
    ])
    res = cargar(c, "ESTADO_FINANCIERO", contenido)
    assert res.carga.estado == Carga.Estado.APLICADA
    assert res.carga.filas_leidas == 6 and res.carga.filas_validas == 3
    assert res.carga.filas_duplicadas == 1 and res.carga.filas_omitidas == 2
    movs = {m.fila: m for m in MovimientoContable.objects.filter(carga=res.carga)}
    assert movs[2].valor == D("1234567.89") and movs[2].recurso == c.R1 and movs[2].centro == c.F1
    assert movs[3].centro == c.F2 and movs[3].cuenta == "510102"
    assert movs[4].recurso == c.R14 and movs[4].centro is None and movs[4].valor == D("2000")
    assert res.carga.valor_total == D("1234567.89") + D("1000") + D("2000")


def test_errores_rechazan_la_carga_completa(cuentas):
    contenido = xlsx(["CUENTA", "CENTRO DE COSTOS", "VALOR"], [
        ["510101", "F1", 100], ["999999", "F1", 50], ["510101", "NO-EXISTE", 20], ["510101", "F1", "abc"]])
    res = cargar(cuentas, "ESTADO_FINANCIERO", contenido)
    assert res.carga.estado == Carga.Estado.CON_ERRORES
    assert MovimientoContable.objects.filter(carga=res.carga).count() == 0
    mensajes = list(ErrorCarga.objects.filter(carga=res.carga, nivel="ERROR").values_list("fila", "mensaje"))
    assert {f for f, _ in mensajes} == {3, 4, 5}
    # Aplicación parcial explícita
    res2 = cargar(cuentas, "ESTADO_FINANCIERO", contenido, nombre="otro.xlsx", aplicar_parcial=True)
    assert res2.carga.estado == Carga.Estado.APLICADA and res2.carga.filas_validas == 1


def test_estructura_invalida(cuentas):
    with pytest.raises(ErrorCargaArchivo, match="VALOR"):
        cargar(cuentas, "ESTADO_FINANCIERO", xlsx(["CUENTA", "CENTRO"], [["5101", "F1"]]))
    with pytest.raises(ErrorCargaArchivo, match="Extensión"):
        cargar(cuentas, "ESTADO_FINANCIERO", b"hola", nombre="malicioso.exe")
    with pytest.raises(ErrorCargaArchivo, match="periodo"):
        cargar(cuentas, "ESTADO_FINANCIERO", xlsx(["CUENTA", "VALOR"], [["5101", 1]]), periodo=None)


def test_archivo_repetido_y_anulacion(cuentas):
    contenido = xlsx(["CUENTA", "CENTRO DE COSTOS", "VALOR"], [["510101", "F1", 100]])
    res = cargar(cuentas, "ESTADO_FINANCIERO", contenido)
    with pytest.raises(ErrorCargaArchivo, match="ya fue aplicado"):
        cargar(cuentas, "ESTADO_FINANCIERO", contenido)
    assert anular_carga(res.carga, None, "Archivo equivocado") == 1
    assert MovimientoContable.objects.filter(carga=res.carga).count() == 0
    cargar(cuentas, "ESTADO_FINANCIERO", contenido)  # tras anular se puede volver a cargar


def test_periodo_cerrado_rechaza(cuentas):
    cuentas.periodo.estado = "CERRADO"
    cuentas.periodo.save()
    with pytest.raises(ErrorCargaArchivo, match="cerrado"):
        cargar(cuentas, "ESTADO_FINANCIERO", xlsx(["CUENTA", "VALOR"], [["510101", 1]]))


def test_mapeo_configurable_por_institucion(cuentas):
    """Nómina con formato propio: encabezado en la fila 3, columnas con otros nombres, coma decimal y recurso fijo."""
    plantilla = PlantillaMapeo.objects.create(
        institucion=cuentas.inst, nombre="Nómina software X", tipo_archivo="NOMINA", fila_encabezado=3,
        mapeo={"documento_tercero": "DocumentoTercero", "nombre_tercero": "NombreTercero",
               "centro_costo": "NombreCentroCosto", "valor": "#E"},
        constantes={"recurso": "R1"}, separador_decimal=",")
    contenido = xlsx(["DocumentoTercero", "NombreTercero", "NombreCentroCosto", "Cargo", "Total devengado"],
                     [["1.000.000.001", "EMPLEADO UNO", "ESE01-TAME-URGENCIAS", "MEDICO", "4.500.000,50"]], filas_previas=2)
    res = cargar(cuentas, "NOMINA", contenido, plantilla=plantilla)
    m = MovimientoContable.objects.get(carga=res.carga)
    assert (m.fuente, m.recurso, m.centro, m.valor, m.tercero_documento) == ("AREA", cuentas.R1, cuentas.F2, D("4500000.50"), "1000000001")
    assert m.fila == 4


def test_mano_obra_valida_suma_100(cuentas):
    contenido = xlsx(["TIPO", "DOCUMENTO", "CENTRO DE COSTOS", "PORCENTAJE"], [
        ["OPS", "900123", "F1", 60], ["OPS", "900123", "F2", 40], ["Empleado", "100", "F1", 70], ["Empleado", "100", "F2", 20]])
    res = cargar(cuentas, "MANO_OBRA", contenido, aplicar_parcial=True)
    assert DistribucionManoObra.objects.filter(documento="900123").count() == 2
    assert not DistribucionManoObra.objects.filter(documento="100").exists()
    assert any("suman 90" in e[3] for e in res.errores)


def test_valores_inductor_permanentes_y_por_origen(cuentas):
    contenido = xlsx(["INDUCTOR", "CENTRO DE COSTOS", "CENTRO ORIGEN", "CANTIDAD", "PERMANENTE"], [
        ["M2", "F1", None, "35,5", "SI"], ["SOLICITUDES", "F2", "I1", 7, None]])
    cargar(cuentas, "VALORES_INDUCTOR", contenido)
    assert ValorInductor.objects.get(inductor=cuentas.M2, centro=cuentas.F1, periodo__isnull=True).cantidad == D("35.5")
    assert ValorInductor.objects.get(inductor=cuentas.SOL, centro=cuentas.F2, centro_origen=cuentas.I1,
                                     periodo=cuentas.periodo).cantidad == 7


def test_facturacion_y_rips_seudonimizan_paciente(cuentas):
    fac = xlsx(["NUMERO FACTURA", "FECHA", "EPS", "CUPS", "CANTIDAD", "VALOR TOTAL", "AMBITO", "TIPO DOCUMENTO PACIENTE",
                "DOCUMENTO PACIENTE"], [["FE1", "15/09/2025", "EPS-X", "890201", 2, "84.000", "AMB", "CC", "1098765432"],
                                        ["FE2", "15/08/2025", "EPS-X", "890201", 1, 42000, "1", "CC", "123"]])
    res = cargar(cuentas, "FACTURACION", fac)
    f = Facturacion.objects.get(carga=res.carga, numero_factura="FE1")
    assert f.valor_unitario == D("42000") and f.cups == cuentas.cups_a and f.eps is None
    assert f.paciente and "1098765432" not in f.paciente
    assert any("fuera del periodo" in a[3] for a in res.advertencias)
    rips = xlsx(["NUMERO FACTURA", "CUPS", "DOCUMENTO PACIENTE"], [["FE1", "890201", "1098765432"]])
    r = Rips.objects.get(carga=cargar(cuentas, "RIPS", rips).carga)
    assert r.paciente == f.paciente and r.cantidad == 1
    assert not res.carga.archivo and len(res.carga.hash_sha256) == 64  # el original con documentos no se conserva


def test_csv_con_punto_y_coma(cuentas):
    contenido = "CUENTA;CENTRO DE COSTOS;VALOR\n510101;F1;1.500,25\n".encode("latin-1")
    res = cargar(cuentas, "ESTADO_FINANCIERO", contenido, nombre="ef.csv")
    assert MovimientoContable.objects.get(carga=res.carga).valor == D("1500.25")


def test_plantilla_estandar_tiene_todas_las_hojas():
    from openpyxl import load_workbook

    from cargas.esquemas import ESQUEMAS

    wb = load_workbook(io.BytesIO(generar_plantilla()))
    assert set(ESQUEMAS) <= set(wb.sheetnames) and "INSTRUCCIONES" in wb.sheetnames
    assert [c.value for c in wb["ESTADO_FINANCIERO"][1]][:2] == ["CUENTA", "NOMBRE CUENTA"]


def test_libro_estandar_completo(cuentas):
    from cargas.procesador import procesar_libro_estandar

    wb = Workbook()
    wb.active.title = "INSTRUCCIONES"
    ef = wb.create_sheet("ESTADO_FINANCIERO")
    ef.append(["CUENTA", "NOMBRE CUENTA", "CENTRO DE COSTOS", "VALOR"])
    ef.append(["510101", "SUELDOS", "F1", 500])
    rips = wb.create_sheet("RIPS")
    rips.append(["NUMERO FACTURA", "CUPS", "DOCUMENTO PACIENTE"])
    rips.append(["FE1", "890201", "123456"])
    wb.create_sheet("NOMINA").append(["DOCUMENTO", "VALOR"])  # solo encabezado: se ignora
    buf = io.BytesIO()
    wb.save(buf)
    salida = procesar_libro_estandar(institucion=cuentas.inst, usuario=None, contenido=buf.getvalue(), nombre="libro.xlsx",
                                     periodo=cuentas.periodo)
    assert [h for h, _ in salida] == ["ESTADO_FINANCIERO", "RIPS"]
    assert all(r.carga.estado == Carga.Estado.APLICADA and not r.carga.archivo for _, r in salida)
