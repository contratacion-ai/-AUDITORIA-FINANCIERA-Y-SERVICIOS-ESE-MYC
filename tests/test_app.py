"""Pruebas de extremo a extremo con datos ficticios (tests/generar_datos_demo.py).  Ejecutar:  python -m pytest -q"""
import io, os, sys, tempfile
from datetime import date
import pytest

DATA = tempfile.mkdtemp(prefix="auditoria_test_")
os.environ["AUDIT_DATA_DIR"] = DATA
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from fastapi.testclient import TestClient                     # noqa: E402
import openpyxl                                               # noqa: E402
import generar_datos_demo as G                                # noqa: E402
from app.main import app                                      # noqa: E402
from app import facturas as F                                 # noqa: E402

cli = TestClient(app)

@pytest.fixture(scope="module")
def cargado():
    rutas = G.generar(os.path.join(DATA, "demo"))
    for tipo, nombre in (("facturacion", "facturacion_ambito_demo.xlsx"), ("facturas", "facturas_generadas_demo.xlsx"), ("farmacia", "farmacia_demo.xlsx")):
        with open(rutas[nombre], "rb") as f:
            r = cli.post(f"/api/cargas/{tipo}", files={"file": (nombre, f, "application/octet-stream")})
        assert r.status_code == 200, r.text
    cargas = cli.get("/api/cargas").json()["filas"]
    assert all(c["estado"] == "ok" for c in cargas), cargas
    return rutas

def _lineas_fuente(ruta, mes_hoja):
    ws = openpyxl.load_workbook(ruta, read_only=True)[mes_hoja]; it = ws.iter_rows(values_only=True); h = next(it)
    return [dict(zip(h, r)) for r in it]

def test_medicamentos_cuadra_con_fuente(cargado):
    filas = _lineas_fuente(cargado["facturacion_ambito_demo.xlsx"], "ENERO 2026")
    esperado = sum(r["VALOR TOTAL"] for r in filas if r["SERVICIO"] in ("MEDICAMENTOS", "MATERIALES E INSUMOS MEDICOS", "SUPLEMENTOS DIETARIOS"))
    k = cli.get("/api/medicamentos/resumen?mes=2026-01").json()["kpis"]
    assert abs(k["valor"] - esperado) < 1
    assert k["medicamentos"] > 0 and k["insumos"] > 0 and k["suplementos"] > 0

def test_precio_no_homogeneo_detectado(cargado):
    p = cli.get("/api/medicamentos/precios?mes=2026-01").json()["filas"]
    g = next(x for x in p if x["codigo"] == "G-0233")
    assert "NO HOMOGÉNEO" in g["estado"] and g["sede_mas_cara"] == "HOSPITAL SAN LORENZO"

def test_hoja_por_sede_y_abc(cargado):
    d = cli.get("/api/medicamentos/sede?mes=2026-01&clase=MEDICAMENTO").json()
    assert "Total general" in d["columnas"] and all(r["clase"] == "MEDICAMENTO" for r in d["filas"])
    it = cli.get("/api/medicamentos/items?mes=2026-01").json()["filas"]
    assert {x["abc"] for x in it} <= {"A", "B", "C"} and it[0]["abc"] == "A"

def test_farmacia_cruce(cargado):
    d = cli.get("/api/medicamentos/farmacia?mes=2026-01").json()
    assert d["total"] > 0 and any(e != "OK" for e in (r["estado"] for r in d["filas"]))

def test_facturas_generadas(cargado):
    n = sum(1 for _ in openpyxl.load_workbook(cargado["facturas_generadas_demo.xlsx"], read_only=True).active.iter_rows(min_row=8))
    r = cli.get("/api/facturas/resumen?mes=2026-01").json()
    assert r["kpis"]["facturas"] == n
    assert abs(r["kpis"]["radicado"] + r["kpis"]["glosado"] + r["kpis"]["pendiente"] - r["kpis"]["valor"]) < 1
    reglas = {x["regla"] for x in r["reglas"]}
    assert {"F05", "F06"} <= reglas
    # el encabezado con +1000 (F06 sembrado) no cuadra con el detalle por ámbito
    c = cli.get("/api/facturas/cruce?mes=2026-01").json()["filas"]
    assert sum("ENCABEZADO ≠ DETALLE" in x["estado"] for x in c) == sum(1 for i in range(n) if i % 37 == 0)
    det = cli.get("/api/facturas/detalle?mes=2026-01&solo_alertas=true").json()["filas"]
    assert det and all(x["alertas"] != "OK" for x in det)

def test_filtro_sede(cargado):
    s = "HOSPITAL SAN ANTONIO"
    r = cli.get(f"/api/facturas/resumen?mes=2026-01&sede={s}").json()
    assert {x["sede"] for x in r["por_sede"]} == {s}
    m = cli.get(f"/api/medicamentos/resumen?mes=2026-01&sede={s}").json()
    assert {x["sede"] for x in m["por_sede"]} == {s}

def test_informe_gerencial_y_excel(cargado):
    g = cli.get("/api/gerencial?mes=2026-01").json()
    assert g["indicadores"] and g["recomendaciones"] and g["narrativa"]
    r = cli.get("/api/export?mes=2026-01"); assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    for h in ("Informe_Gerencial", "FG_Sede", "FG_Alertas", "MED_Sede_Clase", "MED_Medicamentos_x_Sede", "MED_Insumos_x_Sede", "MED_Farmacia"):
        assert h in wb.sheetnames, h

@pytest.mark.parametrize("url", ["/api/resumen", "/api/comparativos", "/api/resumen/empresas", "/api/conciliacion/fact-rips", "/api/conciliacion/tri-eje",
                                 "/api/conciliacion/radicacion", "/api/conciliacion/contratos", "/api/conciliacion/hallazgos", "/api/conciliacion/validacion-cruce",
                                 "/api/costos/resumen", "/api/informes/3_servicio", "/api/filtros", "/api/medicamentos/eps", "/api/facturas/facturadores",
                                 "/api/plantillas/farmacia", "/api/meses", "/"])
def test_endpoints_responden(cargado, url):
    assert cli.get(url).status_code == 200, url

def test_festivos_y_dias_habiles():
    f = F.festivos_co(2026)
    assert date(2026, 4, 2) in f and date(2026, 4, 3) in f and date(2026, 1, 12) in f and date(2026, 6, 29) in f
    import pandas as pd
    d = F.dias_habiles(pd.Series(["2026-01-02"]), pd.Series(["2026-01-16"]))   # festivo 12-ene en medio
    assert d.iloc[0] == 9
