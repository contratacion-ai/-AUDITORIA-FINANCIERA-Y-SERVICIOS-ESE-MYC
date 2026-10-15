"""Tablero de análisis, cargue sin límite de archivos y acceso en red."""

from __future__ import annotations

import io
import json
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import openpyxl
import pandas as pd
import pytest

pytest.importorskip("fastapi")
pytest.importorskip("httpx")
from fastapi.testclient import TestClient  # noqa: E402

from rips_auditoria import tablero as tb  # noqa: E402
from rips_auditoria.aplicativo import COLUMNAS  # noqa: E402
from rips_auditoria.web import red, servidor  # noqa: E402

from . import fabrica  # noqa: E402

LOCAL = ("127.0.0.1", 50000)
REMOTO = ("192.168.1.20", 50001)


@pytest.fixture
def cliente(monkeypatch):
    monkeypatch.setenv("AUDITORIA_ESE_DIR", "/no/existe")
    return TestClient(servidor.app, client=LOCAL)


def _esperar(cliente, tid, maximo=180):
    fin = time.monotonic() + maximo
    while time.monotonic() < fin:
        t = cliente.get(f"/api/cargues/{tid}").json()
        if t["estado"] in ("terminado", "error"):
            return t
        time.sleep(0.3)
    raise AssertionError("el cargue no terminó")


@pytest.fixture
def con_datos(cliente, tmp_path):
    zip_lote = fabrica.zip_varios_hospitales(tmp_path / "lote.zip")
    r = cliente.post("/api/cargues", data={"ruta_local": str(zip_lote), "enviar": "false", "consolidado": "false"})
    t = _esperar(cliente, r.json()["id"])
    assert t["estado"] == "terminado", t["error"]
    return cliente


def _fila(mes, sede, eps="NUEVA EPS", cups="890201", cant=1.0, valor=1000.0, factura="F1", fev=1000.0,
          validacion="CONCILIADA", servicio="Consultas"):
    return dict(zip(COLUMNAS, [mes, eps, "", "CONTRATO X", "", "EVENTO", sede, factura, "", servicio, cups, "",
                               cant, valor, fev, validacion], strict=True))


# ---------------------------------------------------------------------------

def test_cargue_de_mas_de_mil_archivos(cliente):
    """Starlette limita un formulario a 1000 archivos por defecto: una carpeta anual trae miles."""
    archivos = [("archivos", (f"Hospital San Lorenzo/05. Mayo/123__SANITAS/Rips/{n}", c.encode(), "text/plain"))
                for n, c in fabrica.envio_valido().items()]
    archivos += [("archivos", (f"Hospital San Lorenzo/05. Mayo/soportes/LEAME_{i}.txt", b"x", "text/plain"))
                 for i in range(1500)]
    r = cliente.post("/api/cargues", files=archivos, data={"enviar": "false"})
    assert r.status_code == 200, r.text
    t = _esperar(cliente, r.json()["id"])
    assert t["estado"] == "terminado", t["error"]
    assert len(t["entradas"]) == len(archivos)
    assert [(u["hospital"], u["periodo"]) for u in t["resultado"]["unidades"]] == [("Hospital San Lorenzo", "2026-05")]


def test_tablero_sin_datos(cliente):
    assert cliente.get("/api/tablero/opciones").json()["meses"] == []
    d = cliente.get("/api/tablero").json()
    assert d["kpis"]["registros"] == 0 and d["alertas"] == []
    assert cliente.get("/api/tablero/excel").status_code == 404
    assert cliente.get("/tablero").status_code == 200


def test_tablero_despues_del_cargue(con_datos):
    c = con_datos
    o = c.get("/api/tablero/opciones").json()
    assert o["meses"] == ["2026-05", "2026-06"]
    assert o["sede"] == ["HOSPITAL SAN ANTONIO", "HOSPITAL SAN JOSE"]
    d = c.get("/api/tablero").json()
    k = d["kpis"]
    assert k["sedes"] == 2 and k["meses"] == 2 and k["facturas"] == 4
    assert k["usuarios"] == 8 and k["frecuencia_uso"] == 3.5  # indicadores guardados en el cargue
    assert [m["MES"] for m in d["por_mes"]] == ["2026-05", "2026-06"]
    assert d["analisis"]["hallazgos"] and d["analisis"]["recomendaciones"]
    # filtros por sede y mes
    f = c.get("/api/tablero", params={"sede": "HOSPITAL SAN JOSE", "desde": "2026-06"}).json()
    assert f["kpis"]["sedes"] == 1 and f["kpis"]["meses"] == 1 and f["kpis"]["usuarios"] == 2
    # comparativo entre parámetros, de mayor a menor
    comp = c.get("/api/tablero/comparativo", params={"filas": "servicio", "columnas": "sede", "medida": "cantidad"}).json()
    totales = [r["TOTAL"] for r in comp["datos"]]
    assert totales == sorted(totales, reverse=True) and comp["totales"]["TOTAL"] == k["cantidad"]
    assert c.get("/api/tablero/comparativo", params={"filas": "mes", "columnas": "mes"}).status_code == 400


def test_tablas_paginadas_y_csv(con_datos):
    c = con_datos
    for nombre in tb.TABLAS:
        r = c.get(f"/api/tablero/tabla/{nombre}")
        assert r.status_code == 200, nombre
    p1 = c.get("/api/tablero/tabla/consolidado", params={"tam": 5, "pagina": 1}).json()
    p4 = c.get("/api/tablero/tabla/consolidado", params={"tam": 5, "pagina": 4}).json()
    assert p1["total"] == 16 and len(p1["filas"]) == 5 and len(p4["filas"]) == 1
    assert p1["columnas"] == COLUMNAS
    b = c.get("/api/tablero/tabla/consolidado", params={"buscar": "san jose"}).json()
    assert b["total"] == 8
    csv = c.get("/api/tablero/tabla/consolidado", params={"formato": "csv"})
    assert csv.headers["content-type"].startswith("text/csv")
    assert len(pd.read_csv(io.StringIO(csv.content.decode("utf-8-sig")), sep=";")) == 16
    cups = c.get("/api/tablero/tabla/cups").json()
    valores = [f["valor"] for f in cups["filas"]]
    assert valores == sorted(valores, reverse=True)
    assert c.get("/api/tablero/tabla/no_existe").status_code == 404


def test_informe_excel(con_datos):
    r = con_datos.get("/api/tablero/excel")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    assert {"Resumen", "Alertas", "Por mes", "Resumen sedes", "RIPS por EPS", "RIPS por CUPS",
            "CONSOLIDADO RIPS", "Frecuencias"} <= set(wb.sheetnames)
    filas = list(wb["CONSOLIDADO RIPS"].iter_rows(values_only=True))
    assert list(filas[3]) == COLUMNAS and len(filas) == 4 + 16


def test_alertas_y_analisis():
    filas = []
    for mes, valor in (("2026-01", 1000), ("2026-02", 1000), ("2026-04", 3000)):  # falta 2026-03
        filas.append(_fila(mes, "SEDE A", valor=valor, factura=f"A{mes}", fev=valor))
    for mes in ("2026-01", "2026-02", "2026-03", "2026-04"):
        filas.append(_fila(mes, "SEDE B", valor=500, factura=f"B{mes}", fev=900, validacion="DIFERENCIA"))
    df = pd.DataFrame(filas)
    df = tb_df(df)
    al = tb.alertas(df, pd.DataFrame(columns=tb.COLUMNAS_INDICADORES))
    cats = {(a["nivel"], a["categoria"], a["sede"], a["mes"]) for a in al}
    assert ("CRÍTICA", "Sin RIPS", "SEDE A", "2026-03") in cats
    assert any(c[1] == "Conciliación AF vs. servicios" and c[2] == "SEDE B" for c in cats)
    assert al[0]["nivel"] == "CRÍTICA"  # ordenadas por prioridad
    an = tb.analisis(df, pd.DataFrame(columns=tb.COLUMNAS_INDICADORES), al)
    assert any("RIPS faltantes" in r for r in an["recomendaciones"])


def tb_df(df: pd.DataFrame) -> pd.DataFrame:
    """Mismo tratamiento que tb.cargar() a un DataFrame en memoria."""
    import tempfile
    from pathlib import Path

    with tempfile.TemporaryDirectory() as d:
        for mes, g in df.groupby("MES"):
            g.to_csv(Path(d) / f"{mes}.csv", index=False)
        return tb.cargar(d)[0]


def test_indicadores_reemplazan_sede_mes(con_datos, tmp_path):
    ind = tb.cargar(tb.carpeta_acumulado())[1]
    assert len(ind) == 4
    # un segundo cargue del mismo lote no duplica
    zip_lote = fabrica.zip_varios_hospitales(tmp_path / "otra.zip")
    t = _esperar(con_datos, con_datos.post("/api/cargues", data={"ruta_local": str(zip_lote)}).json()["id"])
    assert t["estado"] == "terminado"
    assert len(tb.cargar(tb.carpeta_acumulado())[1]) == 4


def test_enviar_desde_tablero_pide_meses(con_datos):
    assert con_datos.post("/api/tablero/enviar", json={"meses": []}).status_code == 400
    assert con_datos.post("/api/tablero/enviar", json={"meses": ["1999-01"]}).status_code == 400


# ---------------------------------------------------------------------------
# Acceso en red
# ---------------------------------------------------------------------------

def test_equipo_remoto_necesita_clave(monkeypatch):
    monkeypatch.setenv("AUDITORIA_ESE_DIR", "/no/existe")
    remoto = TestClient(servidor.app, client=REMOTO)
    r = remoto.get("/tablero", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].startswith("/acceso?siguiente=/tablero")
    assert remoto.get("/api/estado").status_code == 401
    assert remoto.get("/acceso").status_code == 200
    assert remoto.post("/acceso", data={"clave": "MALA", "siguiente": "/"}).status_code == 401
    clave = red.clave_acceso()
    r = remoto.post("/acceso", data={"clave": clave.lower(), "siguiente": "/tablero"}, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/tablero"
    e = remoto.get("/api/estado").json()
    assert e["equipo_local"] is False and e["red"]["clave"] is None  # la clave sólo se muestra en el servidor
    assert remoto.post("/api/cargues/x/abrir-carpeta").status_code == 400
    # redirección sólo a rutas propias
    otro = TestClient(servidor.app, client=("192.168.1.21", 1))
    r = otro.post("/acceso", data={"clave": clave, "siguiente": "//malicioso.com"}, follow_redirects=False)
    assert r.headers["location"] == "/"


def test_bloqueo_por_intentos(monkeypatch):
    red._fallos.clear()
    c = TestClient(servidor.app, client=("192.168.1.99", 1))
    for _ in range(red.INTENTOS_MAXIMOS):
        assert c.post("/acceso", data={"clave": "X"}).status_code == 401
    assert c.post("/acceso", data={"clave": red.clave_acceso()}).status_code == 429
    red._fallos.clear()


def test_clave_fija_en_configuracion(monkeypatch):
    assert red.clave_acceso({"red": {"clave": "Hospital2026"}}) == "Hospital2026"
    assert red.token("hospital2026") == red.token("HOSPITAL2026")
    assert red.clave_acceso() == red.clave_acceso()  # generada una vez y guardada
    assert len(red.clave_acceso()) == 8


class _V5(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def do_GET(self):
        cuerpo = json.dumps({"ok": True, "ruta": self.path}).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(cuerpo)))
        self.end_headers()
        self.wfile.write(cuerpo)

    def do_POST(self):
        n = int(self.headers["Content-Length"])
        cuerpo = json.dumps({"recibido": n}).encode()
        self.send_response(201)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(cuerpo)))
        self.end_headers()
        self.wfile.write(cuerpo)


def test_puente_al_aplicativo_de_auditoria():
    srv = ThreadingHTTPServer(("127.0.0.1", 0), _V5)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        app = red.ProteccionRed(red.PuenteAuditoria(f"http://127.0.0.1:{srv.server_address[1]}"))
        remoto = TestClient(app, client=REMOTO)
        assert remoto.get("/api/health").status_code == 401
        remoto.post("/acceso", data={"clave": red.clave_acceso()})
        assert remoto.get("/api/informes/x?mes=2026-05").json() == {"ok": True, "ruta": "/api/informes/x?mes=2026-05"}
        r = remoto.post("/api/cargas/rips", files={"file": ("a.xlsx", b"x" * 5000)})
        assert r.status_code == 201 and r.json()["recibido"] > 5000
    finally:
        srv.shutdown()
    cerrado = TestClient(red.PuenteAuditoria("http://127.0.0.1:1"), client=LOCAL)
    assert cerrado.get("/").status_code == 503
