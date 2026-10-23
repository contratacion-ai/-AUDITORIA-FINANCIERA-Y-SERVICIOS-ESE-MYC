"""Plantilla de carga directa (Excel) generada desde el maestro de referencia.

Hojas:
* INSTRUCCIONES: pasos y reglas.
* CARGA: el usuario diligencia PERIODO, CODIGO CECO (lista desplegable), CUENTA, NOMBRE CUENTA, VALOR y
  DESCRIPCION. Las columnas «(AUTO)» se calculan con fórmulas (nombre, CECO de resolución, sede, estado y
  validación de la fila) y la herramienta las ignora al cargar.
* CATALOGO CECO: copia del maestro (código, nombres, llave, sede, tipo y estado).
* RESUMEN: filas listas, filas con error y totales por estado, antes de cargar.

No usa macros (VBA): funciona en Excel, LibreOffice y Google Sheets.
"""
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from . import config as C

FILAS = 3000
AZUL = PatternFill("solid", fgColor="1F4E78")
AZUL_EDITABLE = PatternFill("solid", fgColor="DDEBF7")
GRIS = PatternFill("solid", fgColor="EDEDED")
BLANCO = Font(bold=True, color="FFFFFF")

COLUMNAS = [
    # (encabezado, ancho, editable, nota)
    ("PERIODO", 11, True, "AAAA-MM, p. ej. 2026-01 (texto)"),
    ("CODIGO CECO", 14, True, "Código del CECO del estado financiero (lista)"),
    ("CUENTA", 14, True, "Cuenta contable, solo dígitos (clases 5, 6, 7)"),
    ("NOMBRE CUENTA", 36, True, "Nombre de la cuenta (opcional)"),
    ("VALOR", 18, True, "Valor del mes; negativo para ajustes crédito"),
    ("DESCRIPCION", 30, True, "Detalle opcional"),
    ("NOMBRE CECO (AUTO)", 42, False, ""),
    ("CECO RESOLUCIÓN (AUTO)", 14, False, ""),
    ("LLAVE CECO|SEDE (AUTO)", 16, False, ""),
    ("ESTADO (AUTO)", 13, False, ""),
    ("VALIDACIÓN (AUTO)", 44, False, ""),
]

INSTRUCCIONES = [
    "PLANTILLA DE CARGA DIRECTA – COSTOS POR CENTRO DE COSTOS (CECO)",
    "ESE Departamental de Primer Nivel Moreno y Clavijo",
    "",
    "1. Diligencie la hoja CARGA desde la fila 2. Una fila por CECO y cuenta contable del mes.",
    "2. Columnas azules (editables): PERIODO (AAAA-MM), CODIGO CECO (elíjalo de la lista), CUENTA, NOMBRE CUENTA, VALOR y DESCRIPCION.",
    "3. Las columnas grises (AUTO) se calculan solas: nombre del CECO, código de resolución, llave CECO|SEDE, estado y validación. No las edite.",
    "4. Revise la columna VALIDACIÓN: solo las filas en OK se cargarán. Los mensajes indican qué corregir.",
    "5. Estados: ACTIVO = se carga · VERIFICAR = validar con contabilidad/habilitación (no se carga) · NO APLICA = servicio no habilitado "
    "en primer nivel (no debería tener gasto) · EXCLUIR = partida que no es costo.",
    "6. Revise la hoja RESUMEN antes de cargar: total por estado y número de filas con error.",
    "7. Guarde el archivo (.xlsx) y procéselo: python -m homologacion_ceco procesar PLANTILLA.xlsx --salida resultado.xlsx",
    "   Puede incluirlo en un ZIP con otros archivos; la herramienta reconoce esta plantilla sin configuración adicional.",
    "8. Montos: escriba números (sin texto). Si pega desde otro sistema, use Pegado especial > Valores.",
    "9. Si un CECO no aparece en la lista, no lo invente: solicite su homologación (se agrega al archivo de referencia).",
    "",
    "Validaciones antes de cargar (columna VALIDACIÓN):",
    "• PERIODO con formato AAAA-MM · CECO existente en el catálogo · CUENTA solo dígitos y de clase 5, 6 o 7 · VALOR numérico y distinto de cero",
    "• Estado del CECO: solo ACTIVO se carga; los demás se reportan como retenidos en el resultado.",
]


def generar_plantilla(referencia, salida, filas=FILAS):
    wb = Workbook()
    ins = wb.active
    ins.title = "INSTRUCCIONES"
    for i, linea in enumerate(INSTRUCCIONES, start=1):
        c = ins.cell(row=i, column=1, value=linea)
        if i <= 2:
            c.font = Font(bold=True, size=13 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True)
    ins.column_dimensions["A"].width = 150

    carga = wb.create_sheet("CARGA")
    cat = wb.create_sheet("CATALOGO CECO")
    res = wb.create_sheet("RESUMEN")

    # ------------------------------------------------------------------ catálogo
    enc_cat = ["CODIGO CECO", "NOMBRE CECO", "CECO RESOLUCION", "NOMBRE CECO RESOLUCION", "LLAVE CECO|SEDE", "SEDE / IPS",
               "TIPO CECO", "ESTADO"]
    for j, e in enumerate(enc_cat, start=1):
        c = cat.cell(row=1, column=j, value=e)
        c.font, c.fill = BLANCO, AZUL
    m = referencia.maestro.sort_values("cod_ef")
    for i, r in enumerate(m.itertuples(index=False), start=2):
        cod = int(r.cod_ef) if r.cod_ef.isdigit() else r.cod_ef
        for j, v in enumerate([cod, r.nombre_ef, r.cod_res, r.nombre_res, r.llave, r.sede, r.tipo, r.estado], start=1):
            cat.cell(row=i, column=j, value=v)
    ultima = len(m) + 1
    for j, ancho in enumerate([13, 50, 12, 40, 14, 32, 30, 12], start=1):
        cat.column_dimensions[get_column_letter(j)].width = ancho
    cat.freeze_panes = "A2"
    cat.auto_filter.ref = f"A1:H{ultima}"
    wb.defined_names["CECO_CODIGOS"] = DefinedName("CECO_CODIGOS", attr_text=f"'CATALOGO CECO'!$A$2:$A${ultima}")
    rango = f"'CATALOGO CECO'!$A$2:$H${ultima}"

    # ------------------------------------------------------------------ hoja de carga
    for j, (e, ancho, editable, nota) in enumerate(COLUMNAS, start=1):
        c = carga.cell(row=1, column=j, value=e)
        c.font, c.fill = BLANCO, AZUL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        carga.column_dimensions[get_column_letter(j)].width = ancho
    carga.row_dimensions[1].height = 32
    for i in range(2, filas + 2):
        for j in range(1, 7):
            carga.cell(row=i, column=j).fill = AZUL_EDITABLE
        carga.cell(row=i, column=1).number_format = "@"
        carga.cell(row=i, column=3).number_format = "@"
        carga.cell(row=i, column=5).number_format = "#,##0.00"
        busca = f"IFERROR(VLOOKUP(--B{i},{rango},{{col}},FALSE),\"\")"
        carga.cell(row=i, column=7, value=f'=IF(B{i}="","",IFERROR(VLOOKUP(--B{i},{rango},2,FALSE),"CECO NO EXISTE"))')
        carga.cell(row=i, column=8, value=f'=IF(B{i}="","",{busca.format(col=3)})')
        carga.cell(row=i, column=9, value=f'=IF(B{i}="","",{busca.format(col=5)})')
        carga.cell(row=i, column=10, value=f'=IF(B{i}="","",{busca.format(col=8)})')
        carga.cell(row=i, column=11, value=(
            f'=IF(COUNTA(A{i}:F{i})=0,"",'
            f'IF(OR(LEN(A{i})<>7,MID(A{i},5,1)<>"-",ISERROR(--LEFT(A{i},4)),ISERROR(--RIGHT(A{i},2))),"ERROR: PERIODO debe ser AAAA-MM",'
            f'IF(B{i}="","ERROR: falta CODIGO CECO",'
            f'IF(G{i}="CECO NO EXISTE","ERROR: CECO no existe en el catálogo",'
            f'IF(C{i}="","ERROR: falta CUENTA",'
            f'IF(ISERROR(--C{i}),"ERROR: CUENTA debe tener solo dígitos",'
            f'IF(ISERROR(FIND(LEFT(C{i},1),"{"".join(C.sedes_y_puc()["clases_costo"])}")),"ERROR: CUENTA fuera de las clases de costo",'
            f'IF(NOT(ISNUMBER(E{i})),"ERROR: VALOR debe ser numérico",'
            f'IF(E{i}=0,"AVISO: valor cero (no se carga)",'
            f'IF(J{i}<>"ACTIVO","NO SE CARGA: CECO en estado "&J{i},"OK"))))))))))'))
        for j in range(7, 12):
            carga.cell(row=i, column=j).fill = GRIS
    carga.freeze_panes = "C2"
    carga.auto_filter.ref = f"A1:K{filas + 1}"

    ult = filas + 1
    dv_ceco = DataValidation(type="list", formula1="=CECO_CODIGOS", allow_blank=True, showErrorMessage=True,
                             errorTitle="CECO no válido", error="Elija un código de la lista (hoja CATALOGO CECO).")
    dv_periodo = DataValidation(type="textLength", operator="equal", formula1="7", allow_blank=True, showErrorMessage=True,
                                errorTitle="Periodo", error="Escriba el periodo como AAAA-MM, p. ej. 2026-01.")
    dv_valor = DataValidation(type="decimal", operator="between", formula1="-1E+15", formula2="1E+15", allow_blank=True,
                              showErrorMessage=True, errorTitle="Valor", error="El valor debe ser numérico.")
    dv_cuenta = DataValidation(type="custom", formula1="=ISNUMBER(--C2)", allow_blank=True, showErrorMessage=True,
                               errorTitle="Cuenta", error="La cuenta debe tener solo dígitos.")
    for dv, col in ((dv_periodo, "A"), (dv_ceco, "B"), (dv_cuenta, "C"), (dv_valor, "E")):
        dv.add(f"{col}2:{col}{ult}")
        carga.add_data_validation(dv)
    rojo = PatternFill("solid", fgColor="F8CBAD")
    amarillo = PatternFill("solid", fgColor="FFF2CC")
    verde = PatternFill("solid", fgColor="E2EFDA")
    carga.conditional_formatting.add(f"K2:K{ult}", FormulaRule(formula=['LEFT(K2,5)="ERROR"'], fill=rojo))
    carga.conditional_formatting.add(f"K2:K{ult}", FormulaRule(formula=['OR(LEFT(K2,2)="NO",LEFT(K2,5)="AVISO")'], fill=amarillo))
    carga.conditional_formatting.add(f"K2:K{ult}", CellIsRule(operator="equal", formula=['"OK"'], fill=verde))

    # ------------------------------------------------------------------ resumen previo a la carga
    res.column_dimensions["A"].width = 48
    res.column_dimensions["B"].width = 14
    res.column_dimensions["C"].width = 22
    res["A1"] = "RESUMEN DE VALIDACIÓN ANTES DE CARGAR"
    res["A1"].font = Font(bold=True, size=12)
    filas_res = [
        ("Filas diligenciadas", f'=COUNTIF(CARGA!K2:K{ult},"?*")', f"=SUM(CARGA!E2:E{ult})"),
        ("Filas listas para cargar (OK)", f'=COUNTIF(CARGA!K2:K{ult},"OK")', f'=SUMIFS(CARGA!E2:E{ult},CARGA!K2:K{ult},"OK")'),
        ("Filas con ERROR (corregir)", f'=COUNTIF(CARGA!K2:K{ult},"ERROR*")', f'=SUMIFS(CARGA!E2:E{ult},CARGA!K2:K{ult},"ERROR*")'),
        ("Filas que no se cargarán por estado del CECO", f'=COUNTIF(CARGA!K2:K{ult},"NO SE CARGA*")',
         f'=SUMIFS(CARGA!E2:E{ult},CARGA!K2:K{ult},"NO SE CARGA*")'),
        ("Filas con valor cero", f'=COUNTIF(CARGA!K2:K{ult},"AVISO*")', None),
    ]
    for j, e in enumerate(["CONCEPTO", "FILAS", "VALOR"], start=1):
        c = res.cell(row=3, column=j, value=e)
        c.font, c.fill = BLANCO, AZUL
    for i, (k, n, v) in enumerate(filas_res, start=4):
        res.cell(row=i, column=1, value=k)
        res.cell(row=i, column=2, value=n)
        if v:
            res.cell(row=i, column=3, value=v).number_format = "#,##0.00"
    base = 4 + len(filas_res) + 1
    for j, e in enumerate(["ESTADO DEL CECO", "FILAS", "VALOR"], start=1):
        c = res.cell(row=base, column=j, value=e)
        c.font, c.fill = BLANCO, AZUL
    for i, estado in enumerate(C.ESTADOS, start=base + 1):
        res.cell(row=i, column=1, value=estado)
        res.cell(row=i, column=2, value=f'=COUNTIF(CARGA!J2:J{ult},A{i})')
        res.cell(row=i, column=3, value=f'=SUMIFS(CARGA!E2:E{ult},CARGA!J2:J{ult},A{i})').number_format = "#,##0.00"
    wb.active = 1
    wb.save(salida)
    return salida
