import shutil

import openpyxl
import pytest

from homologacion_ceco import config as C
from homologacion_ceco.referencia import ErrorReferencia, cargar_referencia
from homologacion_ceco.reglas import ESTRUCTURA, MAESTRO, NOMBRE, Homologador


def test_referencia_conteos_y_llave(ref):
    m = ref.maestro
    assert len(m) == 549
    assert m["estado"].value_counts().to_dict() == {"ACTIVO": 399, "NO APLICA": 96, "VERIFICAR": 51, "EXCLUIR": 3}
    assert len(ref.resolucion) == 184 and len(ref.utilizados) == 46
    fila = ref.buscar("1020101")
    assert fila["llave"] == "B01|HSLA" and fila["nombre_res"] == "CONSULTA URGENCIAS" and fila["tipo"] == C.PRODUCTIVO
    # La llave replica la fórmula D&"|"&LEFT(I,FIND(" ",I&" ")-1): una sede sin espacio queda completa.
    assert ref.buscar("1100606")["llave"] == "M05|PSLP"
    assert set(m["tipo"]) == set(C.TIPOS)


def test_control_de_calidad_recalculado(ref):
    cc = ref.control_calidad().set_index("indicador")["valor"]
    activos = ref.maestro[ref.maestro["estado"] == "ACTIVO"]
    assert cc["TOTAL CECO E.F."] == 549
    assert cc["Códigos homologados no válidos en resolución"] == 0
    assert cc['ACTIVOS asignados a CECO fuera de "Utilizados"'] == int((~activos["en_utilizados"]).sum()) == 43
    sin_asignar = set(ref.utilizados["cod_res"]) - set(activos["cod_res"])
    assert cc['CECO de resolución "Utilizados" sin ningún CECO E.F. asignado'] == len(sin_asignar) == 10
    assert sum(cc[t] for t in C.TIPOS) == 399


def test_referencia_editada_se_recalcula(tmp_path, ref):
    """Si el usuario cambia la columna D (código homologado) el nombre y la llave se recalculan sin abrir Excel."""
    copia = tmp_path / "ref.xlsx"
    shutil.copy(ref.ruta, copia)
    wb = openpyxl.load_workbook(copia)
    ws = wb["C COSTOS ESTADO FINANCIERO"]
    assert ws["A4"].value == 1020201
    ws["D4"] = "A24"
    ws["J5"] = "PENDIENTE"
    ws["D6"] = "ZZZ"
    wb.save(copia)
    r = cargar_referencia(copia)
    assert r.buscar("1020201")["llave"] == "A24|HSLA"
    assert r.buscar("1020201")["nombre_res"] == "GINECOLOGÍA Y OBSTETRICIA"
    tipos = set(r.hallazgos["tipo"])
    assert {"ESTADO NO VALIDO", "CODIGO NO VALIDO"} <= tipos
    h = Homologador(r)
    assert h.homologar("1020202").estado == C.VERIFICAR      # estado no válido → no se carga
    assert h.homologar("1020203").estado == C.VERIFICAR      # código no válido → no se carga


def test_referencia_invalida(tmp_path):
    with pytest.raises(ErrorReferencia):
        cargar_referencia(tmp_path / "no_existe.xlsx")
    malo = tmp_path / "malo.xlsx"
    malo.write_bytes(b"no es excel")
    with pytest.raises(ErrorReferencia):
        cargar_referencia(malo)


def test_maestro_manda_y_respeta_estados(ref):
    h = Homologador(ref, activar_nuevos=True)
    assert h.homologar("1020201").origen == MAESTRO and h.homologar("1020201").cargable
    for cod, estado in (("1020304", C.VERIFICAR), ("1020302", C.NO_APLICA), ("3012199", C.EXCLUIR)):
        x = h.homologar(cod)
        assert x.estado == estado and not x.cargable


def test_reglas_reproducen_el_maestro(ref):
    """Control de las reglas: aplicadas a los 549 CECO (ignorando el maestro) proponen el mismo código."""
    d = Homologador(ref).comparar_con_maestro()
    assert d["coincide"].mean() == 1.0
    assert (d["sede_maestro"] == d["sede_regla"]).mean() > 0.97
    assert set(d["origen_regla"]) == {ESTRUCTURA, NOMBRE}


@pytest.mark.parametrize("nombre, cod_res", [
    ("HSXX-CONSULTA EXTERNA", "A10"), ("LABORATORIO CLINICO", "L07"), ("IMAGENOLOGIA", "L02"), ("FARMACIA", "M05"),
    ("MANTENIMIENTO PLANTA FISICA", "P02"), ("SERVICIOS PUBLICOS", "R02"), ("PROMOCION Y MANTENIMIENTO", "N03"),
    ("TRANSPORTE PACIENTES CONDUCTORES", "N08"), ("ESE01-TAME-MANTENIMIENTO PLANTA FISICA", "R03"),
])
def test_reglas_de_nombre(ref, nombre, cod_res):
    p = Homologador(ref).proponer("9999999", nombre)
    assert p.cod_res == cod_res and p.origen == NOMBRE


def test_servicios_publicos_lleva_nota_de_m2(ref):
    assert "m²" in Homologador(ref).proponer("9999999", "HSLA-SERVICIOS PUBLICOS").nota


@pytest.mark.parametrize("nombre", ["HSLA-RECIEN NACIDOS", "HSLA-QUIROFANOS", "HSLA-CONSULTA ESPECIALIZADA"])
def test_servicios_que_requieren_validacion_no_se_resuelven_solos(ref, nombre):
    p = Homologador(ref, activar_nuevos=True).proponer("1990000", nombre)
    assert p.estado == C.VERIFICAR and not p.cargable


def test_ceco_nuevo_por_estructura(ref):
    # proponer() ignora el maestro: grupo 1, sede 04 (HSAT), servicio 0201 → A10 en todas las sedes.
    h = Homologador(ref)
    p = h.proponer("1040201", "")
    assert p.cod_res == "A10" and p.origen == ESTRUCTURA and p.sede == "HSAT - Tame" and p.llave == "A10|HSAT"
    assert p.estado == C.VERIFICAR                                  # por defecto un CECO nuevo no se carga
    assert Homologador(ref, activar_nuevos=True).proponer("1040201", "").estado == C.ACTIVO
    # Sede desconocida (segmento 99): no se carga aunque se activen los nuevos.
    sin_sede = Homologador(ref, activar_nuevos=True).proponer("1990201", "")
    assert sin_sede.estado == C.VERIFICAR and sin_sede.llave == "" and "sede" in sin_sede.nota
    # La estructura aprendida conserva NO APLICA (servicio no habilitado en primer nivel).
    assert Homologador(ref, activar_nuevos=True).proponer("1990302", "").estado == C.NO_APLICA


def test_ceco_sin_regla(ref):
    p = Homologador(ref, activar_nuevos=True).proponer("8880000", "PARTIDA DESCONOCIDA")
    assert p.estado == C.SIN_HOMOLOGAR and not p.cod_res


def test_sede_por_subsede_y_archivo(ref):
    h = Homologador(ref)
    assert h.proponer("1091606", "FARMACIA PRINCIPAL FILIPINAS").sede == "HSAT - Filipinas"
    assert h.proponer("1990201", "CONSULTA EXTERNA", sede_fuente="Fortul").sede == "HSFF - Fortul"
