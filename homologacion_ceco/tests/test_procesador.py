"""Proceso de punta a punta con el lote de ejemplo (valores calculados a mano en docs/HOMOLOGACION_CECO.md)."""
import io
import zipfile

import openpyxl
import pytest

from homologacion_ceco import config as C
from homologacion_ceco.cli import main
from homologacion_ceco.columnas import ErrorMapeo
from homologacion_ceco.ejemplo import ef_enero, lote_ejemplo
from homologacion_ceco.exportar import exportar
from homologacion_ceco.lectores import Lector
from homologacion_ceco.plantilla import generar_plantilla
from homologacion_ceco.procesador import CARGADO, EXCLUIDO, RECHAZADO, RETENIDO, Opciones, procesar

from .conftest import pesos, xlsx


def _por(res, destino):
    return res.registros[res.registros["destino"] == destino]


def _total(df):
    return sum(int(v) for v in df["valor_cent"] if v == v and v is not None)


def test_totales_del_lote(resultado_lote):
    r = resultado_lote
    car = _por(r, CARGADO)
    por_periodo = {p: pesos(_total(g)) for p, g in car.groupby("periodo")}
    assert por_periodo == {"2026-01": 15300000.50, "2026-02": 6600000.00, "2026-03": 6010000.00}
    assert pesos(_total(_por(r, RETENIDO))) == 2550000.00
    assert pesos(_total(_por(r, RECHAZADO))) == 600000.00
    assert pesos(_total(r.consolidado)) == pesos(_total(car)) == 27910000.50
    assert r.conciliacion["concepto"].iloc[-1].endswith("CUADRA")
    assert r.codigo_salida == 2


def test_cada_fila_tiene_un_destino_y_motivo(resultado_lote):
    reg = resultado_lote.registros
    assert reg["destino"].isin([CARGADO, RETENIDO, EXCLUIDO, RECHAZADO]).all()
    no_cargados = reg[reg["destino"] != CARGADO]
    assert (no_cargados["motivo"].fillna("") != "").all()   # ningún registro se pierde en silencio


def test_estados_no_activos_no_se_cargan(resultado_lote):
    reg = resultado_lote.registros
    assert set(reg.loc[reg["destino"] == CARGADO, "estado"]) == {C.ACTIVO}
    ret = reg[reg["destino"] == RETENIDO].set_index("cod_ceco")
    assert ret.loc["1020304", "estado"] == C.VERIFICAR
    assert ret.loc["1020302", "estado"] == C.NO_APLICA
    assert ret.loc["3012199", "estado"] == C.EXCLUIR
    assert ret.loc["1100201", "estado"] == C.VERIFICAR and "CECO nuevo" in ret.loc["1100201", "motivo"]
    assert ret.loc["1100401", "estado"] == C.VERIFICAR                # quirófanos: regla de validación
    alertas = set(resultado_lote.alertas["tipo"])
    assert "NO APLICA CON GASTO" in alertas


def test_limpieza_y_reglas(resultado_lote):
    reg = resultado_lote.registros
    motivos = " | ".join(reg["motivo"].dropna())
    for esperado in ("es subtotal de 1 cuenta(s)", "Cuenta de clase 1", "Fila de total", "Duplicado exacto de",
                     "Sin código de centro de costos", "'ABC' no es un número", "'31/13/2026' no es un periodo"):
        assert esperado in motivos
    # Débito − crédito con formato colombiano y negativos entre paréntesis.
    lab_feb = reg[(reg["cod_ceco"] == "1020501") & (reg["periodo"] == "2026-02") & (reg["destino"] == CARGADO)]
    assert pesos(_total(lab_feb)) == 2300000.00
    energia_feb = reg[(reg["cod_ceco"] == "4022208") & (reg["periodo"] == "2026-02")]
    assert pesos(_total(energia_feb)) == -50000.00


def test_formatos_matriz_y_por_ceco(resultado_lote):
    m = resultado_lote.mapeo
    assert {"LARGO", "MATRIZ", "POR_CECO"} <= set(m["formato"])
    car = _por(resultado_lote, CARGADO)
    marzo = car[car["periodo"] == "2026-03"].groupby("llave")["valor_cent"].sum().map(pesos).to_dict()
    assert marzo == {"A10|HSLA": 3200000.0, "A10|HSAT": 1300000.0, "R02|HSLA": 820000.0, "L02|HSAT": 690000.0}


def test_zip_errores_y_duplicados(resultado_lote):
    a = resultado_lote.archivos.set_index("origen")
    assert a.loc["lote.zip!/corrupto.xlsx", "estado"] == "ERROR"
    assert a.loc["lote.zip!/LEAME.pdf", "estado"] == "OMITIDO"
    assert a.loc["lote.zip!/copia/EF_ENERO_2026.xlsx", "estado"] == "DUPLICADO"
    assert not any("__MACOSX" in o for o in a.index)
    assert "lote.zip!/soportes/soportes_marzo.zip!/matriz_marzo_2026.xlsx" in a.index   # ZIP anidado


def test_trazabilidad_hasta_el_origen(resultado_lote):
    r = resultado_lote
    car = _por(r, CARGADO)
    # Cada línea consolidada = suma de sus registros de detalle, y cada registro apunta a archivo/hoja/fila.
    for c in r.consolidado.itertuples(index=False):
        det = car[car["id_consolidado"] == c.id_consolidado]
        assert int(det["valor_cent"].sum()) == c.valor_cent and len(det) == c.registros
        assert det["origen"].str.startswith("lote.zip").all() and det["fila"].gt(0).all()
    tz = r.trazabilidad
    assert {"MAPEO", "LIMPIEZA", "VALIDACION", "REGLA", "HOMOLOGACION", "CONSOLIDACION"} <= set(tz["etapa"])
    for rid in car["id_registro"]:
        assert (tz.loc[tz["id_registro"] == rid, "etapa"] == "HOMOLOGACION").sum() == 1
    # La fila de Excel registrada es la real: el ACTIVO 1020501 de enero está en la fila 10 del balance.
    lab = car[(car["cod_ceco"] == "1020501") & (car["periodo"] == "2026-01")].iloc[0]
    assert lab["fila"] == 10 and lab["hoja"] == "Balance por CECO"


def test_activar_nuevos(ref):
    res = procesar([], ref, Opciones(activar_nuevos=True), contenidos=[("lote.zip", lote_ejemplo())])
    reg = res.registros.set_index("cod_ceco")
    assert reg.loc["1100401", "destino"] == RETENIDO          # quirófanos siempre VERIFICAR
    assert reg.loc["1100201", "destino"] == RETENIDO          # sin sede determinable → VERIFICAR


def test_mapeo_del_usuario_y_errores_de_estructura(ref):
    contenido = xlsx([["CC X", "AUX", "MONTO MES DEL CENTRO"], [1020201, "51010101", 1000]])
    res = procesar([], ref, Opciones(periodo="2026-04", mapeo={
        "archivos": {"*.xlsx": {"cod_ceco": "#A", "cuenta": "AUX", "valor": "MONTO MES DEL CENTRO"}}}),
        contenidos=[("raro.xlsx", contenido)])
    assert pesos(_total(_por(res, CARGADO))) == 1000.0
    assert set(res.mapeo["metodo"]) == {"MAPEO DEL USUARIO", "TOMADO DEL PARÁMETRO --PERIODO"}
    # Sin mapeo, faltan columnas: la hoja queda en ERROR con el detalle y el proceso sigue.
    res2 = procesar([], ref, Opciones(periodo="2026-04"), contenidos=[("raro.xlsx", contenido), ("ef.xlsx", ef_enero())])
    a = res2.archivos.set_index("origen")
    assert "Faltan columnas" in a.loc["raro.xlsx", "motivo"] or a.loc["raro.xlsx", "estado"] in ("OMITIDO", "ERROR")
    assert len(_por(res2, CARGADO)) == 4
    # Columna fija inexistente: esa hoja queda en ERROR con el detalle (el lote continúa).
    res3 = procesar([], ref, Opciones(mapeo={"archivos": {"ef*": {"cuenta": "NO EXISTE"}}}), contenidos=[("ef.xlsx", ef_enero())])
    ev = res3.archivos[res3.archivos["origen"] == "ef.xlsx"].iloc[-1]
    assert ev["estado"] == "ERROR" and "NO EXISTE" in ev["motivo"]
    # Un campo desconocido en el archivo de mapeo es un error de configuración: se detiene.
    with pytest.raises(ErrorMapeo):
        procesar([], ref, Opciones(mapeo={"sinonimos": {"centro": ["X"]}}), contenidos=[("ef.xlsx", ef_enero())])


def test_usar_movimiento(ref):
    contenido = xlsx([["CECO", "CUENTA", "DEBITO", "CREDITO", "SALDO"], [1020201, "51010101", 300, 100, 5000]])
    normal = procesar([], ref, Opciones(periodo="2026-01"), contenidos=[("a.xlsx", contenido)])
    mov = procesar([], ref, Opciones(periodo="2026-01", usar_movimiento=True), contenidos=[("a.xlsx", contenido)])
    assert pesos(_total(_por(normal, CARGADO))) == 5000.0
    assert pesos(_total(_por(mov, CARGADO))) == 200.0


def test_sin_periodo_se_rechaza(ref):
    res = procesar([], ref, Opciones(), contenidos=[("sin.xlsx", xlsx([["CECO", "CUENTA", "VALOR"], [1020201, "5101", 10]]))])
    assert res.registros.iloc[0]["destino"] == RECHAZADO and "Sin periodo" in res.registros.iloc[0]["motivo"]


def test_limites_zip(ref):
    grande = io.BytesIO()
    with zipfile.ZipFile(grande, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("ef.xlsx", ef_enero())
    tablas, eventos = Lector({"max_total_zip_mb": 0}).leer_bytes("g.zip", grande.getvalue())
    assert not tablas and eventos[0].estado == "ERROR" and "supera" in eventos[0].motivo
    tablas, eventos = Lector().leer_bytes("roto.zip", b"PK\x03\x04 roto")
    assert eventos[0].estado == "ERROR"


def test_csv_cp1252_y_tabulado(ref):
    contenido = "CECO\tCUENTA\tVALOR\tFECHA\n1020201\t51010101\t1.000,50\t2026-07-01\n".encode("cp1252")
    res = procesar([], ref, Opciones(), contenidos=[("t.txt", contenido)])
    assert pesos(_total(_por(res, CARGADO))) == 1000.50


def test_exportar_libro(tmp_path, resultado_lote):
    salida, extras = exportar(resultado_lote, tmp_path / "res.xlsx")
    assert not extras
    wb = openpyxl.load_workbook(salida)
    for hoja in ("RESUMEN", "EF HOMOLOGADO", "EF POR CECO", "CARGA HERRAMIENTA", "DETALLE CARGADO", "RETENIDOS",
                 "RECHAZOS Y EXCLUSIONES", "ALERTAS", "CECO NUEVOS", "MAPEO COLUMNAS", "ARCHIVOS", "TRAZABILIDAD",
                 "CONTROL REFERENCIA", "SUPUESTOS"):
        assert hoja in wb.sheetnames
    carga = list(wb["CARGA HERRAMIENTA"].iter_rows(values_only=True))
    enc = carga[0]
    i_valor = enc.index("VALOR")
    assert round(sum(f[i_valor] for f in carga[1:]), 2) == 27910000.50
    assert {"CUENTA", "CENTRO DE COSTOS", "SEDE", "VALOR", "DESCRIPCION"} <= set(enc)
    # Un texto de origen que empieza por "=" se guarda como texto, nunca como fórmula.
    textos = [c for fila in wb["DETALLE CARGADO"].iter_rows() for c in fila if c.value == "=1+1"]
    assert textos and all(c.data_type == "s" for c in textos)
    ef = [f for f in wb["EF HOMOLOGADO"].iter_rows(values_only=True) if f[1] == "TOTAL COSTOS Y GASTOS HOMOLOGADOS (CARGADOS)"]
    assert [round(f[-1] if f[-1] is not None else [v for v in f if v is not None][-1], 2) for f in ef] == [15300000.5, 6600000.0, 6010000.0]


def test_plantilla_se_carga_directamente(tmp_path, ref):
    ruta = generar_plantilla(ref, tmp_path / "plantilla.xlsx", filas=50)
    wb = openpyxl.load_workbook(ruta)
    ws = wb["CARGA"]
    assert [c.value for c in ws[1]][:6] == ["PERIODO", "CODIGO CECO", "CUENTA", "NOMBRE CUENTA", "VALOR", "DESCRIPCION"]
    assert len(ws.data_validations.dataValidation) == 4
    assert ws["K2"].value.startswith("=IF(") and wb["CATALOGO CECO"].max_row == 550
    for fila in (["2026-05", 1020201, "51010101", "SUELDO", 1000, "x"], ["2026-05", 1020304, "51010101", "SUELDO", 50, ""],
                 ["2026-05", 1040201, "51110601", "ENERGIA", 25.5, ""]):
        n = ws.max_row + 1 if ws["A2"].value else 2
        n = next(i for i in range(2, 60) if ws.cell(row=i, column=1).value is None)
        for j, v in enumerate(fila, start=1):
            ws.cell(row=n, column=j, value=v)
    wb.save(ruta)
    res = procesar([ruta], ref)
    assert pesos(_total(_por(res, CARGADO))) == 1025.50
    assert pesos(_total(_por(res, RETENIDO))) == 50.0
    a = res.archivos
    assert set(a.loc[a["hoja"] == "CARGA", "estado"]) == {"PROCESADO"}
    assert "ERROR" not in set(a["estado"])                     # instrucciones, catálogo y resumen se omiten sin error
    assert not res.mapeo["columna_origen"].str.contains("AUTO").any()


def test_cli(tmp_path, capsys):
    lote = tmp_path / "lote.zip"
    assert main(["ejemplo", "--salida", str(lote)]) == 0
    assert main(["procesar", str(lote), "--salida", str(tmp_path / "r.xlsx")]) == 2
    assert "CUADRA" in capsys.readouterr().out
    assert main(["plantilla", "--salida", str(tmp_path / "p.xlsx"), "--filas", "10"]) == 0
    assert main(["validar-referencia"]) == 0
    assert main(["--referencia", str(tmp_path / "no.xlsx"), "validar-referencia"]) == 1
    limpio = tmp_path / "ok.xlsx"
    limpio.write_bytes(xlsx([["CECO", "CUENTA", "VALOR", "PERIODO"], [1020201, "51010101", 10, "2026-01"]]))
    assert main(["procesar", str(limpio), "--salida", str(tmp_path / "ok_r.xlsx")]) == 0
