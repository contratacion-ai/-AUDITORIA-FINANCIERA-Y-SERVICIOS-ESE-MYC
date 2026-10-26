"""Reporte CTRP del software contable (estado de resultados por tercero y centro de costo), con datos sintéticos."""
import openpyxl
import pytest

from homologacion_ceco.ejemplo import ctrp_estado_resultados
from homologacion_ceco.exportar import exportar
from homologacion_ceco.normalizar import reparar_mojibake
from homologacion_ceco.procesador import CARGADO, RECHAZADO, Opciones, procesar

from .conftest import pesos


@pytest.fixture(scope="module")
def res(ref):
    return procesar([], ref, Opciones(), contenidos=[("CTRPEstadoResultadosTerCen.xlsx", ctrp_estado_resultados())])


def test_cuenta_y_periodo_del_reporte(res):
    reg = res.registros
    assert set(reg["periodo"]) == {"2026-07"}                                 # del pie del reporte
    car = reg[reg["destino"] == CARGADO]
    assert set(car["cuenta"]) == {"4312960101", "5101010101", "6310500201"}   # vía OidCuentaProceso
    assert car["fila"].tolist() == [8, 9, 16, 17, 27]                         # filas reales del reporte
    assert res.mapeo["metodo"].eq("FORMATO RECONOCIDO").all()


def test_terceros_distintos_no_son_duplicados_y_no_se_copian(res):
    sueldos = res.registros[res.registros["cuenta"] == "5101010101"]
    assert len(sueldos) == 2 and set(sueldos["destino"]) == {CARGADO}
    assert pesos(int(sueldos["valor_cent"].sum())) == 600000.0
    texto_salida = res.registros.astype(str).to_string()
    assert "TERCERO FICTICIO" not in texto_salida and "900000021" not in texto_salida
    assert set(sueldos["documento"]) == {"OidTercero 21", "OidTercero 22"}


def test_cuenta_sin_detalle_por_ceco_queda_rechazada_y_cuadra(res):
    rech = res.registros[res.registros["destino"] == RECHAZADO]
    assert rech["cuenta"].tolist() == ["6310190101"] and pesos(int(rech["valor_cent"].sum())) == 250000.0
    assert "sin detalle por centro de costo" in rech.iloc[0]["motivo"]
    assert "CUENTA SIN DETALLE POR CECO" in set(res.alertas["tipo"])
    co = res.control_origen.set_index("clase")
    assert set(co["estado"]) == {"CUADRA"}
    assert co.loc["4", "reporte"] == -1000000.0 and co.loc["6", "leido"] == 650000.0


def test_estado_de_resultados_e_ingresos(tmp_path, res):
    salida, _ = exportar(res, tmp_path / "r.xlsx")
    wb = openpyxl.load_workbook(salida)
    filas = {f[1].strip(): f for f in wb["EF HOMOLOGADO"].iter_rows(values_only=True) if isinstance(f[1], str)}
    total = lambda nombre: [v for v in filas[nombre] if isinstance(v, (int, float))][-1]  # noqa: E731
    assert total("TOTAL INGRESOS (CARGADOS)") == 1000000                     # ingresos en positivo
    assert total("TOTAL GASTOS Y COSTOS HOMOLOGADOS (CARGADOS)") == 1000000
    assert total("RESULTADO DEL PERIODO (INGRESOS − GASTOS Y COSTOS)") == 0
    assert total("SIN CENTRO DE COSTO O CON ERROR (RECHAZADO)") == 250000
    assert total("TOTAL GASTOS Y COSTOS LEÍDOS (CARGADO + NO CARGADO)") == 1250000   # = clases 5 + 6 del reporte
    assert "SERVICIOS DE SALUD PRESTADOS CON EL MECANISMO DE PAGO POR CAPITACIÓN" in filas    # tildes reparadas
    carga = list(wb["CARGA HERRAMIENTA"].iter_rows(min_row=2, values_only=True))
    assert {f[1][:1] for f in carga} == {"5", "6"}                           # la herramienta ABC solo recibe costos


def test_usar_movimiento_en_ctrp(ref):
    r = procesar([], ref, Opciones(usar_movimiento=True), contenidos=[("c.xlsx", ctrp_estado_resultados())])
    assert set(r.mapeo.loc[r.mapeo["campo"] == "debito", "columna_origen"]) == {"ValorMovimientoDebito"}


@pytest.mark.parametrize("entrada, esperado", [
    ("CAPITACIÃ“N", "CAPITACIÓN"), ("MEDELLÃ\x8dN", "MEDELLÍN"), ("COMPAÃ‘IA", "COMPAÑIA"), ("Â¿", "¿"),
    ("HOSPITALIZACIÓN", "HOSPITALIZACIÓN"), ("Ã solo", "Ã solo"),
])
def test_reparar_mojibake(entrada, esperado):
    assert reparar_mojibake(entrada) == esperado
