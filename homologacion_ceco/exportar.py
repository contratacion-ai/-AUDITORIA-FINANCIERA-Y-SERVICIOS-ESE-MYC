"""Libro de salida único: estado financiero homologado, carga para la herramienta de costos y trazabilidad."""
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import __version__
from . import config as C
from . import estado_financiero as EF
from .procesador import CARGADO, EXCLUIDO, RECHAZADO, RETENIDO

MAX_FILAS_EXCEL = 1_048_000
FMT_PESOS = '#,##0.00;[Red]-#,##0.00'
AZUL = PatternFill("solid", fgColor="1F4E78")
AZUL_CLARO = PatternFill("solid", fgColor="DDEBF7")
GRIS = PatternFill("solid", fgColor="EDEDED")
AMARILLO = PatternFill("solid", fgColor="FFF2CC")
ROJO = PatternFill("solid", fgColor="F8CBAD")
VERDE = PatternFill("solid", fgColor="E2EFDA")
BLANCO_NEGRITA = Font(bold=True, color="FFFFFF")
NEGRITA = Font(bold=True)
FINO = Border(bottom=Side(style="thin", color="BFBFBF"))

COLOR_ESTADO = {C.ACTIVO: VERDE, C.VERIFICAR: AMARILLO, C.NO_APLICA: GRIS, C.EXCLUIR: ROJO, C.SIN_HOMOLOGAR: ROJO}

SUPUESTOS = [
    "El archivo de referencia (maestro) es la fuente de verdad: si un CECO está en el maestro se usan su código de "
    "resolución, sede, tipo y estado; las reglas solo proponen homologación para CECO nuevos.",
    "Solo se cargan CECO en estado ACTIVO. VERIFICAR, NO APLICA, EXCLUIR y SIN HOMOLOGAR quedan en la hoja RETENIDOS "
    "con su motivo y se muestran aparte en el estado financiero.",
    "Los CECO nuevos (no están en el maestro) quedan en VERIFICAR con una propuesta (hoja CECO NUEVOS) salvo que se use "
    "--activar-nuevos; los que una regla marca para validación (recién nacidos, quirófanos, consulta especializada…) "
    "siempre quedan en VERIFICAR.",
    "La llave CECO|SEDE usa el primer token de la sede (misma fórmula de la columna G del maestro): HSAT - Tame y HSAT - "
    "Filipinas comparten la IPS HSAT.",
    "Se leen las clases del estado de resultados (4 ingresos, 5 gastos, 6 y 7 costos); la hoja CARGA HERRAMIENTA solo "
    "lleva gastos y costos (5, 6, 7). Los ingresos se muestran en positivo y se calcula el resultado del periodo. Las "
    "cuentas que son prefijo de otras del mismo CECO y hoja se tratan como subtotales y no se suman, para no duplicar.",
    "Reporte CTRP del software contable (estado de resultados por tercero y centro): la cuenta de cada fila de detalle "
    "se toma de OidCuentaProceso; el periodo, del pie del reporte; el tercero solo por su identificador interno "
    "(OidTercero) para no confundir filas iguales de distintos terceros. Una cuenta sin detalle por CECO queda "
    "rechazada (SIN CENTRO DE COSTO) para que lo leído cuadre con el total del reporte.",
    "Si el archivo trae VALOR (o SALDO / MOVIMIENTO) se usa ese valor; si solo trae DÉBITO y CRÉDITO, valor = débito − crédito.",
    "Montos: se aceptan formatos 1.234.567,89 · 1,234,567.89 · $ (1.000) · 1.234 (miles). Un único separador seguido de "
    "exactamente 3 dígitos se interpreta como separador de miles (costumbre colombiana).",
    "Periodo: columna PERIODO/FECHA/MES del archivo; si no existe, el periodo del nombre de la hoja o del archivo "
    "(p. ej. EF_2026-01.xlsx, costos enero 2026.csv); si tampoco, el parámetro --periodo.",
    "Duplicados: se excluye una fila idéntica a otra (periodo, CECO, cuenta, valor, descripción y comprobante) aunque esté "
    "en otro archivo; un archivo con contenido idéntico a otro ya leído no se procesa dos veces.",
    "No se conservan datos de terceros (NIT/cédula, nombres): el estado financiero por CECO no los necesita.",
]


def _seguro(v):
    """Valor apto para una celda: NaN → vacío, escalares de numpy → Python."""
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    return v.item() if hasattr(v, "item") else v


def _celda(ws, fila, col, valor, **estilo):
    c = ws.cell(row=fila, column=col, value=_seguro(valor))
    if isinstance(c.value, str) and c.data_type == "f":
        c.data_type = "s"  # texto del archivo de origen que empieza por "=": nunca se escribe como fórmula
    for k, v in estilo.items():
        setattr(c, k, v)
    return c


def _tabla(ws, df, columnas, fila=1, titulo=None, montos=(), anchos=None, color_por=None):
    """Escribe un DataFrame con encabezado, filtro y formato. columnas: [(campo, encabezado)]."""
    if titulo:
        _celda(ws, fila, 1, titulo, font=Font(bold=True, size=12))
        fila += 1
    for j, (_, enc) in enumerate(columnas, start=1):
        _celda(ws, fila, j, enc, font=BLANCO_NEGRITA, fill=AZUL, alignment=Alignment(wrap_text=True, vertical="center"))
    inicio = fila
    total = len(df)
    datos = df.head(MAX_FILAS_EXCEL - fila) if total > MAX_FILAS_EXCEL - fila else df
    campos = [c for c, _ in columnas]
    idx_color = campos.index(color_por) if color_por in campos else None
    for valores in datos[campos].itertuples(index=False, name=None):
        fila += 1
        for j, v in enumerate(valores, start=1):
            c = _celda(ws, fila, j, v)
            if campos[j - 1] in montos:
                c.number_format = FMT_PESOS
        if idx_color is not None:
            relleno = COLOR_ESTADO.get(valores[idx_color])
            if relleno:
                ws.cell(row=fila, column=idx_color + 1).fill = relleno
    ws.freeze_panes = ws.cell(row=inicio + 1, column=1)
    if len(datos):
        ws.auto_filter.ref = f"A{inicio}:{get_column_letter(len(columnas))}{fila}"
    for j, (campo, enc) in enumerate(columnas, start=1):
        ancho = (anchos or {}).get(campo) or min(60, max(10, len(str(enc)) + 2,
                                                          *(len(str(v)) + 2 for v in datos[campo].head(200).tolist() if v is not None)))
        ws.column_dimensions[get_column_letter(j)].width = ancho
    return fila, total > len(datos)


def _pesos_df(df, col="valor_cent", nuevo="valor"):
    df = df.copy()
    df[nuevo] = df[col].map(lambda v: None if v is None or pd.isna(v) else int(v) / 100)
    return df


def _ef(ws, tablas, titulo, nombres_sede, resultado, negritas=(0, 1, -1)):
    fila = 1
    _celda(ws, fila, 1, "ESE Departamental de Primer Nivel Moreno y Clavijo", font=Font(bold=True, size=13))
    fila += 1
    _celda(ws, fila, 1, titulo, font=Font(bold=True, size=12))
    fila += 1
    _celda(ws, fila, 1, f"Homologado con {resultado.referencia.ruta.name} · generado {datetime.now():%Y-%m-%d %H:%M}",
           font=Font(italic=True, color="595959"))
    fila += 2
    max_cols = 3
    for periodo, df in tablas.items():
        sedes = [c for c in df.columns if c not in ("nivel", "cuenta", "nombre")]
        max_cols = max(max_cols, 2 + len(sedes))
        _celda(ws, fila, 1, f"PERIODO {periodo}", font=BLANCO_NEGRITA, fill=AZUL)
        for j in range(2, 3 + len(sedes)):
            ws.cell(row=fila, column=j).fill = AZUL
        fila += 1
        encabezados = ["CÓDIGO", "NOMBRE"] + [s if s == EF.TOTAL else f"{s}\n{nombres_sede.get(s, '')}" for s in sedes]
        for j, enc in enumerate(encabezados, start=1):
            _celda(ws, fila, j, enc, font=NEGRITA, fill=AZUL_CLARO, alignment=Alignment(wrap_text=True, vertical="center",
                                                                                          horizontal="center"))
        ws.row_dimensions[fila].height = 45
        for valores in df.itertuples(index=False, name=None):
            fila += 1
            r = dict(zip(df.columns, valores))
            nivel = int(r["nivel"])
            sangria = "   " * max(0, nivel - 1) if nivel > 0 else ""
            _celda(ws, fila, 1, r["cuenta"])
            _celda(ws, fila, 2, sangria + str(r["nombre"] or ""))
            for j, s in enumerate(sedes, start=3):
                v = r[s]
                if not (v is None or pd.isna(v)):
                    _celda(ws, fila, j, v, number_format=FMT_PESOS)
            for j in range(1, 3 + len(sedes)):
                c = ws.cell(row=fila, column=j)
                c.border = FINO
                if nivel in negritas:
                    c.font = NEGRITA
                if nivel == 0:
                    c.fill = AZUL_CLARO
                elif nivel == -2:
                    c.fill = COLOR_ESTADO.get(r["nombre"], GRIS)
        fila += 3
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 55
    for j in range(3, max_cols + 1):
        ws.column_dimensions[get_column_letter(j)].width = 18
    ws.freeze_panes = "C6"


COLS_REGISTRO = [
    ("id_registro", "ID REGISTRO"), ("origen", "ARCHIVO"), ("hoja", "HOJA"), ("fila", "FILA"), ("periodo", "PERIODO"),
    ("cod_ceco", "COD CECO E.F."), ("nombre_ceco_maestro", "NOMBRE CECO E.F. (MAESTRO)"),
    ("nombre_ceco_fuente", "NOMBRE CECO (ARCHIVO)"), ("cuenta", "CUENTA"), ("nombre_cuenta", "NOMBRE CUENTA"),
    ("valor", "VALOR"), ("cod_res", "COD CECO RES"), ("nombre_res", "NOMBRE CECO RES"), ("sede", "SEDE / IPS"),
    ("llave", "LLAVE CECO|SEDE"), ("tipo", "TIPO CECO"), ("estado", "ESTADO"), ("confianza", "CONFIANZA"),
    ("origen_homologacion", "HOMOLOGADO POR"), ("criterio", "CRITERIO"), ("descripcion", "DESCRIPCIÓN"),
    ("documento", "COMPROBANTE"),
]


def exportar(resultado, salida):
    """Escribe el libro de salida. Devuelve la ruta y la lista de CSV complementarios (si una hoja no cupo)."""
    salida = Path(salida)
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "RESUMEN"
    extras = []
    reg = _pesos_df(resultado.registros)
    nombres_sede = EF.nombres_sede(resultado)

    # ------------------------------------------------------------------ RESUMEN
    fila = 1
    _celda(ws, fila, 1, "HOMOLOGACIÓN DE COSTOS POR CENTRO DE COSTOS – ESTADO FINANCIERO HOMOLOGADO", font=Font(bold=True, size=13))
    fila += 1
    _celda(ws, fila, 1, "ESE Departamental de Primer Nivel Moreno y Clavijo", font=Font(italic=True))
    fila += 2
    op = resultado.opciones
    meta = [
        ("Generado", f"{datetime.now():%Y-%m-%d %H:%M}"),
        ("Versión de la herramienta", __version__),
        ("Archivo de referencia", resultado.referencia.ruta.name),
        ("Huella SHA-256 de la referencia", resultado.referencia.huella),
        ("Clases de cuenta cargadas", ", ".join(op.clases)),
        ("Periodo por defecto", op.periodo or "(del archivo)"),
        ("CECO nuevos", "se cargan si la regla los resuelve (--activar-nuevos)" if op.activar_nuevos else "quedan en VERIFICAR"),
        ("Resultado", "TODO CARGADO" if resultado.codigo_salida == 0 else "REVISAR: hay rechazos, retenidos o archivos con error"),
    ]
    for k, v in meta:
        _celda(ws, fila, 1, k, font=NEGRITA)
        _celda(ws, fila, 2, v)
        fila += 1
    fila += 1
    conc = resultado.conciliacion
    fila, _ = _tabla(ws, conc, [("concepto", "CONCILIACIÓN"), ("registros", "REGISTROS"), ("valor", "VALOR")], fila=fila,
                     titulo="Conciliación de valores", montos=("valor",), anchos={"concepto": 70, "registros": 12, "valor": 22})
    fila += 2
    if len(resultado.control_origen):
        fila, _ = _tabla(ws, resultado.control_origen, [
            ("origen", "ARCHIVO"), ("clase", "CLASE"), ("nombre", "NOMBRE"), ("reporte", "TOTAL SEGÚN EL REPORTE (D − C)"),
            ("leido", "TOTAL LEÍDO (D − C)"), ("diferencia", "DIFERENCIA"), ("estado", "ESTADO")], fila=fila,
            titulo="Control contra los totales del reporte de origen", montos=("reporte", "leido", "diferencia"))
        fila += 2
    fila, _ = _tabla(ws, resultado.control_referencia, [("seccion", "SECCIÓN"), ("indicador", "INDICADOR"), ("valor", "N°"),
                                                        ("porcentaje", "%")],
                     fila=fila, titulo="Control de calidad del archivo de referencia", anchos={"seccion": 70})
    ws.freeze_panes = None
    ws.auto_filter.ref = None
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 60

    # ------------------------------------------------------------------ Estados financieros
    _ef(wb.create_sheet("EF HOMOLOGADO"), EF.por_cuenta(resultado),
        "ESTADO DE RESULTADOS POR SEDE / IPS – CENTROS DE COSTOS HOMOLOGADOS A LA RESOLUCIÓN", nombres_sede, resultado,
        negritas=(0, 1, 2, -1))
    _ef(wb.create_sheet("EF POR CECO"), EF.por_ceco(resultado),
        "GASTOS Y COSTOS POR CENTRO DE COSTOS DE RESOLUCIÓN Y SEDE / IPS (CECO ACTIVOS)", nombres_sede, resultado)

    # ------------------------------------------------------------------ Carga para la herramienta de costos
    cons = _pesos_df(resultado.consolidado)
    cons = cons[cons["cuenta"].fillna("").map(EF.es_costo)] if len(cons) else cons  # la herramienta ABC costea gastos y costos
    cons["descripcion"] = cons.apply(lambda r: f"CECO E.F. {r['cecos_ef']} · {r['id_consolidado']}", axis=1) if len(cons) else ""
    cons["utilizados"] = cons["en_utilizados"].map(lambda v: "SÍ" if v else "NO") if len(cons) else ""
    _tabla(wb.create_sheet("CARGA HERRAMIENTA"), cons, [
        ("periodo", "PERIODO"), ("cuenta", "CUENTA"), ("nombre_cuenta", "NOMBRE CUENTA"), ("cod_res", "CENTRO DE COSTOS"),
        ("nombre_res", "NOMBRE CECO RESOLUCION"), ("sede_ips", "SEDE"), ("llave", "LLAVE CECO|SEDE"), ("tipo", "TIPO CECO"),
        ("valor", "VALOR"), ("descripcion", "DESCRIPCION"), ("id_consolidado", "ID CONSOLIDADO"), ("registros", "N REGISTROS"),
        ("cecos_ef", "CECO E.F. ORIGEN"), ("utilizados", "EN LISTA UTILIZADOS"),
    ], montos=("valor",))

    # ------------------------------------------------------------------ Detalle, retenidos, rechazos
    cargados = reg[reg["destino"] == CARGADO]
    _hoja_grande(wb, "DETALLE CARGADO", cargados, COLS_REGISTRO + [("id_consolidado", "ID CONSOLIDADO")], salida, extras,
                 montos=("valor",), color_por="estado")
    _hoja_grande(wb, "RETENIDOS", reg[reg["destino"] == RETENIDO], COLS_REGISTRO[:17] + [("motivo", "MOTIVO")] + COLS_REGISTRO[17:],
                 salida, extras, montos=("valor",), color_por="estado")
    rech = reg[reg["destino"].isin([RECHAZADO, EXCLUIDO])].copy()
    _hoja_grande(wb, "RECHAZOS Y EXCLUSIONES", rech, [
        ("id_registro", "ID REGISTRO"), ("destino", "DESTINO"), ("severidad", "SEVERIDAD"), ("motivo", "MOTIVO"),
        ("origen", "ARCHIVO"), ("hoja", "HOJA"), ("fila", "FILA"), ("periodo", "PERIODO"), ("cod_ceco", "COD CECO E.F."),
        ("cuenta", "CUENTA"), ("nombre_cuenta", "NOMBRE CUENTA"), ("valor", "VALOR"), ("descripcion", "DESCRIPCIÓN"),
    ], salida, extras, montos=("valor",))
    _tabla(wb.create_sheet("ALERTAS"), resultado.alertas, [
        ("tipo", "TIPO"), ("detalle", "DETALLE"), ("id_registro", "ID REGISTRO"), ("origen", "ARCHIVO"), ("hoja", "HOJA"),
        ("fila", "FILA")], anchos={"detalle": 90})
    nuevos = _pesos_df(resultado.nuevos)
    _tabla(wb.create_sheet("CECO NUEVOS"), nuevos, [
        ("cod_ceco", "Codigo"), ("nombre_ef", "Nombre"), ("cod_res", "CENTRO DE COSTOS RESOLUCION HOMOLOGADO (PROPUESTO)"),
        ("nombre_res", "NOMBRE CECO RESOLUCIÓN"), ("llave", "LLAVE HERRAMIENTA (CECO|SEDE)"), ("tipo", "TIPO CECO"),
        ("sede", "SEDE / IPS"), ("estado", "ESTADO PARA COSTEO"), ("confianza", "CONFIANZA"),
        ("criterio", "CRITERIO DE HOMOLOGACIÓN"), ("metodo", "MÉTODO"), ("registros", "REGISTROS"), ("valor", "VALOR RETENIDO"),
    ], montos=("valor",), color_por="estado")

    # ------------------------------------------------------------------ Mapeo, archivos, trazabilidad, referencia
    _tabla(wb.create_sheet("MAPEO COLUMNAS"), resultado.mapeo, [
        ("origen", "ARCHIVO"), ("hoja", "HOJA"), ("formato", "FORMATO"), ("fila_encabezado", "FILA ENCABEZADO"),
        ("campo", "CAMPO COMÚN"), ("columna_origen", "COLUMNA EN EL ARCHIVO"), ("metodo", "MÉTODO"), ("puntaje", "PUNTAJE")])
    _tabla(wb.create_sheet("ARCHIVOS"), resultado.archivos, [
        ("origen", "ARCHIVO"), ("hoja", "HOJA"), ("estado", "ESTADO"), ("motivo", "DETALLE"), ("filas", "FILAS CON DATOS"),
        ("tamano", "BYTES"), ("huella", "SHA-256")], anchos={"motivo": 90})
    _hoja_grande(wb, "TRAZABILIDAD", resultado.trazabilidad, [
        ("id_registro", "ID"), ("etapa", "ETAPA"), ("origen", "ARCHIVO"), ("hoja", "HOJA"), ("fila", "FILA"),
        ("campo", "CAMPO"), ("valor_origen", "VALOR ORIGEN"), ("valor_resultado", "VALOR RESULTADO"), ("detalle", "DETALLE"),
    ], salida, extras, anchos={"detalle": 100})
    _tabla(wb.create_sheet("CONTROL REFERENCIA"), resultado.hallazgos_referencia,
           [("tipo", "HALLAZGO"), ("codigo", "CÓDIGO"), ("detalle", "DETALLE")], anchos={"detalle": 110})
    ws = wb.create_sheet("SUPUESTOS")
    _celda(ws, 1, 1, "Supuestos y reglas aplicadas", font=Font(bold=True, size=12))
    for i, s in enumerate(SUPUESTOS, start=3):
        _celda(ws, i, 1, f"{i - 2}. {s}", alignment=Alignment(wrap_text=True, vertical="top"))
        ws.row_dimensions[i].height = 45
    ws.column_dimensions["A"].width = 140
    wb.save(salida)
    return salida, extras


def _hoja_grande(wb, nombre, df, columnas, salida, extras, **kw):
    ws = wb.create_sheet(nombre)
    _, recortada = _tabla(ws, df, columnas, **kw)
    if recortada:
        csv = salida.with_name(f"{salida.stem}_{nombre.replace(' ', '_')}.csv")
        df[[c for c, _ in columnas]].to_csv(csv, index=False, encoding="utf-8-sig", sep=";")
        extras.append(csv)
        _celda(ws, ws.max_row + 2, 1, f"La hoja supera el límite de filas de Excel; el detalle completo está en {csv.name}",
               font=Font(bold=True, color="C00000"))
